"""Controlled legacy contracts with editable projections and lossless native storage.

Unchanged OOXML (including relationships, styles and pagination) is retained.
Unsupported edits fail closed rather than silently rebuilding a different page.
No arbitrary Word parser, expressions, filesystem references or network reads.
"""
from __future__ import annotations

import base64
import io
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

from .files import validate_docx
from .renderer import _stable_zip, digest
from .schemas import FIELD_CATALOG, Layout, NativeStyle, Style, walk_blocks

TOKEN = re.compile(r"\{\{([A-Z0-9_:.-]+)\}\}", re.I)
SPECIAL = {"CONDITION_LABEL", "PRESSURE_DISPLAY", "LEFT_SIGNATURE", "RIGHT_SIGNATURE"}
RESOURCES = Path(__file__).resolve().parents[2] / "resources" / "templates"
STEMS = {"TRANSFER_ISSUE": ("transfer_issue", 3), "TRANSFER_RETURN": ("transfer_return", 3),
         "REPAIR_PROTOCOL": ("repair_protocol", 6), "PART_REQUEST": ("part_request", 3)}
KEYS = {"PARTS_USED": ("position", "part_number", "description", "quantity"),
        "REPAIR_PARTICIPANTS": ("name", "job_title", "contribution", "duration"),
        "REPAIR_ATTACHMENTS": ("filename", "stage", "description"),
        "REQUEST_LINES": ("position", "part_number", "description", "quantity"),
        "EQUIPMENT_CHECKLIST": ("position", "component", "condition")}
HEADERS = {"PARTS_USED": ("Поз.", "Номер", "Описание", "Количество"),
           "REPAIR_PARTICIPANTS": ("Три имена", "Длъжност", "Участие", "Време"),
           "REPAIR_ATTACHMENTS": ("Файл", "Етап", "Описание")}


def incompatible():
    raise ValueError("studio_native_incompatible")


def _style(paragraph, *, native_paragraph=False):
    # Presentation defaults are merely editable controls. Native inherited
    # formatting stays untouched unless a control is changed explicitly.
    run = next((r for r in paragraph.runs if r.text), None)
    f = run.font if run else paragraph.style.font
    font = f.name or paragraph.style.font.name or "Arial"
    if font not in {"Arial", "Liberation Sans", "DejaVu Sans", "DejaVu Serif", "Times New Roman", "Calibri"}:
        incompatible()
    def spacing(name):
        value = getattr(paragraph.paragraph_format, name)
        if value is None:
            value = getattr(paragraph.style.paragraph_format, name)
        return max(0, min(30, value.mm if value is not None else 0))
    size = f.size or paragraph.style.font.size
    model = NativeStyle if native_paragraph else Style
    return model(font=font, size_pt=max(1 if native_paragraph else 6, min(48, size.pt if size else 10)),
                 color="#000000", spacing_mm=spacing("space_after"), paragraph_before_mm=spacing("space_before"),
                 align={1: "center", 2: "right"}.get(paragraph.alignment, "left"),
                 bold=bool(run.bold) if run else False,
                 italic=bool(run.italic) if run else False).model_dump()


def _frame_style(properties, tag, fallback=None):
    borders = properties.find(qn("w:" + tag)) if properties is not None else None
    edge = next(iter(borders), None) if borders is not None else None
    if edge is None:
        return fallback or {"border_pt": 0, "color": "#000000"}
    width = int(edge.get(qn("w:sz"), "4")) / 8 if edge.get(qn("w:val")) not in {"nil", "none"} else 0
    color = edge.get(qn("w:color"), "000000")
    return {"border_pt": min(width, 3), "color": "#" + (color if re.fullmatch(r"[a-fA-F0-9]{6}", color) else "000000")}


def _project(document, prefix, language="bg"):
    blocks, targets, index = [], {}, 0

    def project(container, region):
        nonlocal index
        result = []
        root = container._element
        elements = root.body if hasattr(root, "body") else root
        for element in elements:
            if element.tag not in {qn("w:p"), qn("w:tbl")}:
                if element.tag not in {qn("w:sectPr"), qn("w:tcPr")}:
                    incompatible()
                continue
            index += 1
            ident = f"{prefix}_{index}"
            targets[ident] = element
            if element.tag == qn("w:p"):
                p = Paragraph(element, container)
                if p._p.xpath(".//w:drawing | .//w:pict"):
                    extents = p._p.xpath(".//wp:extent")
                    if p.text.strip() or len(extents) != 1 or p._p.xpath(".//w:pict"):
                        incompatible()
                    result.append({"id": ident, "type": "native_image", "region": region,
                        "width_mm": int(extents[0].get("cx")) / 36000,
                        "height_mm": int(extents[0].get("cy")) / 36000, "style": _style(p)})
                    continue
                if p._p.xpath(".//w:instrText"):
                    result.append({"id": ident, "type": "page_number", "region": region, "style": _style(p)})
                    continue
                match = re.fullmatch(r"\s*\{\{TABLE:([A-Z_]+)\}\}\s*", p.text)
                if match:
                    code = match[1]
                    if code not in HEADERS:
                        incompatible()
                    from ..documents.repair_table_contract import empty_label, headers
                    keys, labels = KEYS[code], headers(code, language)
                    table_style = _style(p)
                    table_style.update(border_pt=.5, color="#000000", width_percent=100)
                    result.append({"id": ident, "type": "dynamic_table", "source": code, "style": table_style,
                        "native_properties": True, "repeat_header": False,
                        "empty_label": empty_label(code, language),
                        "columns": [{"key": k, "label": v, "width_percent": 100 / len(keys)} for k, v in zip(keys, labels, strict=True)]})
                else:
                    result.append({"id": ident, "type": "native_paragraph", "text": p.text,
                                   "region": region, "style": _style(p, native_paragraph=True)})
            else:
                table = Table(element, container)
                all_text = "\n".join(p.text for row in table.rows for c in row.cells for p in c.paragraphs)
                code = "EQUIPMENT_CHECKLIST" if "{{CHECK_1}}" in all_text else "REQUEST_LINES" if "{{TABLE:REQUEST_LINES}}" in all_text else None
                if code:
                    labels = [c.text for c in table.rows[0].cells]
                    keys = KEYS[code] if len(labels) == 4 or code != "REQUEST_LINES" else ("position", "description", "quantity")
                    widths = [c.width or 1 for c in table.columns]
                    table_style = _style(table.rows[1].cells[0].paragraphs[0])
                    table_style.update(border_pt=.5, color="#000000", width_percent=100)
                    result.append({"id": ident, "type": "dynamic_table", "source": code, "style": table_style,
                        "native_properties": True, "repeat_header": bool(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")),
                        "columns": [{"key": k, "label": label, "width_percent": float(w / sum(widths) * 100)} for k, label, w in zip(keys, labels, widths, strict=True)]})
                    continue
                widths = [float(c.width or 1) for c in table.columns]
                frame = _frame_style(table._tbl.tblPr, "tblBorders", {"border_pt": .5 if table.style and table.style.style_id == "TableGrid" else 0, "color": "#000000"})
                rows = []
                for row in table.rows:
                    cells = []
                    for tc in row._tr.tc_lst:
                        if tc.tcPr is not None and tc.tcPr.find(qn("w:vMerge")) is not None:
                            incompatible()
                        from docx.table import _Cell
                        cell = _Cell(tc, table)
                        cells.append({"colspan": tc.grid_span, "blocks": project(cell, "body"),
                            "padding_mm": None, "vertical_align": {0: "top", 1: "center", 3: "bottom"}.get(cell.vertical_alignment, "top"),
                            "style": _frame_style(tc.tcPr, "tcBorders", frame)})
                    rows.append({"cells": cells, "min_height_mm": min(60, row.height.mm if row.height else 0),
                        "keep_together": bool(row._tr.xpath("./w:trPr/w:cantSplit"))})
                result.append({"id": ident, "type": "grid", "region": region,
                    "widths": [w / sum(widths) * 100 for w in widths], "grid_rows": rows,
                    "style": {**frame, "spacing_mm": 0, "align": {1: "center", 2: "right"}.get(table.alignment, "left")}})
        return result

    blocks.extend(project(document, "body"))
    section = document.sections[0]
    blocks.extend(project(section.header, "header"))
    blocks.extend(project(section.footer, "footer"))
    return blocks, targets


def _request_source(source, language, free):
    from ..documents import company_header, request_layout_v2
    from ..documents.request_issue_documents import _meta, _rows
    from ..documents.request_machine_section import rows as machine_rows
    data = {"language": language, "request_kind": "FREE_FORM" if free else "CATALOG",
        "document_layout_version": 2 if free else 3, "department": "{{DEPARTMENT}}",
        "requested_by_name": "{{REQUESTER_NAME}}", "reason": "{{REMARKS}}",
        "repair_reference": "{{REPAIR_REFERENCE}}", "official_number": "{{DOCUMENT_NUMBER}}",
        "approved_by_name": "{{APPROVER_NAME}}", "lines": [], "machine": None}
    rows = _rows(data) + [["{{TABLE:REQUEST_LINES}}"] + [""] * (2 if free else 3)]
    machine = None if free else {"name": "{{MACHINE_NAME}}", "inventory_number": "{{MACHINE_NUMBER}}",
        "brand": "{{BRAND}}", "model": "{{MODEL}}", "serial_number": "{{SERIAL_NUMBER}}",
        "pressure_bar": None, "batch_reference": "{{BATCH_REFERENCE}}"}
    raw = request_layout_v2.build_docx(data, source, rows=rows, meta=_meta(data),
        date="{{CREATION_DATE}}", resources=company_header.current_resources(), machine_section=machine)
    doc = Document(io.BytesIO(raw))
    if machine:
        # Preserve native numeric formatting through a dedicated typed binding.
        assert len(machine_rows(machine, language)) == 4
        doc.tables[0].cell(2, 3).text = "{{PRESSURE_DISPLAY}}"
        request_layout_v2.paragraph_style(doc.tables[0].cell(2, 3).paragraphs[0], size=10)
    out = io.BytesIO()
    doc.save(out)
    return _stable_zip(out.getvalue())


def convert(version):
    from ..template_engine import source_bytes
    typ, language = version.template.document_type, version.language
    if typ not in STEMS or language not in {"bg", "en", "ru"}:
        incompatible()
    source = source_bytes(version)
    stem, number = STEMS[typ]
    known = RESOURCES / f"{stem}-{language}-v{number}.docx"
    if (not known.is_file() or digest(source) != digest(known.read_bytes())
            or digest(source) != version.source_sha256 or version.layout_contract.get("reference_only")):
        incompatible()
    validate_docx(source)
    variants = {"DEFAULT": source} if typ != "PART_REQUEST" else {
        "MACHINE": _request_source(source, language, False),
        "FREE_FORM": _request_source(source, language, True)}
    blocks = []
    for name, raw in sorted(variants.items()):
        projected, _ = _project(Document(io.BytesIO(raw)), name.lower(), language)
        blocks.extend(projected)
    native = {"adapter": 1, "document_type": typ, "language": language,
        "source_version_id": version.id, "source_version": version.version,
        "source_sha256": version.source_sha256,
        "sources": {k: base64.b64encode(v).decode() for k, v in variants.items()},
        "source_hashes": {k: digest(v) for k, v in variants.items()}}
    section = Document(io.BytesIO(next(iter(variants.values())))).sections[0]
    page = {f"{edge}_mm": getattr(section, f"{edge}_margin").mm for edge in ("top", "bottom", "left", "right")}
    page.update(margin_mm=section.left_margin.mm, header_distance_mm=section.header_distance.mm,
                footer_distance_mm=section.footer_distance.mm)
    return Layout.model_validate({"schema_version": 4, "page": page, "blocks": blocks, "native": native})


def validate(layout, document_type):
    native = layout.native
    if (layout.schema_version != 4 or not isinstance(native, dict) or set(native) != {"adapter", "document_type", "language",
            "source_version_id", "source_version", "source_sha256", "sources", "source_hashes"}
            or native["adapter"] != 1 or native["document_type"] != document_type
            or native["language"] not in {"bg", "en", "ru"}
            or set(native["sources"]) != ({"MACHINE", "FREE_FORM"} if document_type == "PART_REQUEST" else {"DEFAULT"})):
        incompatible()
    actual = {b.id: b for b in walk_blocks(layout.blocks)}
    expected = {}
    for name, encoded in native["sources"].items():
        if not isinstance(encoded, str) or len(encoded) > 4_000_000:
            incompatible()
        try:
            raw = base64.b64decode(encoded, validate=True)
        except ValueError:
            incompatible()
        if digest(raw) != native["source_hashes"].get(name):
            incompatible()
        validate_docx(raw)
        projected, _ = _project(Document(io.BytesIO(raw)), name.lower(), native["language"])
        expected.update({b.id: b for b in walk_blocks(Layout.model_validate({"blocks": projected}).blocks)})
    if actual.keys() != expected.keys():
        incompatible()
    for ident, block in actual.items():
        original = expected[ident]
        if block.type != original.type or block.region != original.region or block.request_variant != original.request_variant:
            incompatible()
        if block.type == "native_paragraph":
            # Typed tokens only; slots and canonical fields cannot be deleted,
            # duplicated or moved to another paragraph/cell.
            tokens = TOKEN.findall(block.text)
            if sorted(tokens) != sorted(TOKEN.findall(original.text)) or "[[ASSETCORE_" in block.text:
                incompatible()
            if any(t not in FIELD_CATALOG and t not in SPECIAL and not t.startswith("TEMPLATE_LANGUAGE:") for t in tokens):
                incompatible()
            if "{{" in TOKEN.sub("", block.text) or "}}" in TOKEN.sub("", block.text):
                incompatible()
        elif block.type == "grid":
            if block.repeat_rows != original.repeat_rows or block.keep_together != original.keep_together:
                incompatible()
            if len(block.grid_rows) != len(original.grid_rows):
                incompatible()
            for row, old in zip(block.grid_rows, original.grid_rows, strict=True):
                if row.page_break_before != old.page_break_before:
                    incompatible()
                if len(row.cells) != len(old.cells):
                    incompatible()
                for cell, before in zip(row.cells, old.cells, strict=True):
                    if cell.colspan != before.colspan or [b.id for b in cell.blocks] != [b.id for b in before.blocks]:
                        incompatible()
        elif block.type == "dynamic_table":
            if not block.native_properties:
                incompatible()
            if block.source != original.source or [c.key for c in block.columns] != [c.key for c in original.columns]:
                incompatible()
            if block.group_by is not None:
                incompatible()
            if any("[[ASSETCORE_" in c.label or "{{" in c.label for c in block.columns):
                raise ValueError("studio_reserved_marker")
            if "[[ASSETCORE_" in block.empty_label or "{{" in block.empty_label:
                raise ValueError("studio_reserved_marker")
    if [b.id for b in layout.blocks] != [b.id for b in _expected_roots(native)]:
        incompatible()


def _expected_roots(native):
    for name, raw in sorted(native["sources"].items()):
        blocks, _ = _project(Document(io.BytesIO(base64.b64decode(raw))), name.lower(), native["language"])
        yield from Layout.model_validate({"blocks": blocks}).blocks


def variant(data):
    free = data.get("request_kind") == "FREE_FORM"
    machine = data.get("machine")
    if free and machine is not None or not free and not machine:
        raise ValueError("studio_request_variant_mismatch")
    if data.get("document_layout_version") not in (None, 2 if free else 3):
        raise ValueError("studio_request_variant_mismatch")
    return "FREE_FORM" if free else "MACHINE"


def image_bytes(layout, block_id):
    if layout.native is None:
        incompatible()
    for name, encoded in layout.native["sources"].items():
        document = Document(io.BytesIO(base64.b64decode(encoded, validate=True)))
        blocks, targets = _project(document, name.lower(), layout.native["language"])
        found = next((b for b in walk_blocks(Layout.model_validate({"blocks": blocks}).blocks)
                      if b.id == block_id and b.type == "native_image"), None)
        if found is None:
            continue
        element = targets[block_id]
        root = element.getroottree().getroot()
        part = document.sections[0].header.part if root.tag == qn("w:hdr") else document.sections[0].footer.part if root.tag == qn("w:ftr") else document.part
        blips = element.xpath(".//a:blip")
        if len(blips) != 1:
            incompatible()
        image = part.related_parts[blips[0].get(qn("r:embed"))]
        if image.content_type not in {"image/png", "image/jpeg"}:
            incompatible()
        return image.blob, image.content_type
    incompatible()


def _property(parent, tag, attributes):
    for old in list(parent.findall(qn("w:" + tag))):
        parent.remove(old)
    node = OxmlElement("w:" + tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    parent.append(node)
    return node


def _replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _paragraph_delta(paragraph, current, old, width=188):
    if current.align != old.align:
        paragraph.alignment = {"left": 0, "center": 1, "right": 2}[current.align]
    for name, attr in (("spacing_mm", "space_after"), ("paragraph_before_mm", "space_before")):
        if getattr(current, name) != getattr(old, name):
            setattr(paragraph.paragraph_format, attr, Mm(getattr(current, name)))
    for run in paragraph.runs:
        if current.font != old.font:
            run.font.name = current.font
        if current.size_pt != old.size_pt:
            run.font.size = Pt(current.size_pt)
        if current.color != old.color:
            run.font.color.rgb = RGBColor.from_string(current.color[1:])
        if current.bold != old.bold:
            run.bold = current.bold
        if current.italic != old.italic:
            run.italic = current.italic
    if current.background != old.background:
        _property(paragraph._p.get_or_add_pPr(), "shd", {"fill": current.background.lstrip("#") if current.background != "transparent" else "auto"})
    if current.border_pt != old.border_pt:
        borders = _property(paragraph._p.get_or_add_pPr(), "pBdr", {})
        for edge in ("top", "left", "bottom", "right"):
            _property(borders, edge, {"val": "single" if current.border_pt else "nil", "sz": max(1, round(current.border_pt * 8)), "color": current.color[1:]})
    if current.width_percent != old.width_percent:
        unused = width * (1 - current.width_percent / 100)
        paragraph.paragraph_format.left_indent = Mm(unused if current.align == "right" else unused / 2 if current.align == "center" else 0)
        paragraph.paragraph_format.right_indent = Mm(unused if current.align == "left" else unused / 2 if current.align == "center" else 0)


def _cell_properties(cell, current, previous):
    props = cell._tc.get_or_add_tcPr()
    if current.padding_mm is not None and current.padding_mm != previous.padding_mm:
        margin = _property(props, "tcMar", {})
        for edge in ("top", "left", "bottom", "right"):
            _property(margin, edge, {"w": round(current.padding_mm / 25.4 * 1440), "type": "dxa"})
    if current.vertical_align != previous.vertical_align:
        _property(props, "vAlign", {"val": current.vertical_align})
    if current.style.background != previous.style.background:
        _property(props, "shd", {"fill": current.style.background.lstrip("#") if current.style.background != "transparent" else "auto"})
    if (current.style.border_pt, current.style.color) != (previous.style.border_pt, previous.style.color):
        borders = _property(props, "tcBorders", {})
        for edge in ("top", "left", "bottom", "right"):
            _property(borders, edge, {"val": "single" if current.style.border_pt else "nil",
                "sz": round(current.style.border_pt * 8), "color": current.style.color.lstrip("#")})


def _table_properties(element, block, old, document):
    from .renderer import _table_borders
    table = Table(element, document)
    # Missing native tblGrid is legal in the legacy paragraph-table renderer.
    grid = element.find(qn("w:tblGrid"))
    total = sum(int(c.get(qn("w:w"))) for c in grid) if grid is not None else round(188 / 25.4 * 1440)
    widths = block.widths if block.type == "grid" else [c.width_percent for c in block.columns]
    before = old.widths if old.type == "grid" else [c.width_percent for c in old.columns]
    if widths != before or block.style.width_percent != old.style.width_percent:
        if grid is None:
            grid = OxmlElement("w:tblGrid")
            element.insert(1, grid)
        for child in list(grid):
            grid.remove(child)
        actual = total * block.style.width_percent / old.style.width_percent
        for width in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(round(actual * width / 100)))
            grid.append(column)
        for row in table.rows:
            column_index = 0
            from docx.table import _Cell
            for tc in row._tr.tc_lst:
                cell = _Cell(tc, table)
                cell.width = Mm(actual * sum(widths[column_index:column_index + tc.grid_span]) / 100 * 25.4 / 1440)
                column_index += tc.grid_span
    if block.style != old.style:
        if block.style.align != old.style.align:
            _property(table._tbl.tblPr, "jc", {"val": block.style.align})
        if block.style.border_pt != old.style.border_pt or block.style.color != old.style.color:
            _table_borders(table, block.style)
        if block.type == "dynamic_table":
            for node in element.iter(qn("w:p")):
                paragraph = Paragraph(node, document)
                _paragraph_delta(paragraph, block.style, old.style)
        if block.style.spacing_mm != old.style.spacing_mm or block.style.paragraph_before_mm != old.style.paragraph_before_mm:
            gap = OxmlElement("w:p")
            element.addnext(gap)
            p = Paragraph(gap, document)
            p.paragraph_format.space_after = Mm(block.style.spacing_mm)
            p.paragraph_format.space_before = Mm(block.style.paragraph_before_mm)
            p.paragraph_format.line_spacing = Pt(1)
    if block.type == "grid":
        from docx.table import _Cell
        for row, current, previous in zip(table.rows, block.grid_rows, old.grid_rows, strict=True):
            if current.page_break_before != previous.page_break_before:
                incompatible()
            if current.min_height_mm != previous.min_height_mm:
                _property(row._tr.get_or_add_trPr(), "trHeight", {"val": round(current.min_height_mm / 25.4 * 1440), "hRule": "atLeast"})
            if current.keep_together != previous.keep_together:
                _property(row._tr.get_or_add_trPr(), "cantSplit", {"val": "1" if current.keep_together else "0"})
            for tc, cell, before_cell in zip(row._tr.tc_lst, current.cells, previous.cells, strict=True):
                _cell_properties(_Cell(tc, table), cell, before_cell)
    else:
        for index, tr in enumerate(element.findall(qn("w:tr"))):
            if index == 0 and block.repeat_header != old.repeat_header:
                _property(tr.get_or_add_trPr(), "tblHeader", {"val": "1" if block.repeat_header else "0"})
            if block.keep_rows_together != old.keep_rows_together:
                _property(tr.get_or_add_trPr(), "cantSplit", {"val": "1" if block.keep_rows_together else "0"})
            if block.min_height_mm is not None and block.min_height_mm != old.min_height_mm:
                _property(tr.get_or_add_trPr(), "trHeight", {"val": round(block.min_height_mm / 25.4 * 1440), "hRule": "atLeast"})
            if block.cell_padding_mm is not None and block.cell_padding_mm != old.cell_padding_mm:
                for tc in tr.tc_lst:
                    margin = _property(tc.get_or_add_tcPr(), "tcMar", {})
                    for edge in ("top", "left", "bottom", "right"):
                        _property(margin, edge, {"w": round(block.cell_padding_mm / 25.4 * 1440), "type": "dxa"})


def render(manifest, values, tables, *, source=False):
    from ..template_engine import _insert_table_after, _replace_paragraph
    layout = Layout.model_validate(manifest["layout"])
    validate(layout, manifest["document_type"])
    native = layout.native
    if native["language"] != manifest["language"]:
        incompatible()
    name = values.get("REQUEST_VARIANT") if manifest["document_type"] == "PART_REQUEST" else "DEFAULT"
    if name not in native["sources"]:
        raise ValueError("studio_request_variant_mismatch")
    raw = base64.b64decode(native["sources"][name])
    document = Document(io.BytesIO(raw))
    section = document.sections[0]
    section.page_width, section.page_height = (Mm(210), Mm(297)) if layout.page.orientation == "portrait" else (Mm(297), Mm(210))
    for edge in ("top", "bottom", "left", "right"):
        value = getattr(layout.page, edge + "_mm")
        if value is not None:
            setattr(section, edge + "_margin", Mm(value))
    for region in ("header", "footer"):
        value = getattr(layout.page, region + "_distance_mm")
        if value is not None:
            setattr(section, region + "_distance", Mm(value))
    projected, targets = _project(document, name.lower(), manifest["language"])
    original = {b.id: b for b in walk_blocks(Layout.model_validate({"blocks": projected}).blocks)}
    actual = {b.id: b for b in walk_blocks(layout.blocks)}
    bindings = {**values, "PRESSURE_DISPLAY": f"{values['PRESSURE_BAR']:g} bar" if isinstance(values.get("PRESSURE_BAR"), (int, float)) else str(values.get("PRESSURE_BAR") or "")}
    for key, value in bindings.items():
        if "[[ASSETCORE_" in str(value) and (key not in {"LEFT_SIGNATURE", "RIGHT_SIGNATURE"} or value != "[[ASSETCORE_" + key + "]]"):
            raise ValueError("studio_reserved_marker")
        if not source and "{{" in str(value):
            raise ValueError("studio_reserved_marker")
    for ident, old in original.items():
        block, element = actual[ident], targets[ident]
        if block.type in {"native_paragraph", "page_number"}:
            paragraph = Paragraph(element, document)
            if block.type == "native_paragraph" and block.text != old.text:
                _replace_text(paragraph, block.text)
            if block.style != old.style:
                _paragraph_delta(paragraph, block.style, old.style)
            if not source and block.type == "native_paragraph":
                optional = next((f for f in ("REMARKS", "REPAIR_REFERENCE") if "{{" + f + "}}" in block.text), None)
                if name != "DEFAULT" and optional and not bindings.get(optional):
                    element.getparent().remove(element)
                    continue
                _replace_paragraph(paragraph, bindings)
        elif block.type == "native_image":
            paragraph = Paragraph(element, document)
            if block.image_id is not None:
                image = manifest["images"].get(str(block.image_id))
                if image is None:
                    incompatible()
                for child in list(element):
                    if child.tag != qn("w:pPr"):
                        element.remove(child)
                paragraph.add_run().add_picture(io.BytesIO(base64.b64decode(image["content"], validate=True)), width=Mm(block.width_mm), height=Mm(block.height_mm))
            elif (block.width_mm, block.height_mm) != (old.width_mm, old.height_mm):
                for extent in element.xpath(".//wp:extent | .//a:xfrm/a:ext"):
                    extent.set("cx", str(round(block.width_mm * 36000)))
                    extent.set("cy", str(round(block.height_mm * 36000)))
            if block.style != old.style:
                _paragraph_delta(paragraph, block.style, old.style)
        elif block.type == "grid":
            _table_properties(element, block, old, document)
        elif block.type == "dynamic_table" and not source:
            raw_rows = (tables or {}).get(block.source)
            if raw_rows is None or len(raw_rows) > 501:
                incompatible()
            keys = KEYS[block.source]
            records = raw_rows if not raw_rows or isinstance(raw_rows[0], dict) else [dict(zip(keys, row, strict=True)) for row in raw_rows[1:]]
            if not records and block.empty_label:
                records = [{block.columns[0].key: block.empty_label}]
            native_components = None
            if block.source == "EQUIPMENT_CHECKLIST":
                from ..documents.transfer_documents import CHECKLIST
                # The approved v3 wording differs slightly from the UI's
                # canonical HPWJ labels. Keep that verified print vocabulary
                # only for the complete matching HPWJ checklist. Other category
                # snapshots always retain their recorded labels and length.
                if [r.get("component") for r in records] == CHECKLIST[manifest["language"]]:
                    native_components = [r.cells[1].text for r in Table(element, document).rows[1:]]
            rows = [[c.label for c in block.columns]]
            for i, record in enumerate(records, 1):
                record = dict(record)
                if any("[[ASSETCORE_" in str(v) or "{{" in str(v) for v in record.values()):
                    raise ValueError("studio_reserved_marker")
                if block.source == "EQUIPMENT_CHECKLIST":
                    if native_components is not None:
                        record["component"] = native_components[i - 1]
                    record["position"] = f"{i}."
                    record["condition"] = " · ".join(str(record.get(k) or "") for k in ("condition", "measurement", "note") if record.get(k) not in (None, ""))
                if "quantity" in record and record.get("unit"):
                    record["quantity"] = f"{record['quantity']} {record['unit']}"
                rows.append([str(record[c.key]) if record.get(c.key) is not None else "" for c in block.columns])
            if element.tag == qn("w:p"):
                parent, position = element.getparent(), element.getparent().index(element)
                _insert_table_after(Paragraph(element, document), rows)
                element = parent[position]
            else:
                table = Table(element, document)
                originals = list(table.rows)[1:]
                if not originals:
                    incompatible()
                for cell, label in zip(table.rows[0].cells, rows[0], strict=True):
                    if cell.text != label:
                        _replace_text(cell.paragraphs[0], label)
                for i, row_data in enumerate(rows[1:]):
                    copied = deepcopy(originals[min(i, len(originals) - 1)]._tr)
                    if block.source == "REQUEST_LINES":
                        _property(copied.get_or_add_trPr(), "cantSplit", {"val": "1" if all(len(v) < 1500 and v.count("\n") < 30 for v in row_data) else "0"})
                    table._tbl.append(copied)
                    row = _Row(copied, table)
                    for cell, value in zip(row.cells, row_data, strict=True):
                        paragraph = cell.paragraphs[0]
                        if paragraph.runs:
                            paragraph.runs[0].text = value
                            for run in paragraph.runs[1:]:
                                run.text = ""
                        else:
                            paragraph.add_run(value)
                for row in originals:
                    table._tbl.remove(row._tr)
            _table_properties(element, block, old, document)
            if block.source == "REQUEST_LINES":
                for tr, row_data in zip(element.findall(qn("w:tr"))[1:], rows[1:], strict=True):
                    if any(len(v) >= 1500 or v.count("\n") >= 30 for v in row_data):
                        _property(tr.get_or_add_trPr(), "cantSplit", {"val": "0"})
    out = io.BytesIO()
    document.save(out)
    result = _stable_zip(out.getvalue(), manifest)
    validate_docx(result)
    return result
