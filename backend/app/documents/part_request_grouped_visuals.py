"""One official visual page per captured role, artifact and source page."""

from __future__ import annotations

import hashlib
import io
from collections import defaultdict
from pathlib import PureWindowsPath
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Mm, Pt
from pydantic import ValidationError
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, Spacer
from sqlalchemy.orm import Session

from ..models import PartRequest, PartVisualArtifact
from ..part_requests.visual_schemas import VisualSnapshotOut
from ..part_requests.visual_snapshots import _integrity_error, snapshot_response
from .common import _language
from .part_request_visual_appendix import (
    LABELS,
    MAX_IMAGE_BYTES,
    MAX_SOURCE_BYTES,
    MAX_SOURCE_TOTAL_BYTES,
    MAX_VISUAL_BLOCKS,
    VisualBlock,
    VisualLine,
    VisualPage,
    VisualPlan,
    _render_page,
)
from .rendering import _pdf_styles, _set_run_font

ROLE_ORDER = {"EXPLODED_SCHEME": 0, "SPARE_PARTS_LIST": 1, "LEGACY_UNCLASSIFIED": 2}


def _key(reference: dict) -> tuple[str, str, int]:
    role = reference.get("visual_role") or reference["source_metadata"].get("visual_role")
    if role is None:
        role = "LEGACY_UNCLASSIFIED"
    if role not in ROLE_ORDER:
        raise _integrity_error()
    return role, reference["artifact_sha256"], reference["page_number"]


def _source_bytes(db: Session, digest: str, cache: dict[str, bytes]) -> bytes:
    if digest not in cache:
        artifact = db.get(PartVisualArtifact, digest)
        if artifact is None or artifact.byte_length > MAX_SOURCE_BYTES:
            raise _integrity_error()
        if (artifact.byte_length != len(artifact.content)
                or hashlib.sha256(artifact.content).hexdigest() != digest):
            raise _integrity_error()
        if sum(map(len, cache.values())) + artifact.byte_length > MAX_SOURCE_TOTAL_BYTES:
            raise _integrity_error()
        cache[digest] = artifact.content
    return cache[digest]


def prepare_appendix(db: Session, request: PartRequest) -> VisualPlan:
    """Verify every immutable snapshot before rendering or writing official rows."""
    grouped: dict[tuple[str, str, int], list[dict]] = defaultdict(list)
    line_values: list[dict] = []
    line_keys: dict[int, dict[tuple[str, str, int], list[dict]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for sequence, line in enumerate(sorted(request.lines, key=lambda item: item.id), 1):
        try:
            evidence = snapshot_response(line)
        except (ValueError, TypeError, AttributeError, KeyError) as exc:
            raise _integrity_error() from exc
        snapshot = evidence["snapshot"]
        catalog = snapshot["catalog"] if snapshot else {}
        if snapshot:
            try:
                VisualSnapshotOut.model_validate(snapshot)
            except ValidationError as exc:
                raise _integrity_error() from exc
            if snapshot["line_id"] != line.id or snapshot["schema_version"] != 1:
                raise _integrity_error()
            occurrences = snapshot["visual_references"]
            if sorted(value["ordinal"] for value in occurrences) != list(
                range(1, snapshot["occurrence_count"] + 1)
            ):
                raise _integrity_error()
            references = [
                {**value, "geometry": {
                    name: value[name] for name in ("x", "y", "width", "height")
                }}
                for value in occurrences
            ] + [
                {**value, "ordinal": None, "geometry": None}
                for value in catalog.get("visual_pages", [])
            ]
            for reference in references:
                key = _key(reference)
                contribution = {
                    "line_id": line.id,
                    "line_sequence": sequence,
                    "part_number": str(catalog.get("requested_part_number") or catalog.get("part_number") or ""),
                    "position": str(catalog.get("position") or ""),
                    "snapshot_id": snapshot["id"],
                    "snapshot_sha256": snapshot["sha256"],
                    "occurrence_ordinal": reference["ordinal"],
                    "geometry": reference["geometry"],
                    "source_kind": reference.get("source_kind"),
                    "hotspot_id": reference.get("hotspot_id"),
                    "catalog_visual_source_id": reference.get("catalog_visual_source_id")
                    or reference["source_metadata"].get("catalog_visual_source_id"),
                    "catalog_visual_part_map_id": reference.get("catalog_visual_part_map_id"),
                    "source_metadata": reference["source_metadata"],
                }
                grouped[key].append(contribution)
                line_keys[line.id][key].append(contribution)
        line_values.append({
            "line_id": line.id,
            "sequence": sequence,
            "part_number": str(catalog.get("requested_part_number") or catalog.get("part_number") or line.part_number or ""),
            "description": str(catalog.get("description") or line.description),
            "position": str(catalog.get("position") or line.position or ""),
            "state": evidence["state"],
            "snapshot_id": snapshot["id"] if snapshot else None,
            "snapshot_sha256": snapshot["sha256"] if snapshot else None,
            "captured_at": snapshot["captured_at"] if snapshot else None,
        })

    if len(grouped) > MAX_VISUAL_BLOCKS:
        raise _integrity_error()
    artifacts: dict[str, bytes] = {}
    pages: list[VisualPage] = []
    total_image_bytes = 0
    for role, digest, page_number in sorted(
        grouped, key=lambda key: (ROLE_ORDER[key[0]], key[1], key[2])
    ):
        contributions = sorted(
            grouped[(role, digest, page_number)],
            key=lambda value: (value["line_sequence"],
                               value["occurrence_ordinal"] or 0,
                               value["hotspot_id"] or 0),
        )
        media_types = {value["source_metadata"]["media_type"] for value in contributions}
        if len(media_types) != 1:
            raise _integrity_error()
        source = _source_bytes(db, digest, artifacts)
        if any(value["source_metadata"]["byte_length"] != len(source)
               for value in contributions):
            raise _integrity_error()
        geometry_markers = {}
        for contribution in contributions:
            geometry = contribution["geometry"]
            if geometry is not None:
                geometry_key = tuple(geometry[name] for name in ("x", "y", "width", "height"))
                if geometry_key not in geometry_markers:
                    geometry_markers[geometry_key] = len(geometry_markers) + 1
                contribution["marker"] = geometry_markers[geometry_key]
            else:
                contribution["marker"] = None
        render_references = [
            {"ordinal": marker, **dict(zip(("x", "y", "width", "height"), geometry, strict=True))}
            for geometry, marker in geometry_markers.items()
        ]
        image, width, height = _render_page(
            source, next(iter(media_types)), page_number, render_references
        )
        total_image_bytes += len(image)
        if total_image_bytes > MAX_IMAGE_BYTES:
            raise _integrity_error()
        filenames = dict.fromkeys(
            PureWindowsPath(value["source_metadata"]["filename"]).name
            for value in contributions if value["source_metadata"].get("filename")
        )
        revisions = dict.fromkeys(
            value for item in contributions
            if (value := item["source_metadata"].get("revision_label")
                or item["source_metadata"].get("document_revision"))
        )
        pages.append(VisualPage(
            role, digest, page_number, image, width, height,
            " / ".join(filenames) or None,
            " / ".join(revisions) or None,
            tuple(contributions),
        ))
    page_index = {(p.role, p.artifact_sha256, p.page_number): p for p in pages}
    lines = []
    for value in line_values:
        blocks = []
        for key, contributions in sorted(line_keys[value["line_id"]].items(),
                                         key=lambda item: (ROLE_ORDER[item[0][0]], item[0][1], item[0][2])):
            page = page_index[key]
            blocks.append(VisualBlock(
                page.role, page.artifact_sha256, page.page_number,
                tuple(item["occurrence_ordinal"] for item in contributions
                      if item["occurrence_ordinal"] is not None),
                page.image, page.width, page.height, page.filename, page.revision,
                tuple(item["position"] for item in contributions
                      if item["occurrence_ordinal"] is not None),
            ))
        lines.append(VisualLine(**value, blocks=tuple(blocks)))
    return VisualPlan(tuple(lines), tuple(pages))


def appendix_manifest(plan: VisualPlan) -> dict:
    return {
        "renderer_version": 2,
        "lines": [
            {
                "line_id": line.line_id,
                "visual_snapshot_id": line.snapshot_id,
                "visual_snapshot_sha256": line.snapshot_sha256,
                "blocks": [
                    {"visual_role": block.role, "artifact_sha256": block.artifact_sha256,
                     "page_number": block.page_number,
                     "occurrence_ordinals": list(block.ordinals)}
                    for block in line.blocks
                ],
            }
            for line in plan.lines
        ],
        "pages": [
            {"visual_role": page.role, "artifact_sha256": page.artifact_sha256,
             "page_number": page.page_number, "contributions": list(page.contributions)}
            for page in plan.pages
        ],
    }


def _page_heading(page: VisualPage, labels: dict[str, str]) -> str:
    return labels[{
        "EXPLODED_SCHEME": "scheme",
        "SPARE_PARTS_LIST": "list",
        "LEGACY_UNCLASSIFIED": "unclassified",
    }[page.role]]


def _page_caption(page: VisualPage, labels: dict[str, str]) -> str:
    markers: dict[int, dict[str, None]] = {}
    positions: dict[str, None] = {}
    for value in page.contributions:
        position = value["position"] or value["part_number"]
        if not position:
            continue
        if value["marker"] is None:
            positions[position] = None
        else:
            markers.setdefault(value["marker"], {})[position] = None
    parts = [
        f"{labels['markers']}: " + ", ".join(
            f"{marker}: {' / '.join(values)}" for marker, values in markers.items()
        )
    ] if markers else []
    marked = {position for values in markers.values() for position in values}
    unmarked = [position for position in positions if position not in marked]
    if unmarked:
        parts.append(f"{labels['position']}: {', '.join(unmarked)}")
    return " · ".join(parts)


def append_docx(docx: bytes, plan: VisualPlan, language: str) -> bytes:
    if not plan.lines:
        return docx
    labels = LABELS[_language(language)]
    document = Document(io.BytesIO(docx))
    section = document.sections[-1]
    max_width = min(Mm(188), section.page_width - section.left_margin - section.right_margin)
    max_height = min(Mm(205), section.page_height - section.top_margin - section.bottom_margin - Mm(55))
    unavailable = [line for line in plan.lines if not line.blocks]
    if unavailable:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        _set_run_font(heading.add_run(labels["unavailable"]), 9, True)
        for line in unavailable:
            note = document.add_paragraph()
            note.paragraph_format.space_after = Pt(2)
            _set_run_font(note.add_run(
                f"{labels['line']} {line.sequence}: {labels[line.state]}"
            ), 8)
    for page in plan.pages:
        document.add_page_break()
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(5)
        heading.paragraph_format.keep_with_next = True
        _set_run_font(heading.add_run(_page_heading(page, labels)), 11, True)
        caption_text = _page_caption(page, labels)
        if caption_text:
            caption = document.add_paragraph()
            caption.paragraph_format.space_after = Pt(5)
            caption.paragraph_format.keep_with_next = True
            _set_run_font(caption.add_run(caption_text), 8)
        ratio = page.width / page.height
        width = min(max_width, int(max_height * ratio))
        height = min(max_height, int(max_width / ratio))
        document.add_paragraph().add_run().add_picture(
            io.BytesIO(page.image), width=width, height=height
        )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_flowables(plan: VisualPlan, language: str) -> list:
    labels = LABELS[_language(language)]
    _, _, _, _, title, small = _pdf_styles()
    flowables = []
    if any(not line.blocks for line in plan.lines):
        flowables.append(Paragraph(labels["unavailable"], small))
    for line in plan.lines:
        if not line.blocks:
            flowables.append(Paragraph(
                f"{labels['line']} {line.sequence}: {labels[line.state]}", small
            ))
    for page in plan.pages:
        flowables.extend([
            PageBreak(), Paragraph(_page_heading(page, labels), title), Spacer(1, 4 * mm),
        ])
        caption_text = _page_caption(page, labels)
        if caption_text:
            flowables.extend([Paragraph(escape(caption_text), small), Spacer(1, 3 * mm)])
        scale = min(188 * mm / page.width, 235 * mm / page.height)
        flowables.append(PdfImage(
            io.BytesIO(page.image), width=page.width * scale,
            height=page.height * scale,
        ))
    return flowables
