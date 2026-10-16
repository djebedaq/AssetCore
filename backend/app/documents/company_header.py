"""Approved transfer identity, retained as individual immutable source resources."""

import hashlib
import io

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .common import ASSETS
from .rendering import _rina_image

VERSION = "ODESSOS_TRANSFER_HEADER_V1"


def current_resources():
    certification = _rina_image()
    if certification is None:
        raise ValueError("approved_company_certification_missing")
    return {
        "krz": (ASSETS / "krz_logo.png").read_bytes(),
        "odessos": (ASSETS / "odessos_logo.png").read_bytes(),
        "certification": certification.getvalue(),
    }


def manifest(resources=None):
    resources = resources if resources is not None else current_resources()
    return {
        "version": VERSION,
        "parts": {
            name: {"sha256": hashlib.sha256(content).hexdigest(), "byte_length": len(content)}
            for name, content in resources.items()
        },
    }


def verify(resources, evidence):
    if evidence.get("version") != VERSION or manifest(resources) != evidence:
        raise ValueError("approval_company_header_integrity_failed")


def add_docx(section, resources):
    """Native framed header; no logo redrawing or changes to controlled binaries."""
    from docx.enum.table import WD_TABLE_ALIGNMENT

    section.header.is_linked_to_previous = False
    for child in list(section.header._element):
        section.header._element.remove(child)
    table = section.header.add_table(rows=1, cols=3, width=Mm(198))
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    widths = (27.5, 134, 36.5)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    properties.append(borders)
    table.rows[0].height = Mm(31)
    for cell, width in zip(table.rows[0].cells, widths, strict=True):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = 1
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
    left, center, right = table.rows[0].cells
    left.paragraphs[0].add_run().add_picture(io.BytesIO(resources["krz"]), width=Mm(18))
    est = left.add_paragraph("EST 1955")
    est.alignment = 1
    est.paragraph_format.space_before = est.paragraph_format.space_after = Pt(0)
    est.paragraph_format.line_spacing = 1
    est.runs[0].font.name, est.runs[0].font.size, est.runs[0].bold = "Arial", Pt(9), True
    center.paragraphs[0].add_run().add_picture(io.BytesIO(resources["odessos"]), width=Mm(128))
    right.paragraphs[0].add_run().add_picture(io.BytesIO(resources["certification"]), width=Mm(25))
    tail = section.header.add_paragraph()
    tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    tail.add_run().font.size = Pt(1)
