"""Render part-request visual evidence from the immutable 01A aggregate only."""

from __future__ import annotations

import hashlib
import io
import math
from dataclasses import dataclass
from pathlib import PureWindowsPath

import fitz
from docx import Document
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw
from pydantic import ValidationError
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, Spacer
from sqlalchemy.orm import Session

from ..models import PartRequest, PartVisualArtifact
from ..part_requests.visual_schemas import VisualSnapshotOut
from ..part_requests.visual_snapshots import _integrity_error, snapshot_response
from .common import _language
from .rendering import _pdf_styles, _set_run_font

RENDER_DPI = 180
MAX_PAGE_PIXELS = 8_000_000
MAX_IMAGE_BYTES = 48_000_000
MAX_VISUAL_BLOCKS = 120
MAX_SOURCE_BYTES = 80_000_000
MAX_SOURCE_TOTAL_BYTES = 160_000_000

LABELS = {
    "bg": {
        "title": "Визуално приложение към заявката",
        "line": "Ред",
        "part": "Част №",
        "position": "Позиция",
        "source": "Източник",
        "page": "стр.",
        "revision": "Ревизия",
        "captured": "Заснето",
        "markers": "Отбелязани области",
        "no_visual_reference_at_capture": "При заснемането на този ред не е имало потвърден визуален източник.",
        "legacy_snapshot_unavailable": "За този исторически ред няма неизменима визуална снимка.",
        "no_catalog_binding": "Редът няма потвърдена връзка с каталог и визуален източник.",
    },
    "en": {
        "title": "Visual reference appendix",
        "line": "Line",
        "part": "Part No.",
        "position": "Position",
        "source": "Source",
        "page": "p.",
        "revision": "Revision",
        "captured": "Captured",
        "markers": "Marked regions",
        "no_visual_reference_at_capture": "No verified visual reference was available when this line was captured.",
        "legacy_snapshot_unavailable": "No immutable historical visual snapshot is available for this line.",
        "no_catalog_binding": "This line has no verified catalog binding or visual source.",
    },
    "ru": {
        "title": "Визуальное приложение к заявке",
        "line": "Строка",
        "part": "Деталь №",
        "position": "Позиция",
        "source": "Источник",
        "page": "стр.",
        "revision": "Ревизия",
        "captured": "Сохранено",
        "markers": "Отмеченные области",
        "no_visual_reference_at_capture": "На момент сохранения строки проверенного визуального источника не было.",
        "legacy_snapshot_unavailable": "Для этой исторической строки нет неизменимого визуального снимка.",
        "no_catalog_binding": "У строки нет подтверждённой связи с каталогом и визуальным источником.",
    },
}


@dataclass(frozen=True)
class VisualBlock:
    artifact_sha256: str
    page_number: int
    ordinals: tuple[int, ...]
    image: bytes
    width: int
    height: int
    filename: str | None
    revision: str | None
    positions: tuple[str, ...]


@dataclass(frozen=True)
class VisualLine:
    line_id: int
    sequence: int
    part_number: str
    description: str
    position: str
    state: str
    snapshot_id: int | None
    snapshot_sha256: str | None
    captured_at: str | None
    blocks: tuple[VisualBlock, ...]


def _render_page(
    content: bytes, media_type: str, page_number: int, references: list[dict]
) -> tuple[bytes, int, int]:
    if media_type not in {"application/pdf", "image/png", "image/jpeg", "image/tiff"}:
        raise _integrity_error()
    filetype = {
        "application/pdf": "pdf",
        "image/png": "png",
        "image/jpeg": "jpeg",
        "image/tiff": "tiff",
    }[media_type]
    try:
        with fitz.open(stream=content, filetype=filetype) as source:
            if page_number < 1 or page_number > source.page_count:
                raise _integrity_error()
            page = source.load_page(page_number - 1)
            rect = page.rect  # Display orientation, including the captured PDF rotation.
            if rect.width <= 0 or rect.height <= 0:
                raise _integrity_error()
            scale = min(RENDER_DPI / 72, math.sqrt(MAX_PAGE_PIXELS / (rect.width * rect.height)))
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False
            )
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        draw = ImageDraw.Draw(image)
        stroke = max(4, round(min(image.size) / 300))
        for reference in references:
            x, y, width, height = (reference[key] for key in ("x", "y", "width", "height"))
            if (
                not all(math.isfinite(value) for value in (x, y, width, height))
                or min(x, y) < 0
                or min(width, height) <= 0
                or x + width > 1.000001
                or y + height > 1.000001
            ):
                raise _integrity_error()
            box = (
                round(x * image.width),
                round(y * image.height),
                round((x + width) * image.width),
                round((y + height) * image.height),
            )
            draw.rectangle(box, outline="white", width=stroke + 4)
            draw.rectangle(box, outline="#c81923", width=stroke)
            badge = str(reference["ordinal"])
            badge_width = max(22, 10 + 7 * len(badge))
            bx = min(max(0, box[0]), image.width - badge_width)
            by = max(0, box[1] - 23)
            draw.rectangle((bx, by, bx + badge_width, by + 20), fill="#c81923")
            draw.text((bx + 5, by + 3), badge, fill="white")
        output = io.BytesIO()
        image.save(output, format="PNG", compress_level=6)
        return output.getvalue(), image.width, image.height
    except Exception as exc:
        raise _integrity_error() from exc


def prepare_appendix(db: Session, request: PartRequest) -> tuple[VisualLine, ...]:
    """Validate the whole immutable aggregate before any official row is written."""
    result: list[VisualLine] = []
    total_image_bytes = 0
    visual_blocks = 0
    artifacts: dict[str, bytes] = {}
    total_source_bytes = 0
    for sequence, line in enumerate(sorted(request.lines, key=lambda item: item.id), 1):
        try:
            evidence = snapshot_response(line)
        except (ValueError, TypeError, AttributeError, KeyError) as exc:
            raise _integrity_error() from exc
        state = evidence["state"]
        snapshot = evidence["snapshot"]
        if snapshot is not None:
            try:
                VisualSnapshotOut.model_validate(snapshot)
            except ValidationError as exc:
                raise _integrity_error() from exc
            if snapshot["line_id"] != line.id or snapshot["schema_version"] != 1:
                raise _integrity_error()
            catalog = snapshot["catalog"]
            groups: dict[tuple[str, int], list[dict]] = {}
            for reference in snapshot["visual_references"]:
                groups.setdefault(
                    (reference["artifact_sha256"], reference["page_number"]), []
                ).append(reference)
            if sorted(
                reference["ordinal"] for group in groups.values() for reference in group
            ) != list(range(1, snapshot["occurrence_count"] + 1)):
                raise _integrity_error()
            blocks = []
            for (digest, page), references in groups.items():
                visual_blocks += 1
                if visual_blocks > MAX_VISUAL_BLOCKS:
                    raise _integrity_error()
                if digest not in artifacts:
                    artifact = db.get(PartVisualArtifact, digest)
                    if artifact is None or artifact.byte_length > MAX_SOURCE_BYTES:
                        raise _integrity_error()
                    total_source_bytes += artifact.byte_length
                    if total_source_bytes > MAX_SOURCE_TOTAL_BYTES:
                        raise _integrity_error()
                    if (
                        artifact.byte_length != len(artifact.content)
                        or hashlib.sha256(artifact.content).hexdigest() != digest
                    ):
                        raise _integrity_error()
                    artifacts[digest] = artifact.content
                metadata = references[0]["source_metadata"]
                if any(
                    ref["source_metadata"]["media_type"] != metadata["media_type"]
                    for ref in references
                ):
                    raise _integrity_error()
                filenames = tuple(
                    dict.fromkeys(
                        PureWindowsPath(ref["source_metadata"]["filename"]).name
                        for ref in references
                        if ref["source_metadata"].get("filename")
                    )
                )
                revisions = tuple(
                    dict.fromkeys(
                        value
                        for ref in references
                        if (
                            value := ref["source_metadata"].get("revision_label")
                            or ref["source_metadata"].get("document_revision")
                        )
                    )
                )
                image, width, height = _render_page(
                    artifacts[digest], metadata["media_type"], page, references
                )
                total_image_bytes += len(image)
                if total_image_bytes > MAX_IMAGE_BYTES:
                    raise _integrity_error()
                blocks.append(
                    VisualBlock(
                        digest,
                        page,
                        tuple(ref["ordinal"] for ref in references),
                        image,
                        width,
                        height,
                        " / ".join(filenames) or None,
                        " / ".join(revisions) or None,
                        tuple(
                            str(
                                ref["source_metadata"].get("label") or catalog.get("position") or ""
                            )
                            for ref in references
                        ),
                    )
                )
            part_number = str(
                catalog.get("requested_part_number") or catalog.get("part_number") or ""
            )
            description = str(catalog.get("description") or "")
            position = str(catalog.get("position") or "")
            captured_at = snapshot["captured_at"]
            snapshot_id, snapshot_sha256 = snapshot["id"], snapshot["sha256"]
        else:
            blocks = []
            part_number = line.part_number or ""
            description = line.description
            position = line.position or ""
            captured_at = None
            snapshot_id = snapshot_sha256 = None
        result.append(
            VisualLine(
                line.id,
                sequence,
                part_number,
                description,
                position,
                state,
                snapshot_id,
                snapshot_sha256,
                captured_at,
                tuple(blocks),
            )
        )
    return tuple(result)


def appendix_manifest(lines: tuple[VisualLine, ...]) -> dict:
    return {
        "renderer_version": 1,
        "lines": [
            {
                "line_id": line.line_id,
                "visual_snapshot_id": line.snapshot_id,
                "visual_snapshot_sha256": line.snapshot_sha256,
                "blocks": [
                    {
                        "artifact_sha256": block.artifact_sha256,
                        "page_number": block.page_number,
                        "occurrence_ordinals": list(block.ordinals),
                    }
                    for block in line.blocks
                ],
            }
            for line in lines
        ],
    }


def _line_label(line: VisualLine, labels: dict[str, str]) -> str:
    return f"{labels['line']} {line.sequence} · {labels['part']} {line.part_number or '—'} · {line.description}"


def _block_caption(block: VisualBlock, line: VisualLine, labels: dict[str, str]) -> str:
    parts = [
        f"{labels['position']}: {line.position or '—'}",
        f"{labels['source']}: {block.filename or '—'}",
        f"{labels['page']} {block.page_number}",
    ]
    if block.revision:
        parts.append(f"{labels['revision']}: {block.revision}")
    if line.captured_at:
        parts.append(f"{labels['captured']}: {line.captured_at[:19]}")
    markers = ", ".join(
        f"{ordinal}: {position or '—'}"
        for ordinal, position in zip(block.ordinals, block.positions, strict=True)
    )
    parts.append(f"{labels['markers']}: {markers}")
    return " · ".join(parts)


def append_docx(docx: bytes, lines: tuple[VisualLine, ...], language: str) -> bytes:
    if not lines:
        return docx
    labels = LABELS[_language(language)]
    document = Document(io.BytesIO(docx))
    section = document.sections[-1]
    max_width = min(Mm(175), section.page_width - section.left_margin - section.right_margin)
    max_height = min(
        Mm(205), section.page_height - section.top_margin - section.bottom_margin - Mm(50)
    )
    for line in lines:
        blocks = line.blocks or (None,)
        for block in blocks:
            document.add_page_break()
            heading = document.add_paragraph()
            heading.paragraph_format.space_after = Pt(8)
            heading.paragraph_format.keep_with_next = True
            _set_run_font(heading.add_run(labels["title"]), 13, True)
            identity = document.add_paragraph()
            identity.paragraph_format.space_after = Pt(5)
            identity.paragraph_format.keep_with_next = True
            _set_run_font(identity.add_run(_line_label(line, labels)), 9, True)
            caption = document.add_paragraph()
            caption.paragraph_format.space_after = Pt(6)
            caption.paragraph_format.keep_with_next = True
            _set_run_font(
                caption.add_run(
                    _block_caption(block, line, labels) if block else labels[line.state]
                ),
                8,
            )
            if block:
                ratio = block.width / block.height
                width = min(max_width, int(max_height * ratio))
                height = min(max_height, int(max_width / ratio))
                image_paragraph = document.add_paragraph()
                image_paragraph.paragraph_format.space_after = Pt(0)
                image_paragraph.add_run().add_picture(
                    io.BytesIO(block.image), width=width, height=height
                )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_flowables(lines: tuple[VisualLine, ...], language: str) -> list:
    if not lines:
        return []
    labels = LABELS[_language(language)]
    _, _, _, label, title, small = _pdf_styles()
    flowables = []
    from xml.sax.saxutils import escape

    for line in lines:
        for block in line.blocks or (None,):
            flowables.extend(
                [
                    PageBreak(),
                    Paragraph(escape(labels["title"]), title),
                    Spacer(1, 5 * mm),
                    Paragraph(escape(_line_label(line, labels)), label),
                    Spacer(1, 2 * mm),
                    Paragraph(
                        escape(
                            _block_caption(block, line, labels) if block else labels[line.state]
                        ),
                        small,
                    ),
                    Spacer(1, 5 * mm),
                ]
            )
            if block:
                scale = min(175 * mm / block.width, 205 * mm / block.height)
                flowables.append(
                    PdfImage(
                        io.BytesIO(block.image),
                        width=block.width * scale,
                        height=block.height * scale,
                    )
                )
    return flowables
