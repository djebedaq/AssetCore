"""KK reference typography with one DOCX layout for official and preview PDF."""

import io

import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.pdfgen import canvas as pdf_canvas

from ..template_engine import convert_docx_to_pdf
from . import company_header
from .part_request_grouped_visuals import append_docx
from .rendering import _clear_body, _pdf_styles, _set_repeat_table_header
from .request_machine_section import add_docx as add_machine_table

APPROVER = {"bg": "Одобрил", "en": "Approved by", "ru": "Утвердил"}


def paragraph_style(paragraph, *, size=12, center=False):
    paragraph.alignment = 1 if center else 0
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for run in paragraph.runs:
        run.font.name, run.font.size = "DejaVu Serif", Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)


def borders(table, visible=True):
    existing = table._tbl.tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        table._tbl.tblPr.remove(existing)
    edges = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        for key, value in {
            "val": "single" if visible else "nil",
            "sz": "4",
            "color": "000000",
        }.items():
            edge.set(qn(f"w:{key}"), value)
        edges.append(edge)
    table._tbl.tblPr.append(edges)


def build_docx(
    data, template_content, *, rows, meta, date, resources, appendix=None, machine_section=None
):
    document = Document(io.BytesIO(template_content)) if template_content else Document()
    _clear_body(document)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = section.bottom_margin = Mm(12.7)
    section.header_distance, section.top_margin = Mm(12.7), Mm(48)
    section.different_first_page_header_footer = False
    document.settings.odd_and_even_pages_header_footer = False
    company_header.add_docx(section, resources)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Serif", Pt(12)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    if title_style.element.pPr is not None:
        for border in list(title_style.element.pPr.findall(qn("w:pBdr"))):
            title_style.element.pPr.remove(border)
    title = document.add_paragraph("TECHNICAL SPECIFICATION", "Title")
    paragraph_style(title, center=True)
    title.runs[0].bold = True
    title.paragraph_format.keep_with_next = True
    subtitle = document.add_paragraph("FOR SPARE PARTS SUPPLY")
    paragraph_style(subtitle, center=True)
    subtitle.runs[0].italic = True
    subtitle.paragraph_format.space_after = Pt(5)
    subtitle.paragraph_format.keep_with_next = True
    for text in meta:
        p = document.add_paragraph(text)
        paragraph_style(p, size=10)
        p.paragraph_format.keep_with_next = True
    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Mm(4)
    if machine_section is not None:
        add_machine_table(
            document,
            machine_section,
            data["language"],
            paragraph_style=paragraph_style,
            borders=borders,
        )
        gap = document.add_paragraph()
        gap.paragraph_format.space_after = Pt(0)
        gap.paragraph_format.line_spacing = Mm(4)
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    widths = [14.2, 140.4, 30] if len(rows[0]) == 3 else [14.2, 38.2, 102.2, 30]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    borders(table)
    for index, values in enumerate(rows):
        row = table.add_row()
        row.height, row.height_rule = Mm(12 if index == 0 else 10), WD_ROW_HEIGHT_RULE.AT_LEAST
        split = OxmlElement("w:cantSplit")
        split.set(
            qn("w:val"), "1" if all(len(v) < 1500 and v.count("\n") < 30 for v in values) else "0"
        )
        row._tr.get_or_add_trPr().append(split)
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            paragraph_style(cell.paragraphs[0], center=True)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
    _set_repeat_table_header(table.rows[0])
    gap = document.add_paragraph()
    gap.paragraph_format.line_spacing = Mm(9)
    gap.paragraph_format.space_after = Pt(0)
    footer = document.add_table(rows=1, cols=3)
    footer.alignment, footer.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    borders(footer, visible=False)
    values = (
        f"ENQ №: {data.get('official_number') or ''}",
        f"DATE: {date}",
        f"{APPROVER[data['language']]}: {data.get('approved_by_name') or ''}",
    )
    for column, cell, width, value in zip(
        footer.columns, footer.rows[0].cells, (60, 48, 76.6), values, strict=True
    ):
        column.width, cell.width, cell.text = Mm(width), Mm(width), value
        paragraph_style(cell.paragraphs[0], size=11)
    output = io.BytesIO()
    document.save(output)
    return (
        append_docx(output.getvalue(), appendix, data["language"])
        if appendix
        else output.getvalue()
    )


def build_pdf(data, docx, *, preview=False, watermark=None):
    content = convert_docx_to_pdf(docx)
    if content is None:
        raise ValueError("request_pdf_converter_unavailable")
    if not preview:
        return content
    # Only the preview overlay differs from the same DOCX print layout.
    body = _pdf_styles()[2]
    with fitz.open(stream=content, filetype="pdf") as document:
        for page in document:
            output = io.BytesIO()
            overlay = pdf_canvas.Canvas(output, pagesize=(page.rect.width, page.rect.height))
            overlay.setFillColorRGB(0.65, 0.15, 0.15)
            overlay.setFillAlpha(0.25)
            overlay.setFont(body.fontName, 22)
            overlay.translate(page.rect.width / 2, page.rect.height / 2)
            overlay.rotate(38)
            overlay.drawCentredString(0, 0, watermark)
            overlay.save()
            with fitz.open(stream=output.getvalue(), filetype="pdf") as stamp:
                page.show_pdf_page(page.rect, stamp, 0)
        return document.tobytes(garbage=4, deflate=True)
