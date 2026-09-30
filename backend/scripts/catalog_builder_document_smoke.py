"""Actual Builder-to-LibreOffice QA in memory; no configured database writes."""

from __future__ import annotations

import hashlib
import io
import secrets

import fitz
from app import models as m
from app.catalog import service as runtime
from app.catalog_admin import publication
from app.catalog_admin.service import bind_asset
from app.documents import part_request_documents
from app.documents.part_request_grouped_visuals import prepare_appendix
from app.industrial_api import create_multi_part_request
from app.industrial_schemas import MultiPartRequestCreate
from app.part_requests.service import load_request
from app.seed import seed_database
from app.settings import settings
from docx import Document
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session


def main() -> None:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    m.Base.metadata.create_all(engine)
    converted = []
    original = part_request_documents.convert_docx_to_pdf

    def actual_conversion(content):
        pdf = original(content)
        if pdf is None:
            raise RuntimeError("Builder final DOCX did not convert through LibreOffice")
        converted.append((content, pdf))
        return pdf

    part_request_documents.convert_docx_to_pdf = actual_conversion
    try:
        with Session(engine) as db:
            settings.owner_email = settings.owner_email or "qa-only@assetcore.invalid"
            settings.owner_initial_password = settings.owner_initial_password or secrets.token_urlsafe(32)
            settings.owner_job_title = settings.owner_job_title or "QA оператор"
            seed_database(db)
            actor = db.scalar(select(m.User).where(m.User.is_system_owner.is_(True)))
            actor.first_name, actor.middle_name, actor.last_name = "Тест", "Само", "Проверка"
            actor.full_name, actor.job_title, actor.profile_status = "Тест Само Проверка", "QA оператор", "PROFILE_COMPLETE"
            category = m.AssetCategory(code="QA_BUILDER_DOC", name_bg="QA", name_en="QA", name_ru="QA",
                                       capabilities=["HAS_PARTS_CATALOG"])
            db.add(category)
            db.flush()
            catalog = m.CatalogDefinition(code="QA_BUILDER_DOC", asset_category_id=category.id,
                                           name_bg="QA", name_en="QA", name_ru="QA", created_by_id=actor.id)
            db.add(catalog)
            db.flush()
            revision = m.CatalogRevision(catalog_id=catalog.id, revision_code="A", status="DRAFT", created_by_id=actor.id)
            db.add(revision)
            db.flush()
            assembly = m.CatalogRevisionAssembly(revision_id=revision.id, code="QA",
                                                   name_bg="QA", name_en="QA", name_ru="QA", created_by_id=actor.id)
            db.add(assembly)
            db.flush()
            with fitz.open() as pdf:
                pdf.new_page().insert_text((60, 80), "QA-ONLY Builder scheme")
                pdf.new_page().insert_text((60, 80), "QA-ONLY Builder exact spare list")
                raw = pdf.tobytes()
            artifact = m.CatalogRevisionArtifact(assembly_id=assembly.id, title="QA-ONLY",
                filename="qa-only.pdf", media_type="application/pdf", content=raw,
                sha256=hashlib.sha256(raw).hexdigest(), page_count=2, created_by_id=actor.id)
            db.add(artifact)
            db.flush()
            scheme = m.CatalogRevisionVisualPage(artifact_id=artifact.id, page_number=1,
                role="EXPLODED_SCHEME", created_by_id=actor.id)
            spare = m.CatalogRevisionVisualPage(artifact_id=artifact.id, page_number=2,
                role="SPARE_PARTS_LIST", created_by_id=actor.id)
            part = m.CatalogRevisionPart(assembly_id=assembly.id, position="P" * 80,
                part_number="QA-ONLY-PART", name_bg="QA тестова част", name_en="QA test part",
                name_ru="QA тестовая деталь", created_by_id=actor.id)
            db.add_all([scheme, spare, part])
            db.flush()
            db.add(m.CatalogRevisionPartPageMap(part_id=part.id, visual_page_id=spare.id, created_by_id=actor.id))
            db.add(m.CatalogRevisionPositionHotspot(visual_page_id=scheme.id, position=part.position,
                x=.1, y=.1, width=.1, height=.1, provenance="MANUAL_VERIFIED", is_verified=True,
                verified_by_id=actor.id, verified_at=m.utcnow(), version=1, created_by_id=actor.id))
            db.commit()
            preview = publication.readiness(db, revision.id)
            assert preview["ready"], preview["errors"]
            publication.publish(db, actor, revision.id, preview["publication_digest"], None, True)
            live = db.scalar(select(m.PartCatalog).where(m.PartCatalog.builder_part_id == part.id))
            assert not live.compatible_machine_numbers
            for number in ("QA-ONLY-LATE-A", "QA-ONLY-LATE-B"):
                machine = m.Machine(inventory_number=number, name="QA-ONLY Builder document asset",
                    category=category.code, category_id=category.id, brand="QA", model="QA")
                db.add(machine)
                db.commit()
                bind_asset(db, actor, catalog.id, machine.id)
                assert runtime.machine_catalog(db, machine.id)["dataset_version"] == f"CATALOG_BUILDER_R{revision.id}"
                created = create_multi_part_request(MultiPartRequestCreate(machine_id=machine.id,
                    lines=[{"catalog_part_id": live.id, "description": "QA-ONLY", "quantity": 1}]), user=actor, db=db)
                request = load_request(db, created["id"])
                request.status, request.decided_by_id = m.PartRequestStatus.APPROVED.value, actor.id
                db.commit()
                appendix = prepare_appendix(db, request)
                assert [(page.role, page.page_number) for page in appendix.pages] == [
                    ("EXPLODED_SCHEME", 1), ("SPARE_PARTS_LIST", 2)]
                records = part_request_documents.make_part_request_documents(db, request, actor.id)
                for record in records:
                    assert hashlib.sha256(record.content).hexdigest() == record.sha256
                docx, output_pdf = converted[-1]
                doc = Document(io.BytesIO(docx))
                assert any([cell.text for cell in table.rows[0].cells] ==
                    ["Поз.", "PART №", "Описание", "Количество"] for table in doc.tables)
                assert any(cell.text == "P" * 80 for table in doc.tables for row in table.rows for cell in row.cells)
                with fitz.open(stream=output_pdf, filetype="pdf") as rendered:
                    assert rendered.page_count >= 3
                    assert "P" * 80 in "".join("".join(page.get_text().split()) for page in rendered)
                    assert "Разглобена схема" in rendered[-2].get_text()
                    assert "Списък резервни части" in rendered[-1].get_text()
                    assert all(page.get_images() for page in (rendered[-2], rendered[-1]))
            assert len(converted) == 2
    finally:
        part_request_documents.convert_docx_to_pdf = original
        engine.dispose()
    print("builder_late_binding_80_char_official_docx_pdf_libreoffice=passed")


if __name__ == "__main__":
    main()
