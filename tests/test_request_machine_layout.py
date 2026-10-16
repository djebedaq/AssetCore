"""Future machine sections and historical compatibility; disposable QA data only."""

import copy
import hashlib
import io
import json
import os
import zipfile
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from app.documents import company_header
from app.documents import request_issue_documents as renderer
from app.documents.request_machine_section import rows, snapshot
from app.models import (
    CategoryFieldDefinition,
    GeneratedDocument,
    Machine,
    MachineFieldValue,
    PartCatalog,
    PartRequest,
    PartRequestLine,
    RepairKit,
    User,
)
from app.part_requests.official import capture_data
from app.part_requests.service import load_request
from app.template_engine import convert_docx_to_pdf
from docx import Document
from sqlalchemy import select
from test_requests01 import approve, create
from visual_snapshot_cases import future_catalog, request_payload


def document_xml(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return archive.read("word/document.xml")


def records(factory, request_id):
    with factory() as db:
        return {
            row.format: (row.content, row.sha256)
            for row in db.scalars(
                select(GeneratedDocument).where(GeneratedDocument.part_request_id == request_id)
            )
        }


@pytest.mark.parametrize("language", ["bg", "en", "ru"])
def test_machine_rows_keep_four_cells_missing_values_and_zero(language):
    values = rows({"name": "QA", "pressure_bar": 0}, language)
    assert [len(row) for row in values] == [4] * 4
    assert values[0][1:] == ["QA", values[0][2], ""]
    assert values[2][1] == "" and values[2][3] == "0 bar"
    assert values[3][1:] == ["", "", ""]
    assert rows({}, language)[2][3] == ""
    assert len({row[0] for row in values}) == 4


@pytest.mark.parametrize(
    "active,same_category,code,expected",
    [
        (True, True, "batch_reference", "QA-BATCH"),
        (True, True, "BATCH_REFERENCE", "QA-BATCH"),
        (False, True, "batch_reference", None),
        (True, False, "batch_reference", None),
        (True, True, "transfer_batch", None),
    ],
)
def test_only_explicit_machine_batch_is_captured(active, same_category, code, expected):
    field = SimpleNamespace(
        id=7, category_id=1 if same_category else 2, code=code, is_active=active
    )
    machine = SimpleNamespace(
        name="QA",
        inventory_number="QA",
        brand=None,
        model=None,
        serial_number=None,
        pressure_bar=None,
        category_id=1,
        custom_values=[SimpleNamespace(field=field, value="QA-BATCH")],
    )
    captured = snapshot(machine)
    assert captured["batch_reference"] == expected
    assert ("batch_reference_source" in captured) == (expected is not None)


@pytest.mark.parametrize("kind", ["CATALOG", "UNKNOWN_PART", "REPAIR_KIT"])
@pytest.mark.parametrize("linked", [False, True])
def test_section_depends_on_machine_relation_not_request_kind(session_factory, kind, linked):
    with session_factory() as db:
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        machine = db.scalar(select(Machine).where(Machine.inventory_number == "10"))
        kit = RepairKit(code="QA-LAYOUT", name="QA layout test", created_by_id=actor.id)
        db.add(kit)
        db.flush()
        request = PartRequest(
            part_name="QA layout test",
            requested_by_id=actor.id,
            machine_id=machine.id if linked else None,
            repair_kit_id=kit.id if kind == "REPAIR_KIT" else None,
        )
        request.lines.append(
            PartRequestLine(
                description="QA row",
                quantity=1,
                is_unknown_part=kind == "UNKNOWN_PART",
            )
        )
        db.add(request)
        db.flush()
        data = capture_data(load_request(db, request.id))
        assert data["request_kind"] == kind
        assert data["document_layout_version"] == (3 if linked else 2)
        doc = Document(io.BytesIO(renderer.build_docx(data, None)))
        assert len(doc.tables) == (3 if linked else 2)
        if linked:
            assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == rows(
                data["machine"], "bg"
            )
            assert all(machine.serial_number not in p.text for p in doc.paragraphs)
        assert [cell.text for cell in doc.tables[-2].rows[0].cells] == [
            "Pos",
            "PART No",
            "DESCRIPTION",
            "Q'TY ordered",
        ]
        text = document_xml(renderer.build_docx(data, None)).decode()
        for forbidden in ("Решение", "Съставил", "Приел /", "Статус на подписите"):
            assert forbidden not in text


def test_free_form_is_identical_and_historical_versions_keep_flattened_line(
    client, auth_headers, session_factory
):
    item = create(client, auth_headers)
    with session_factory() as db:
        data = capture_data(load_request(db, item["id"]))
    assert data["document_layout_version"] == 2 and data["machine"] is None
    baseline = renderer.build_docx(data, None)
    future = renderer.build_docx({**data, "document_layout_version": 3}, None)
    assert document_xml(baseline) == document_xml(future)
    with (
        fitz.open(stream=convert_docx_to_pdf(baseline), filetype="pdf") as a,
        fitz.open(stream=convert_docx_to_pdf(future), filetype="pdf") as b,
    ):
        assert len(a) == len(b)
        assert all(
            x.get_pixmap().samples == y.get_pixmap().samples for x, y in zip(a, b, strict=True)
        )
    data["request_kind"] = "CATALOG"
    data["machine"] = {"name": "QA old machine", "inventory_number": "QA-OLD", "pressure_bar": 0}
    data["lines"][0].update(position="1", part_number="QA", linked_part_number=None)
    for version in (None, 1, 2):
        old = renderer.build_docx(
            {**data, "document_layout_version": version}, None if version == 2 else baseline
        )
        doc = Document(io.BytesIO(old))
        assert len(doc.tables) == (2 if version == 2 else 1)
        assert "QA old machine · QA-OLD" in [p.text for p in doc.paragraphs]
    with pytest.raises(ValueError, match="unsupported_request_document_layout_version"):
        renderer.build_docx({**data, "document_layout_version": 99}, None)


def test_machine_retry_uses_approval_values_and_preserves_ready_bytes(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory, visuals=False)
    with session_factory() as db:
        machine = db.get(Machine, data["machine_id"])
        field = CategoryFieldDefinition(
            category_id=machine.category_id, code="batch_reference", label_bg="QA партида"
        )
        db.add(field)
        db.flush()
        db.add(MachineFieldValue(machine_id=machine.id, field_id=field.id, value="QA-BATCH-FROZEN"))
        db.commit()
    response = client.post(
        "/api/part-requests/multi", headers=auth_headers, json=request_payload(data)
    )
    assert response.status_code == 201
    item = response.json()
    original_pdf = renderer.build_pdf
    monkeypatch.setattr(
        renderer, "build_pdf", lambda *_a, **_kw: (_ for _ in ()).throw(ValueError("QA retry"))
    )
    item = approve(client, auth_headers, item)
    assert item["generation_status"] == "FAILED"
    with session_factory() as db:
        request = load_request(db, item["id"])
        frozen = copy.deepcopy(request.approval_snapshot.data)
        digest = request.approval_snapshot.sha256
        assert frozen["document_layout_version"] == 3
        assert frozen["machine"]["batch_reference"] == "QA-BATCH-FROZEN"
        request.machine.name, request.machine.pressure_bar = "QA changed later", 999
        request.machine.custom_values[0].value = "QA changed batch"
        db.commit()

    def unavailable(*_a, **_kw):
        raise AssertionError("Active resources must not be read during retry")

    monkeypatch.setattr(company_header, "current_resources", unavailable)
    monkeypatch.setattr(renderer, "_template_version", unavailable)
    monkeypatch.setattr(renderer, "build_pdf", original_pdf)
    retried = client.post(f"/api/part-requests/{item['id']}/documents", headers=auth_headers)
    assert retried.status_code == 201 and retried.json()["generation_status"] == "READY"
    before = records(session_factory, item["id"])
    doc = Document(io.BytesIO(before["docx"][0]))
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == rows(frozen["machine"], "bg")
    with session_factory() as db:
        request = load_request(db, item["id"])
        assert (
            request.approval_snapshot.data == frozen and request.approval_snapshot.sha256 == digest
        )
        actor = db.get(User, data["actor_id"])
        renderer.archive(db, request, actor)
        db.commit()
    assert before == records(session_factory, item["id"])
    for content, sha in before.values():
        assert hashlib.sha256(content).hexdigest() == sha


def test_catalog_preview_approval_and_review_export(client, auth_headers, session_factory):
    """Use verified seed machine/part/evidence; export only to an explicitly requested .tmp folder."""
    with session_factory() as db:
        machine = db.scalar(select(Machine).where(Machine.inventory_number == "10"))
        part = db.scalar(select(PartCatalog).where(PartCatalog.part_number == "E0113540"))
        body = {
            "machine_id": machine.id,
            "department": "QA — ПРИМЕР ЗА ПРЕГЛЕД",
            "reason": "QA пример; не е производствена заявка.",
            "lines": [{"catalog_part_id": part.id, "description": part.description, "quantity": 1}],
        }
    response = client.post("/api/part-requests/multi", headers=auth_headers, json=body)
    assert response.status_code == 201, response.status_code
    item = response.json()
    preview = client.get(f"/api/part-requests/{item['id']}/preview", headers=auth_headers)
    assert preview.status_code == 200
    with fitz.open(stream=preview.content, filetype="pdf") as pdf:
        assert len(pdf) >= 3
        assert all("ПРОЕКТ" in page.get_text() for page in pdf)
        assert "G39300298" in pdf[0].get_text() and "500 bar" in pdf[0].get_text()
        assert "KK1000" not in pdf[0].get_text()
        assert not renderer._date({})
    submitted = client.post(f"/api/part-requests/{item['id']}/submit", headers=auth_headers)
    assert submitted.status_code == 200
    item = approve(client, auth_headers, item)
    assert item["generation_status"] == "READY"
    output = records(session_factory, item["id"])
    with session_factory() as db:
        frozen = load_request(db, item["id"]).approval_snapshot.data
    doc = Document(io.BytesIO(output["docx"][0]))
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == rows(frozen["machine"], "bg")
    footer = " ".join(c.text for c in doc.tables[2].rows[0].cells)
    assert (
        "KK1000" in footer
        and renderer._date(frozen) in footer
        and frozen["approved_by_name"] in footer
    )
    independent = convert_docx_to_pdf(output["docx"][0])
    with (
        fitz.open(stream=output["pdf"][0], filetype="pdf") as pdf,
        fitz.open(stream=independent, filetype="pdf") as check,
    ):
        assert len(pdf) == len(check) >= 3
        assert all(
            a.get_pixmap().samples == b.get_pixmap().samples
            for a, b in zip(pdf, check, strict=True)
        )
    target = os.environ.get("ASSETCORE_MACHINE_LAYOUT_QA_OUTPUT")
    if target:
        target = Path(target).resolve()
        assert target.is_relative_to((Path(__file__).resolve().parents[1] / ".tmp").resolve())
        target.mkdir(parents=True, exist_ok=True)
        (target / "catalog-preview.pdf").write_bytes(preview.content)
        for format_, (content, _) in output.items():
            (target / f"catalog-KK1000.{format_}").write_bytes(content)
        (target / "catalog-snapshot.json").write_text(
            json.dumps(frozen, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        for name, content in (
            ("catalog-preview", preview.content),
            ("catalog-KK1000", output["pdf"][0]),
        ):
            with fitz.open(stream=content, filetype="pdf") as pdf:
                for index, page in enumerate(pdf, 1):
                    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(
                        target / f"{name}-{index}.png"
                    )
