"""PARTS-DOC-01B: official output uses retained visual evidence only."""

from __future__ import annotations

import base64
import hashlib
import io
import zipfile

import fitz
import pytest
from app.documents import part_request_documents
from app.documents.part_request_visual_appendix import LABELS, _render_page, prepare_appendix
from app.models import (
    DocumentTemplateVersion,
    GeneratedDocument,
    OfficialDocument,
    OfficialDocumentVersion,
    PartCatalog,
    PartHotspot,
    PartRequestLine,
    PartVisualArtifact,
    TechnicalDocument,
    User,
)
from app.part_requests.service import load_request
from app.part_requests.visual_snapshots import create_request_line
from app.template_engine import source_bytes
from docx import Document
from fastapi import HTTPException
from PIL import Image
from sqlalchemy import func, select, text
from visual_snapshot_cases import future_catalog, request_payload


def _request(client, headers, data, *, lines=None):
    payload = request_payload(data)
    if lines is not None:
        payload["lines"] = lines
    created = client.post("/api/part-requests/multi", headers=headers, json=payload)
    assert created.status_code == 201, created.text
    request = created.json()
    approved = client.post(
        f"/api/part-requests/{request['id']}/decision",
        headers=headers,
        json={"decision": "APPROVED", "note": "QA approval"},
    )
    assert approved.status_code == 200, approved.text
    return request


def _generate(client, headers, request, language="bg"):
    return client.post(
        f"/api/part-requests/{request['id']}/documents?language={language}",
        headers=headers,
    )


def _records(factory, request_id):
    with factory() as db:
        records = db.scalars(
            select(GeneratedDocument).where(GeneratedDocument.part_request_id == request_id)
        ).all()
        return {
            record.format: (record.content, record.sha256, record.snapshot) for record in records
        }


def test_appendix_labels_have_bg_en_ru_key_parity():
    assert set(LABELS) == {"bg", "en", "ru"}
    assert set(LABELS["bg"]) == set(LABELS["en"]) == set(LABELS["ru"])


def test_future_catalog_official_docx_and_fallback_pdf(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory)
    request = _request(client, auth_headers, data)
    with session_factory() as db:
        template_hashes = {
            version.id: version.source_sha256
            for version in db.scalars(select(DocumentTemplateVersion)).all()
        }
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    generated = _generate(client, auth_headers, request)
    assert generated.status_code == 201, generated.text
    records = _records(session_factory, request["id"])
    assert set(records) == {"docx", "pdf"}
    for content, digest, _ in records.values():
        assert digest == hashlib.sha256(content).hexdigest()
    with zipfile.ZipFile(io.BytesIO(records["docx"][0])) as archive:
        document_xml = archive.read("word/document.xml").decode()
        relations = archive.read("word/_rels/document.xml.rels").decode()
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert "Визуално приложение към заявката" in document_xml
    assert "QA-991122" in document_xml
    assert "qa-source.pdf" in document_xml
    assert 'w:type="page"' in document_xml
    assert 'TargetMode="External"' not in relations
    assert "{{" not in document_xml and "data:image/" not in document_xml
    assert "D:\\" not in document_xml and "C:\\" not in document_xml
    with session_factory() as db:
        official = db.scalar(
            select(OfficialDocument).where(
                OfficialDocument.document_number == request["request_reference"]
            )
        )
        version = db.get(OfficialDocumentVersion, official.current_version_id)
        template = db.get(DocumentTemplateVersion, version.template_version_id)
        original_image_count = len(Document(io.BytesIO(source_bytes(template))).inline_shapes)
    assert len(Document(io.BytesIO(records["docx"][0])).inline_shapes) == original_image_count + 2
    assert len(media) >= original_image_count + 2
    manifest = records["docx"][2]["visual_appendix"]
    assert manifest["renderer_version"] == 1
    assert {
        (block["artifact_sha256"], block["page_number"]) for block in manifest["lines"][0]["blocks"]
    } == {(data["sha256"], 1), (data["sha256"], 2)}
    assert sorted(
        len(block["occurrence_ordinals"]) for block in manifest["lines"][0]["blocks"]
    ) == [2, 2]
    with fitz.open(stream=records["pdf"][0], filetype="pdf") as pdf:
        assert pdf.page_count >= 3
        pdf_text = "\n".join(page.get_text() for page in pdf)
        assert "Визуално приложение към заявката" in pdf_text
        assert "QA-991122" in pdf_text
        assert "qa-source.pdf" in pdf_text
        for page in list(pdf)[-2:]:
            images = page.get_images(full=True)
            assert len(images) == 1
            for image in images:
                for rect in page.get_image_rects(image[0]):
                    assert page.rect.contains(rect)
    with session_factory() as db:
        official = db.scalar(
            select(OfficialDocument).where(
                OfficialDocument.document_number == request["request_reference"]
            )
        )
        version = db.get(OfficialDocumentVersion, official.current_version_id)
        assert version.docx_sha256 == records["docx"][1]
        assert version.pdf_sha256 == records["pdf"][1]
        assert {
            v.id: v.source_sha256 for v in db.scalars(select(DocumentTemplateVersion)).all()
        } == template_hashes
    duplicate = _generate(client, auth_headers, request)
    assert duplicate.status_code == 409
    assert duplicate.json()["detail"]["code"] == "part_request_protocol_already_generated"


def test_captured_page_grouping_and_marker_pixels(client, auth_headers, session_factory):
    data = future_catalog(session_factory)
    request = _request(client, auth_headers, data)
    with session_factory() as db:
        lines = prepare_appendix(db, load_request(db, request["id"]))
    assert len(lines) == 1
    assert [(block.page_number, block.ordinals) for block in lines[0].blocks] == [
        (2, (1, 4)),
        (1, (2, 3)),
    ]
    for block in lines[0].blocks:
        with Image.open(io.BytesIO(block.image)) as image:
            assert image.width == block.width and image.height == block.height
            for x, y in (
                [(0.6, 0.6), (0.1, 0.2)] if block.page_number == 2 else [(0.2, 0.2), (0.3, 0.2)]
            ):
                red = image.getpixel((round(x * image.width), round(y * image.height)))
                assert red[0] > red[1] * 1.5 and red[0] > red[2] * 1.5


def test_live_catalog_changes_do_not_change_appendix(client, auth_headers, session_factory):
    data = future_catalog(session_factory)
    request = _request(client, auth_headers, data)
    with session_factory() as db:
        baseline = prepare_appendix(db, load_request(db, request["id"]))
        before = [hashlib.sha256(block.image).hexdigest() for block in baseline[0].blocks]
    with session_factory() as db:
        part = db.get(PartCatalog, data["part_id"])
        part.description = "QA changed after capture"
        part.source_document = "later.pdf"
        db.get(PartRequestLine, request["lines"][0]["id"]).description = "QA changed request line"
        source = db.get(TechnicalDocument, data["document_id"])
        source.uploaded_content = b"later source"
        source.sha256 = hashlib.sha256(source.uploaded_content).hexdigest()
        db.commit()
    with session_factory() as db:
        current = prepare_appendix(db, load_request(db, request["id"]))
        after = [hashlib.sha256(block.image).hexdigest() for block in current[0].blocks]
    assert before == after
    assert current[0].part_number == "QA-991122"
    assert current[0].description == "Synthetic QA rotor"
    assert all(block.filename == "qa-source.pdf" for block in current[0].blocks)


@pytest.mark.parametrize("corrupt", ["snapshot", "artifact"])
def test_corrupt_immutable_source_aborts_official_generation(
    client, auth_headers, session_factory, corrupt
):
    data = future_catalog(session_factory)
    request = _request(client, auth_headers, data)
    with session_factory() as db:
        table = "part_visual_snapshots" if corrupt == "snapshot" else "part_visual_artifacts"
        triggers = (
            db.execute(
                text("SELECT name FROM sqlite_master WHERE type='trigger' AND tbl_name=:table"),
                {"table": table},
            )
            .scalars()
            .all()
        )
        for trigger in triggers:
            db.execute(text(f'DROP TRIGGER "{trigger}"'))
        if corrupt == "snapshot":
            db.execute(text("UPDATE part_visual_snapshots SET occurrence_count=occurrence_count+1"))
        else:
            artifact = db.get(PartVisualArtifact, data["sha256"])
            changed = b"X" + artifact.content[1:]
            db.execute(
                text("UPDATE part_visual_artifacts SET content=:content WHERE sha256=:sha"),
                {"content": changed, "sha": data["sha256"]},
            )
        db.commit()
    response = _generate(client, auth_headers, request)
    assert response.status_code == 409, response.text
    with session_factory() as db:
        assert db.scalar(select(func.count()).select_from(GeneratedDocument)) == 0
        assert db.scalar(select(func.count()).select_from(OfficialDocument)) == 0
        assert db.scalar(select(func.count()).select_from(OfficialDocumentVersion)) == 0


def test_rotated_landscape_page_marker_geometry():
    with fitz.open() as document:
        page = document.new_page(width=400, height=200)
        page.insert_text((30, 30), "QA rotated drawing")
        page.set_rotation(90)
        content = document.tobytes()
    png, width, height = _render_page(
        content,
        "application/pdf",
        1,
        [{"ordinal": 1, "x": 0.25, "y": 0.4, "width": 0.1, "height": 0.1}],
    )
    assert height > width
    with Image.open(io.BytesIO(png)) as image:
        red = image.getpixel((round(width * 0.25), round(height * 0.4)))
        assert red[0] > red[1] * 1.5


@pytest.mark.parametrize(
    "page,geometry",
    [
        (2, {"x": 0.2, "y": 0.3, "width": 0.1, "height": 0.1}),
        (1, {"x": 0.9, "y": 0.3, "width": 0.2, "height": 0.1}),
        (1, {"x": 0.2, "y": 0.3, "width": 0.0, "height": 0.1}),
    ],
)
def test_unreproducible_page_or_invalid_captured_geometry_fails_closed(page, geometry):
    with fitz.open() as document:
        document.new_page()
        content = document.tobytes()
    with pytest.raises(HTTPException) as failure:
        _render_page(content, "application/pdf", page, [{"ordinal": 1, **geometry}])
    assert failure.value.status_code == 409


@pytest.mark.parametrize(
    "language,title",
    [
        ("bg", "Визуално приложение към заявката"),
        ("en", "Visual reference appendix"),
        ("ru", "Визуальное приложение к заявке"),
    ],
)
def test_no_visual_reference_is_truthful_and_localized(
    client, auth_headers, session_factory, monkeypatch, language, title
):
    data = future_catalog(session_factory, visuals=False)
    request = _request(client, auth_headers, data)
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    generated = _generate(client, auth_headers, request, language)
    assert generated.status_code == 201, generated.text
    records = _records(session_factory, request["id"])
    with zipfile.ZipFile(io.BytesIO(records["docx"][0])) as archive:
        xml = archive.read("word/document.xml").decode()
    assert title in xml
    assert not records["docx"][2]["visual_appendix"]["lines"][0]["blocks"]
    with fitz.open(stream=records["pdf"][0], filetype="pdf") as pdf:
        assert title in "\n".join(page.get_text() for page in pdf)


def test_legacy_and_no_binding_have_distinct_historical_notes(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory, visuals=False)
    manual = {"description": "QA manual part", "quantity": 1}
    request = _request(client, auth_headers, data, lines=[manual, manual])
    with session_factory() as db:
        stored = load_request(db, request["id"])
        sorted(stored.lines, key=lambda line: line.id)[0].catalog_part_id = data["part_id"]
        db.commit()
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    assert _generate(client, auth_headers, request).status_code == 201
    records = _records(session_factory, request["id"])
    states = [line["blocks"] for line in records["docx"][2]["visual_appendix"]["lines"]]
    assert states == [[], []]
    with zipfile.ZipFile(io.BytesIO(records["docx"][0])) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "няма неизменима визуална снимка" in xml
    assert "няма потвърдена връзка с каталог" in xml


def test_two_lines_same_source_page_keep_separate_part_identity(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory)
    with session_factory() as db:
        second = PartCatalog(
            source_record_key="QA-OTHER-SOURCE:1",
            source_id="QA-OTHER-SOURCE",
            brand="QA-SYNTHETIC",
            position="QA-Q8",
            part_number="QA-OTHER-02",
            description="QA second part",
            is_verified=True,
            verification_status="VERIFIED_QA",
            source_document="qa-source.pdf",
            source_page=2,
            compatible_machine_numbers=["QA-FUTURE-9001"],
        )
        db.add(second)
        db.flush()
        db.add(
            PartHotspot(
                part_id=second.id,
                technical_document_id=data["document_id"],
                page_number=2,
                x=0.75,
                y=0.45,
                width=0.05,
                height=0.05,
                label="QA-Q8",
                is_verified=True,
                created_by_id=data["actor_id"],
            )
        )
        db.commit()
        second_id = second.id
    lines = [
        {"catalog_part_id": data["part_id"], "description": "QA first part", "quantity": 1},
        {"catalog_part_id": second_id, "description": "QA second part", "quantity": 1},
    ]
    request = _request(client, auth_headers, data, lines=lines)
    with session_factory() as db:
        appendix = prepare_appendix(db, load_request(db, request["id"]))
    assert [line.part_number for line in appendix] == ["QA-991122", "QA-OTHER-02"]
    assert len(appendix[0].blocks) == 2 and len(appendix[1].blocks) == 1
    assert appendix[0].blocks[0].page_number == appendix[1].blocks[0].page_number == 2
    assert appendix[0].blocks[0].image != appendix[1].blocks[0].image
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    assert _generate(client, auth_headers, request).status_code == 201
    records = _records(session_factory, request["id"])
    with fitz.open(stream=records["pdf"][0], filetype="pdf") as pdf:
        text_content = "\n".join(page.get_text() for page in pdf)
        assert "QA-991122" in text_content and "QA-OTHER-02" in text_content


def test_multiple_artifacts_and_large_request_remain_bounded(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory)
    with fitz.open() as source:
        source.new_page(width=500, height=300).insert_text((40, 50), "QA second artifact")
        content = source.tobytes()
    other_sha = hashlib.sha256(content).hexdigest()
    with session_factory() as db:
        document = TechnicalDocument(
            brand="QA-SYNTHETIC",
            category="QA",
            title="QA second artifact",
            file_path="qa-only/second-artifact.pdf",
            uploaded_content=content,
            uploaded_filename="qa-second.pdf",
            media_type="application/pdf",
            sha256=other_sha,
            page_count=1,
        )
        db.add(document)
        db.flush()
        db.add(
            PartHotspot(
                part_id=data["part_id"],
                technical_document_id=document.id,
                page_number=1,
                x=0.4,
                y=0.3,
                width=0.1,
                height=0.1,
                label="QA second artifact",
                is_verified=True,
                created_by_id=data["actor_id"],
            )
        )
        db.commit()
    request = _request(client, auth_headers, data)
    with session_factory() as db:
        stored = load_request(db, request["id"])
        actor = db.get(User, data["actor_id"])
        for _ in range(8):
            create_request_line(
                db,
                stored.id,
                {"catalog_part_id": data["part_id"], "description": "QA extra line", "quantity": 1},
                actor,
            )
        db.commit()
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    generated = _generate(client, auth_headers, request)
    assert generated.status_code == 201, generated.text
    records = _records(session_factory, request["id"])
    manifest = records["docx"][2]["visual_appendix"]
    assert len(manifest["lines"]) == 9
    assert all(
        {block["artifact_sha256"] for block in line["blocks"]} == {data["sha256"], other_sha}
        for line in manifest["lines"]
    )
    with fitz.open(stream=records["pdf"][0], filetype="pdf") as pdf:
        assert pdf.page_count >= 28


def test_unknown_part_uses_link_time_snapshot_without_rewriting_original(
    client, auth_headers, session_factory, monkeypatch
):
    data = future_catalog(session_factory)
    photo = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(photo, format="PNG")
    created = client.post(
        "/api/part-requests/unknown",
        headers=auth_headers,
        json={
            "machine_id": data["machine_id"],
            "assembly": "QA original assembly",
            "description": "QA original unknown description",
            "quantity": 1,
            "photo": {
                "filename": "qa-unknown.png",
                "media_type": "image/png",
                "content_base64": base64.b64encode(photo.getvalue()).decode(),
            },
        },
    )
    assert created.status_code == 201, created.text
    request = created.json()
    line_id = request["lines"][0]["id"]
    linked = client.post(
        f"/api/part-requests/{request['id']}/lines/{line_id}/link-catalog-part",
        headers=auth_headers,
        json={"catalog_part_id": data["part_id"]},
    )
    assert linked.status_code == 200, linked.text
    assert linked.json()["lines"][0]["description"] == "QA original unknown description"
    with session_factory() as db:
        appendix = prepare_appendix(db, load_request(db, request["id"]))
    assert appendix[0].state == "verified_visual_references"
    assert appendix[0].part_number == "QA-991122"
    assert appendix[0].description == "Synthetic QA rotor"
    assert appendix[0].captured_at == linked.json()["lines"][0]["linked_at"]
    submitted = client.post(f"/api/part-requests/{request['id']}/submit", headers=auth_headers)
    assert submitted.status_code == 200, submitted.text
    approved = client.post(
        f"/api/part-requests/{request['id']}/decision",
        headers=auth_headers,
        json={"decision": "APPROVED", "note": "QA approval"},
    )
    assert approved.status_code == 200, approved.text
    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", lambda _: None)
    assert _generate(client, auth_headers, request).status_code == 201
    records = _records(session_factory, request["id"])
    assert records["docx"][2]["visual_appendix"]["lines"][0]["visual_snapshot_sha256"]
