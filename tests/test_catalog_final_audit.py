"""Cross-feature acceptance checks using only disposable QA records."""

import base64
import hashlib
import io
import shutil

import pytest
from app.catalog import service as runtime_service
from app.catalog.importer import import_authoritative_catalog
from app.catalog.runtime_context import require_compatible_part
from app.catalog.sources import CATALOG_VERSION
from app.documents import part_request_documents
from app.documents.part_request_grouped_visuals import prepare_appendix
from app.models import (
    CatalogDefinition,
    CatalogDiagram,
    CatalogRevision,
    GeneratedDocument,
    Machine,
    PartCatalog,
    PartVisualSnapshot,
    RepairKit,
    TechnicalDocument,
    TechnicalDocumentRevision,
    User,
)
from app.part_requests.service import load_request
from app.seed import seed_database
from catalog_review_helpers import verify_http_revision
from docx import Document
from fastapi import HTTPException
from sqlalchemy import func, inspect, select
from test_catalog_builder_parts import BASE, part, source, workspace


def _draft(client, headers, factory):
    catalog_id, revision_id, assembly_id, _ = workspace(client, headers, factory, include_empty_group=False)
    artifact_id, spare_id, scheme_id, _ = source(client, headers, assembly_id)
    created = part(client, headers, assembly_id, position="P" * 80).json()
    response = client.post(f"{BASE}/parts/{created['id']}/source-pages", headers=headers,
                           json={"visual_page_ids": [spare_id]})
    assert response.status_code == 201, response.text
    return catalog_id, revision_id, assembly_id, artifact_id, created["id"], scheme_id


def _publish(client, headers, revision_id):
    verify_http_revision(client, headers, revision_id)
    preview = client.get(f"{BASE}/revisions/{revision_id}/publication-readiness", headers=headers)
    assert preview.status_code == 200 and preview.json()["ready"], preview.text
    response = client.post(f"{BASE}/revisions/{revision_id}/publish", headers=headers, json={
        "expected_publication_digest": preview.json()["publication_digest"],
        "expected_current_published_revision_id": preview.json()["current_published_revision_id"],
        "confirmed": True,
    })
    assert response.status_code == 200, response.text


def _machine(factory, catalog_id, number):
    with factory() as db:
        catalog = db.get(CatalogDefinition, catalog_id)
        machine = Machine(inventory_number=number, name="Disposable QA asset",
                          category="QA_PARTS_CATEGORY", category_id=catalog.asset_category_id,
                          brand="QA", model="QA")
        db.add(machine)
        db.commit()
        return machine.id


def _columns(row):
    return {column.key: getattr(row, column.key) for column in inspect(type(row)).column_attrs}


@pytest.mark.parametrize("retired", [False, True])
def test_legacy_document_revision_cannot_change_builder_evidence(
    client, auth_headers, session_factory, retired,
):
    _, revision_id, _, artifact_id, _, _ = _draft(client, auth_headers, session_factory)
    _publish(client, auth_headers, revision_id)
    with session_factory() as db:
        document = db.scalar(select(TechnicalDocument).where(
            TechnicalDocument.builder_artifact_id == artifact_id))
        document_id, before = document.id, _columns(document)
        original_content = document.uploaded_content
    if retired:
        clone = client.post(f"{BASE}/revisions/{revision_id}/clone", headers=auth_headers,
                            json={"revision_code": "B"})
        assert clone.status_code == 201, clone.text
        _publish(client, auth_headers, clone.json()["id"])
    response = client.post(f"/api/technical-library/{document_id}/revisions", headers=auth_headers,
                           json={"brand": "QA", "category": "QA", "title": "QA replacement",
                                 "filename": "replacement.pdf", "media_type": "application/pdf",
                                 "content_base64": base64.b64encode(original_content).decode()})
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["code"] == "catalog_published_content_immutable"
    with session_factory() as db:
        assert _columns(db.get(TechnicalDocument, document_id)) == before
        assert db.scalar(select(func.count(TechnicalDocumentRevision.id)).where(
            TechnicalDocumentRevision.document_id == document_id)) == 1


def test_builder_search_includes_alternative_number(client, auth_headers, session_factory, monkeypatch):
    catalog_id, revision_id, _, _, part_id, _ = _draft(client, auth_headers, session_factory)
    response = client.patch(f"{BASE}/parts/{part_id}", headers=auth_headers,
                            json={"alternative_part_number": "QA-ALTERNATIVE-ONLY",
                                  "name_bg": "Тестова част", "name_ru": "Тестовая деталь"})
    assert response.status_code == 200, response.text
    _publish(client, auth_headers, revision_id)
    machine_id = _machine(session_factory, catalog_id, "QA_SEARCH")
    assert client.post(f"{BASE}/catalogs/{catalog_id}/assets/{machine_id}",
                       headers=auth_headers).status_code == 201
    def unavailable_legacy_manifest():
        raise AssertionError("Published Builder runtime entered the legacy manifest")

    monkeypatch.setattr(runtime_service, "load_manifest", unavailable_legacy_manifest)
    runtime = client.get(f"/api/catalog/v2/machines/{machine_id}", headers=auth_headers)
    assert runtime.status_code == 200 and runtime.json()["supported"], runtime.text
    response = client.get(f"/api/catalog/v2/search?machine_id={machine_id}&q=QA-ALTERNATIVE-ONLY",
                          headers=auth_headers)
    assert response.status_code == 200, response.text
    with session_factory() as db:
        live_id = db.scalar(select(PartCatalog.id).where(PartCatalog.builder_part_id == part_id))
        # A captured unbound selection is authoritative for the whole request;
        # it must not be reselected independently for a later request line.
        with pytest.raises(HTTPException):
            require_compatible_part(db, db.get(Machine, machine_id), db.get(PartCatalog, live_id), selected=None)
        require_compatible_part(db, db.get(Machine, machine_id), db.get(PartCatalog, live_id))
    assert [row["id"] for row in response.json()] == [live_id]
    assert response.json()[0]["description_ru"] == "Тестовая деталь"
    assert response.json()[0]["description_bg"] == "Тестова част"
    assert response.json()[0]["description_en"] == "Test part"
    assert response.json()[0]["translation_version"] == "BUILDER"


def test_v2_whole_kit_cannot_bypass_machine_compatibility(client, auth_headers, session_factory):
    with session_factory() as db:
        kit = db.scalar(select(RepairKit).where(RepairKit.source_version == CATALOG_VERSION))
        allowed = set(kit.components[0].part.compatible_machine_numbers)
        machine_id = db.scalar(select(Machine.id).where(Machine.inventory_number.not_in(allowed)))
        kit_id = kit.id
    assert machine_id is not None
    response = client.post("/api/part-requests/multi", headers=auth_headers, json={
        "machine_id": machine_id, "repair_kit_id": kit_id, "repair_kit_mode": "KIT",
        "lines": [{"description": "QA incompatible kit", "quantity": 1}],
    })
    assert response.status_code == 409, response.text


def test_importer_never_archives_builder_document_even_if_bytes_are_missing(
    client, auth_headers, session_factory,
):
    _, revision_id, _, artifact_id, _, _ = _draft(client, auth_headers, session_factory)
    _publish(client, auth_headers, revision_id)
    with session_factory() as db:
        document = db.scalar(select(TechnicalDocument).where(
            TechnicalDocument.builder_artifact_id == artifact_id))
        document.uploaded_content = None
        db.commit()
        document_id = document.id
        before = _columns(document)
        import_authoritative_catalog(db, db.get(User, 1))
        db.commit()
        db.expire_all()
        assert _columns(db.get(TechnicalDocument, document_id)) == before


def test_late_binding_shared_revision_documents_and_repeated_seed(
    client, auth_headers, session_factory, monkeypatch, tmp_path,
):
    converted = []
    converter = part_request_documents.convert_docx_to_pdf

    def actual_conversion(content):
        result = converter(content)
        converted.append(result is not None)
        return result

    monkeypatch.setattr(part_request_documents, "convert_docx_to_pdf", actual_conversion)
    catalog_id, revision_id, _, _, part_id, scheme_id = _draft(client, auth_headers, session_factory)
    hotspot = client.post(f"{BASE}/visual-pages/{scheme_id}/hotspots", headers=auth_headers,
                          json={"position": "P" * 80, "x": .2, "y": .3, "width": .1, "height": .1})
    assert hotspot.status_code == 201, hotspot.text
    verified = client.post(f"{BASE}/hotspots/{hotspot.json()['id']}/verify", headers=auth_headers,
                           json={"expected_version": 1})
    assert verified.status_code == 200, verified.text
    _publish(client, auth_headers, revision_id)
    with session_factory() as db:
        live = db.scalar(select(PartCatalog).where(PartCatalog.builder_part_id == part_id))
        live_id, live_before = live.id, _columns(live)
        assert live.compatible_machine_numbers in (None, [])
    machine_a = _machine(session_factory, catalog_id, "QA_LATE_A")
    assert client.post(f"{BASE}/catalogs/{catalog_id}/assets/{machine_a}",
                       headers=auth_headers).status_code == 201
    # B is created after publication and matches A's model without being bound.
    machine_b = _machine(session_factory, catalog_id, "QA_LATE_B")
    assert not client.get(f"/api/catalog/v2/machines/{machine_b}", headers=auth_headers).json()["supported"]
    request_body = {"machine_id": machine_b, "lines": [{"catalog_part_id": live_id, "quantity": 1,
                    "description": "QA shared part"}], "submit_for_approval": True}
    assert client.post("/api/part-requests/multi", headers=auth_headers, json=request_body).status_code == 409
    assert client.post(f"{BASE}/catalogs/{catalog_id}/assets/{machine_b}",
                       headers=auth_headers).status_code == 201
    request_ids = []
    for machine_id in (machine_a, machine_b):
        runtime = client.get(f"/api/catalog/v2/machines/{machine_id}", headers=auth_headers)
        assert runtime.status_code == 200, runtime.text
        assert runtime.json()["dataset_version"] == f"CATALOG_BUILDER_R{revision_id}"
        request_body["machine_id"] = machine_id
        response = client.post("/api/part-requests/multi", headers=auth_headers, json=request_body)
        assert response.status_code == 201, response.text
        request_id = response.json()["id"]
        request_ids.append(request_id)
        assert response.json()["lines"][0]["position"] == "P" * 80
        assert client.post(f"/api/part-requests/{request_id}/decision", headers=auth_headers,
                           json={"decision": "APPROVED"}).status_code == 200
        generated = client.post(f"/api/part-requests/{request_id}/documents?language=bg",
                                headers=auth_headers)
        assert generated.status_code == 201, generated.text
    with session_factory() as db:
        assert _columns(db.get(PartCatalog, live_id)) == live_before
        snapshot_before = {row.id: _columns(row) for row in db.scalars(select(PartVisualSnapshot))}
        documents_before = {row.id: _columns(row) for row in db.scalars(select(GeneratedDocument))}
        for request_id in request_ids:
            appendix = prepare_appendix(db, load_request(db, request_id))
            assert [page.role for page in appendix.pages] == ["EXPLODED_SCHEME", "SPARE_PARTS_LIST"]
            records = db.scalars(select(GeneratedDocument).where(
                GeneratedDocument.part_request_id == request_id)).all()
            assert {row.format for row in records} == {"docx", "pdf"}
            for row in records:
                assert hashlib.sha256(row.content).hexdigest() == row.sha256
                (tmp_path / f"builder-request-{request_id}.{row.format}").write_bytes(row.content)
                if row.format == "docx":
                    document = Document(io.BytesIO(row.content))
                    assert any([cell.text for cell in table.rows[0].cells] ==
                               ["Поз.", "PART №", "Описание", "Количество"] for table in document.tables)
                    assert any(cell.text == "P" * 80 for table in document.tables
                               for row in table.rows for cell in row.cells)
        if shutil.which("soffice"):
            assert converted == [True, True]
        # Initialization and import may touch only their own operational records.
        for _ in range(2):
            seed_database(db)
            import_authoritative_catalog(db, db.get(User, 1))
            db.commit()
        assert _columns(db.get(PartCatalog, live_id)) == live_before
    clone = client.post(f"{BASE}/revisions/{revision_id}/clone", headers=auth_headers,
                        json={"revision_code": "B"})
    assert clone.status_code == 201, clone.text
    next_id = clone.json()["id"]
    assemblies = client.get(f"{BASE}/revisions/{next_id}/assemblies", headers=auth_headers).json()
    next_part = client.get(f"{BASE}/assemblies/{assemblies[0]['id']}/parts", headers=auth_headers).json()[0]
    assert client.patch(f"{BASE}/parts/{next_part['id']}", headers=auth_headers,
                        json={"name_en": "QA revision two"}).status_code == 200
    _publish(client, auth_headers, next_id)
    for machine_id in (machine_a, machine_b):
        assert client.get(f"/api/catalog/v2/machines/{machine_id}", headers=auth_headers).json()["dataset_version"] == f"CATALOG_BUILDER_R{next_id}"
    with session_factory() as db:
        assert db.get(CatalogRevision, revision_id).status == "RETIRED"
        assert {row.id: _columns(row) for row in db.scalars(select(PartVisualSnapshot))} == snapshot_before
        assert {row.id: _columns(row) for row in db.scalars(select(GeneratedDocument))} == documents_before
        assert db.scalar(select(func.count(CatalogDiagram.id)).where(
            CatalogDiagram.builder_revision_id == revision_id)) == 1


def test_multiple_exact_pages_variants_and_foreign_machine_ids(client, auth_headers, session_factory):
    catalog_id, revision_id, assembly_id, artifact_id, part_id, scheme_id = _draft(
        client, auth_headers, session_factory)
    # One part has two exact list-page references; the same PDF page can have
    # explicit different roles. It is never inferred from its filename.
    extra = client.post(f"{BASE}/artifacts/{artifact_id}/visual-pages", headers=auth_headers,
                        json={"role": "SPARE_PARTS_LIST", "page_numbers": [2]})
    assert extra.status_code == 201, extra.text
    assert client.post(f"{BASE}/parts/{part_id}/source-pages", headers=auth_headers,
                       json={"visual_page_ids": [extra.json()[0]["id"]]}).status_code == 201
    variant = part(client, auth_headers, assembly_id, position="P" * 80, number="QA_VARIANT").json()
    assert client.post(f"{BASE}/parts/{variant['id']}/source-pages", headers=auth_headers,
                       json={"visual_page_ids": [extra.json()[0]["id"]]}).status_code == 201
    kit = client.post(f"{BASE}/assemblies/{assembly_id}/repair-kits", headers=auth_headers,
                      json={"code": "QA_SCOPE_KIT", "name_bg": "QA scope kit",
                            "source_visual_page_id": extra.json()[0]["id"]})
    assert kit.status_code == 201, kit.text
    assert client.post(f"{BASE}/repair-kits/{kit.json()['id']}/components", headers=auth_headers,
                       json={"part_id": part_id, "quantity": 1}).status_code == 201
    for x in (.1, .4):
        hotspot = client.post(f"{BASE}/visual-pages/{scheme_id}/hotspots", headers=auth_headers,
                              json={"position": "P" * 80, "x": x, "y": .2, "width": .1, "height": .1})
        assert hotspot.status_code == 201, hotspot.text
        assert client.post(f"{BASE}/hotspots/{hotspot.json()['id']}/verify", headers=auth_headers,
                           json={"expected_version": 1}).status_code == 200
    _publish(client, auth_headers, revision_id)
    machine_id = _machine(session_factory, catalog_id, "QA_MULTIPAGE")
    foreign_machine = _machine(session_factory, catalog_id, "QA_FOREIGN")
    assert client.post(f"{BASE}/catalogs/{catalog_id}/assets/{machine_id}",
                       headers=auth_headers).status_code == 201
    runtime = client.get(f"/api/catalog/v2/machines/{machine_id}", headers=auth_headers).json()
    source_id = runtime["assemblies"][0]["source_id"]
    detail = client.get(f"/api/catalog/v2/assemblies/{source_id}?machine_id={machine_id}",
                        headers=auth_headers).json()
    diagram_id = detail["diagrams"][0]["id"]
    hotspots = client.get(f"/api/catalog/v2/diagrams/{diagram_id}/hotspots?machine_id={machine_id}",
                          headers=auth_headers)
    assert hotspots.status_code == 200, hotspots.text
    assert len(hotspots.json()) == 2
    ids = {row["id"] for row in detail["parts"]}
    assert all({row["id"] for row in hotspot["variants"]} == ids for hotspot in hotspots.json())
    # Another explicit catalog of the same category does not grant access.
    with session_factory() as db:
        category_id = db.get(CatalogDefinition, catalog_id).asset_category_id
    foreign_catalog = client.post(f"{BASE}/catalogs", headers=auth_headers, json={
        "code": "QA_FOREIGN_CATALOG", "asset_category_id": category_id,
        "name_bg": "QA", "name_en": "QA", "name_ru": "QA",
    }).json()["id"]
    assert client.post(f"{BASE}/catalogs/{foreign_catalog}/assets/{foreign_machine}",
                       headers=auth_headers).status_code == 201
    foreign_revision = client.post(f"{BASE}/catalogs/{foreign_catalog}/revisions", headers=auth_headers,
                                   json={"revision_code": "A"}).json()["id"]
    foreign_assembly = client.post(f"{BASE}/revisions/{foreign_revision}/assemblies", headers=auth_headers,
                                   json={"code": "PUMP", "name_bg": "QA", "name_en": "QA", "name_ru": "QA"}).json()["id"]
    _, foreign_spare, _, _ = source(client, auth_headers, foreign_assembly, marker="QA foreign")
    foreign_part = part(client, auth_headers, foreign_assembly).json()["id"]
    assert client.post(f"{BASE}/parts/{foreign_part}/source-pages", headers=auth_headers,
                       json={"visual_page_ids": [foreign_spare]}).status_code == 201
    _publish(client, auth_headers, foreign_revision)
    runtime_kits = client.get(f"/api/catalog/v2/repair-kits?machine_id={machine_id}",
                              headers=auth_headers)
    assert runtime_kits.status_code == 200 and len(runtime_kits.json()) == 1, runtime_kits.text
    foreign_kit_request = client.post("/api/part-requests/multi", headers=auth_headers, json={
        "machine_id": foreign_machine, "repair_kit_id": runtime_kits.json()[0]["id"],
        "repair_kit_mode": "KIT", "lines": [{"description": "QA", "quantity": 1}],
    })
    assert foreign_kit_request.status_code == 409, foreign_kit_request.text
    assert client.get(f"/api/catalog/v2/assemblies/{source_id}?machine_id={foreign_machine}",
                      headers=auth_headers).status_code in (404, 409)
    assert client.get(f"/api/catalog/v2/diagrams/{diagram_id}/hotspots?machine_id={foreign_machine}",
                      headers=auth_headers).status_code in (404, 409)
    foreign_parts = client.get(f"/api/catalog/parts?machine_id={foreign_machine}", headers=auth_headers).json()
    assert len(foreign_parts) == 1 and foreign_parts[0]["id"] not in ids
    body = {"machine_id": foreign_machine,
            "lines": [{"catalog_part_id": value, "description": "QA", "quantity": 1} for value in ids]}
    assert client.post("/api/part-requests/multi", headers=auth_headers, json=body).status_code == 409
    body["machine_id"] = machine_id
    request = client.post("/api/part-requests/multi", headers=auth_headers, json=body)
    assert request.status_code == 201, request.text
    with session_factory() as db:
        appendix = prepare_appendix(db, load_request(db, request.json()["id"]))
        assert [page.role for page in appendix.pages] == ["EXPLODED_SCHEME", "SPARE_PARTS_LIST", "SPARE_PARTS_LIST"]
        assert {page.page_number for page in appendix.pages[1:]} == {1, 2}
        # Each of the two requested variants contributes both occurrences.
        assert len(appendix.pages[0].contributions) == 4
        assert {item["marker"] for item in appendix.pages[0].contributions} == {1, 2}
        assert len({(item["line_id"], item["occurrence_ordinal"])
                    for item in appendix.pages[0].contributions}) == 4
