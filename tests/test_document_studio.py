"""Foundation security and preservation contracts; disposable data only."""

import base64
import copy
import hashlib
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from app.document_studio import routes
from app.document_studio.schemas import DraftUpdate, IdentityUpdate, Layout, initial_layout
from app.models import (
    AuditLog,
    CompanyProfile,
    DocumentEditorDraft,
    DocumentTemplate,
    DocumentTemplateVersion,
    Machine,
    User,
)
from app.permissions import Permission, has_permission
from app.security import create_access_token
from fastapi import HTTPException
from PIL import Image
from pydantic import ValidationError
from sqlalchemy import func, select


def create_template(client, headers):
    response = client.post('/api/studio/templates', headers=headers, json={
        'code': 'QA_STUDIO', 'document_type': 'PART_REQUEST',
        'name_bg': 'Тестов шаблон', 'name_en': 'QA template', 'name_ru': 'Тестовый шаблон'})
    assert response.status_code == 201, response.text
    return response.json()['id']


def create_draft(client, headers, template_id):
    response = client.post(f'/api/studio/templates/{template_id}/drafts', headers=headers,
                           json={'title': 'QA draft', 'language': 'bg', 'design': 'blank'})
    assert response.status_code == 201, response.text
    return response.json()


def image_payload(content=None, **extra):
    if content is None:
        stream = io.BytesIO()
        Image.new('RGBA', (80, 40), '#183153').save(stream, format='PNG')
        content = stream.getvalue()
    return {'filename': 'synthetic.png', 'media_type': 'image/png', 'purpose': 'PRIMARY',
            'content_base64': base64.b64encode(content).decode(), **extra}


def test_company_is_neutral_and_revision_protected(client, auth_headers, session_factory):
    profile = client.get('/api/studio/company', headers=auth_headers).json()
    assert profile['revision'] == 0
    assert profile['identity']['legal_name'] == 'AssetCore'
    assert profile['identity']['primary_logo_id'] is None
    with session_factory() as db:
        assert db.scalar(select(func.count(CompanyProfile.id))) == 0
    profile['identity'].update(legal_name='Synthetic QA company', registration_number='INT-QA-001',
                               country='QA', timezone='Europe/Sofia', default_language='en')
    payload = {'expected_revision': 0, 'identity': profile['identity']}
    response = client.put('/api/studio/company', headers=auth_headers, json=payload)
    assert response.status_code == 200, response.text
    assert response.json()['revision'] == 1
    assert client.put('/api/studio/company', headers=auth_headers, json=payload).status_code == 409
    payload['expected_revision'] = 1
    assert client.put('/api/studio/company', headers=auth_headers, json=payload).json()['revision'] == 2
    assert client.put('/api/studio/company', headers=auth_headers, json=payload).status_code == 409
    with session_factory() as db:
        assert db.scalar(select(func.count(AuditLog.id)).where(AuditLog.entity_type == 'company_profile')) == 2


@pytest.mark.parametrize('change', [
    {'registration_number': 'x' * 101}, {'email': 'bad'}, {'website': 'javascript:alert(1)'},
    {'website': 'https://user:secret@example.invalid'}, {'primary_color': 'url(http://unsafe)'},
    {'timezone': '../../private'}, {'primary_logo_id': 777}, {'identity_image_ids': [1, 1]},
    {'default_language': 'xx'}, {'legal_name': ''}, {'unexpected': True},
])
def test_invalid_identity_rejected(client, auth_headers, change):
    identity = client.get('/api/studio/company', headers=auth_headers).json()['identity']
    identity.update(change)
    assert client.put('/api/studio/company', headers=auth_headers,
                      json={'expected_revision': 0, 'identity': identity}).status_code == 422


def test_safe_images_authorized_normalized_and_selected(client, auth_headers, viewer_headers):
    payload = image_payload()
    raw = base64.b64decode(payload['content_base64']) + b'<script>unsafe trailing payload</script>'
    payload['content_base64'] = base64.b64encode(raw).decode()
    uploaded = client.post('/api/studio/images', headers=auth_headers, json=payload)
    assert uploaded.status_code == 201, uploaded.text
    image = uploaded.json()
    endpoint = '/api' + image['download_endpoint']
    result = client.get(endpoint, headers=auth_headers)
    assert result.status_code == 200
    assert result.headers['content-type'] == 'image/png'
    assert 'no-store' in result.headers['cache-control'] and 'private' in result.headers['cache-control']
    assert b'<script>' not in result.content
    assert hashlib.sha256(result.content).hexdigest() == image['sha256']
    assert client.get(endpoint).status_code == 401
    assert client.get(endpoint, headers=viewer_headers).status_code == 403
    identity = client.get('/api/studio/company', headers=auth_headers).json()['identity']
    identity.update(primary_logo_id=image['id'], identity_image_ids=[image['id']])
    saved = client.put('/api/studio/company', headers=auth_headers,
                       json={'expected_revision': 0, 'identity': identity})
    assert saved.status_code == 200
    assert saved.json()['identity']['primary_logo_id'] == image['id']


@pytest.mark.parametrize('payload', [
    image_payload(b'<svg onload="alert(1)"></svg>'), image_payload(b'not an image'),
    image_payload(media_type='image/svg+xml'), image_payload(filename='../image.png'),
    image_payload(filename='image.jpg'), image_payload(content_base64='bad!'),
    image_payload(b'X' * 2_000_001),
])
def test_hostile_images_rejected(client, auth_headers, payload):
    response = client.post('/api/studio/images', headers=auth_headers, json=payload)
    assert response.status_code == 422, response.text
    assert client.get('/api/studio/images', headers=auth_headers).json() == []


@pytest.mark.parametrize('role,read,write', [('administrator', 200, 200), ('director', 200, 403),
                                          ('mechanic', 403, 403), ('observer', 403, 403)])
def test_permissions_without_parallel_owner_role(client, session_factory, role, read, write):
    with session_factory() as db:
        owner = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        user = User(email=f'{role}@qa.invalid', full_name='Synthetic QA actor', password_hash=owner.password_hash,
                    role=role, is_active=True, is_system_owner=False, first_name='QA', last_name='Actor',
                    job_title='QA', profile_status='PROFILE_COMPLETE')
        db.add(user)
        db.commit()
        token = create_access_token(user)
        assert not has_permission(user, Permission.SETTINGS_MANAGE) if role != 'administrator' else has_permission(user, Permission.COMPANY_MANAGE)
    headers = {'Authorization': f'Bearer {token}'}
    for path in ('company', 'templates', 'images', 'schema'):
        assert client.get('/api/studio/' + path, headers=headers).status_code == read
    identity = routes.CompanyIdentity().model_dump()
    assert client.put('/api/studio/company', headers=headers,
                      json={'expected_revision': 0, 'identity': identity}).status_code == write
    if role != 'administrator':
        assert client.post('/api/studio/templates/1/drafts', headers=headers,
                           json={'title': 'QA'}).status_code == 403
        assert client.put('/api/studio/drafts/1', headers=headers, json={
            'expected_revision': 1, 'title': 'QA', 'layout': initial_layout('PART_REQUEST').model_dump()}).status_code == 403
        assert client.post('/api/studio/images', headers=headers, json=image_payload()).status_code == 403


def test_draft_roundtrip_and_immutable_template_separation(client, auth_headers, session_factory):
    template_id = create_template(client, auth_headers)
    draft = create_draft(client, auth_headers, template_id)
    assert draft['status'] == 'DRAFT' and draft['renderer_enabled'] is True
    layout = copy.deepcopy(draft['layout'])
    layout['page']['orientation'] = 'landscape'
    payload = {'expected_revision': 1, 'title': 'QA updated', 'layout': layout}
    saved = client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers, json=payload)
    assert saved.status_code == 200, saved.text
    assert saved.json()['revision'] == 2
    assert saved.json()['layout']['page']['orientation'] == 'landscape'
    assert client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers, json=payload).status_code == 409
    assert client.get(f"/api/studio/drafts/{draft['id']}", headers=auth_headers).json()['title'] == 'QA updated'
    with session_factory() as db:
        assert db.scalar(select(func.count(DocumentTemplateVersion.id)).where(DocumentTemplateVersion.template_id == template_id)) == 0
    assert client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=auth_headers).status_code == 422


@pytest.mark.parametrize('mutation', ['signature', 'system', 'expression', 'javascript', 'absolute', 'duplicate', 'width', 'wrong_workflow'])
def test_strict_layout_contract(client, auth_headers, mutation):
    draft = create_draft(client, auth_headers, create_template(client, auth_headers))
    layout = copy.deepcopy(draft['layout'])
    if mutation == 'signature':
        layout['blocks'] = [b for b in layout['blocks'] if b['type'] != 'signature']
    elif mutation == 'system':
        layout['blocks'] = [b for b in layout['blocks'] if b.get('field') != 'DOCUMENT_NUMBER']
    elif mutation == 'expression':
        layout['blocks'][0]['field'] = '__import__("os").system("command")'
    elif mutation == 'javascript':
        layout['script'] = 'alert(1)'
    elif mutation == 'absolute':
        layout['positioning'] = 'absolute'
    elif mutation == 'duplicate':
        layout['blocks'].append(layout['blocks'][0])
    else:
        layout['blocks'].append({'id': 'rows', 'type': 'dynamic_table',
            'source': 'REQUEST_LINES' if mutation == 'width' else 'PARTS_USED',
            'columns': [{'key': 'description', 'label': 'QA', 'width_percent': 50 if mutation == 'width' else 100}]})
    assert client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers,
                      json={'expected_revision': 1, 'title': 'QA', 'layout': layout}).status_code == 422


def test_typed_flow_table_and_unknown_versions():
    layout = initial_layout('PART_REQUEST').model_dump()
    layout['blocks'].append({'id': 'request_rows', 'type': 'dynamic_table', 'source': 'REQUEST_LINES',
        'columns': [{'key': 'description', 'label': 'QA description', 'width_percent': 70},
                    {'key': 'quantity', 'label': 'QA quantity', 'width_percent': 30}]})
    parsed = Layout.model_validate(layout)
    assert parsed.blocks[-1].repeat_header and parsed.blocks[-1].flow
    layout['schema_version'] = 2
    assert Layout.model_validate(layout).schema_version == 2
    layout['schema_version'] = 3
    assert Layout.model_validate(layout).schema_version == 3
    layout['schema_version'] = 4
    assert Layout.model_validate(layout).schema_version == 4
    layout['schema_version'] = 5
    with pytest.raises(ValidationError):
        Layout.model_validate(layout)


def test_docx_reference_preserves_bytes_and_cannot_publish(client, auth_headers):
    from docx import Document
    stream = io.BytesIO()
    Document().save(stream)
    original = stream.getvalue()
    template_id = create_template(client, auth_headers)
    response = client.post(f'/api/studio/templates/{template_id}/import', headers=auth_headers,
        json={'filename': 'synthetic.docx', 'content_base64': base64.b64encode(original).decode(),
              'language': 'bg', 'change_note': 'Synthetic QA reference'})
    assert response.status_code == 201, response.text
    version = response.json()
    assert version['structural_conversion'] is False
    assert client.get(f"/api/studio/versions/{version['id']}/source", headers=auth_headers).content == original
    assert client.post(f"/api/document-template-versions/{version['id']}/publish", headers=auth_headers).status_code == 409


@pytest.mark.parametrize('name,data', [
    ('../unsafe.xml', b'<unsafe/>'), ('word/vbaProject.bin', b'active'),
    ('word/_rels/unsafe.xml.rels', b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>'),
    ('word/active.xml', b'<!DOCTYPE test [<!ENTITY x SYSTEM "file:///private">]><test>&x;</test>'),
    ('word/dde.xml', b'<instrText>DDEAUTO unsafe</instrText>'),
    ('word/renamed-active.xml', b'<Override ContentType="application/vnd.ms-office.vbaProject" PartName="/word/code.dat"/>'),
    ('word/renamed-ole.xml', b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="code.dat"/>'),
    ('word/unsafe-vml.xml', b'<imagedata src="https://example.invalid/image.png"/>'),
    ('word/renamed-vector.xml', b'<svg xmlns="http://www.w3.org/2000/svg" onload="alert(1)"/>'),
])
def test_hostile_docx_blocked(client, auth_headers, name, data):
    from docx import Document
    stream = io.BytesIO()
    Document().save(stream)
    with zipfile.ZipFile(stream, 'a') as archive:
        archive.writestr(name, data)
    response = client.post('/api/studio/templates/1/import', headers=auth_headers, json={
        'filename': 'synthetic.docx', 'content_base64': base64.b64encode(stream.getvalue()).decode(),
        'language': 'bg', 'change_note': 'QA blocked'})
    assert response.status_code == 422


def test_existing_business_material_unchanged(client, auth_headers, session_factory):
    def evidence():
        with session_factory() as db:
            machines = [(m.id, m.inventory_number, m.serial_number, m.brand, m.model) for m in db.scalars(select(Machine).order_by(Machine.id))]
            versions = [(v.id, v.version, v.source_sha256, v.is_published, v.layout_contract,
                         hashlib.sha256(v.source_content or b'').hexdigest())
                        for v in db.scalars(select(DocumentTemplateVersion).order_by(DocumentTemplateVersion.id))]
            return json.dumps([machines, versions], sort_keys=True)
    before = evidence()
    identity = client.get('/api/studio/company', headers=auth_headers).json()['identity']
    identity['legal_name'] = 'Synthetic QA only'
    assert client.put('/api/studio/company', headers=auth_headers,
        json={'expected_revision': 0, 'identity': identity}).status_code == 200
    create_draft(client, auth_headers, 1)
    assert evidence() == before


def test_two_draft_writers_only_one_commits(session_factory):
    with session_factory() as db:
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        template = db.scalar(select(DocumentTemplate))
        draft = routes.create_draft(template.id, routes.DraftCreate(title='QA race'), actor, db)
        actor_id = actor.id
    barrier = Barrier(2)

    def worker(index):
        with session_factory() as db:
            user = db.get(User, actor_id)
            barrier.wait(timeout=10)
            try:
                routes.update_draft(draft['id'], DraftUpdate(expected_revision=1,
                    title=f'QA {index}', layout=Layout.model_validate(draft['layout'])), user, db)
                return 200
            except HTTPException as exc:
                return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(worker, [1, 2])) == [200, 409]
    with session_factory() as db:
        assert db.get(DocumentEditorDraft, draft['id']).revision == 2


def test_two_first_company_writers_only_one_commits(session_factory):
    barrier = Barrier(2)

    def worker(index):
        with session_factory() as db:
            user = db.scalar(select(User).where(User.is_system_owner.is_(True)))
            barrier.wait(timeout=10)
            try:
                routes.update_company(IdentityUpdate(expected_revision=0,
                    identity={'legal_name': f'Synthetic QA {index}'}), user, db)
                return 200
            except HTTPException as exc:
                return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(worker, [1, 2])) == [200, 409]


def test_company_and_drafts_preserve_signed_document_archive(
    client, auth_headers, session_factory, machine_ids, issue_payload, finalize_signatures,
):
    from app.models import Base

    issued = client.post('/api/transfers/bulk-issue', headers=auth_headers,
                         json=issue_payload(next(iter(machine_ids.values()))))
    assert issued.status_code == 201, issued.text
    finalize_signatures(client, issued)
    tables = ('protocol_documents', 'official_document_versions', 'document_signatures',
              'generated_documents', 'document_participants', 'signature_sessions')

    def fingerprints():
        with session_factory() as db:
            result = {}
            for name in tables:
                rows = db.execute(select(Base.metadata.tables[name])).all()
                result[name] = hashlib.sha256(repr(rows).encode()).hexdigest()
            assert db.scalar(select(func.count()).select_from(Base.metadata.tables['document_signatures'])) > 0
            return result

    before = fingerprints()
    identity = client.get('/api/studio/company', headers=auth_headers).json()['identity']
    identity['legal_name'] = 'Synthetic QA archive preservation'
    assert client.put('/api/studio/company', headers=auth_headers,
        json={'expected_revision': 0, 'identity': identity}).status_code == 200
    create_draft(client, auth_headers, 1)
    assert fingerprints() == before


def test_company_timezone_on_hosts_without_system_iana_database():
    import zoneinfo

    previous = zoneinfo.TZPATH
    try:
        zoneinfo.reset_tzpath([])
        zoneinfo.ZoneInfo.clear_cache()
        identity = routes.CompanyIdentity(timezone='Europe/Sofia')
        assert identity.timezone == 'Europe/Sofia'
        with pytest.raises(ValidationError):
            routes.CompanyIdentity(timezone='QA/Unknown')
    finally:
        zoneinfo.reset_tzpath(previous)
        zoneinfo.ZoneInfo.clear_cache()
