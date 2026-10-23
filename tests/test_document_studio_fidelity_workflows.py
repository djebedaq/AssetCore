"""Four actual workflows using isolated synthetic equipment and frozen archives."""
import io
import json
import os
import shutil
import zipfile
from pathlib import Path

import fitz
import pytest
from app.document_studio.renderer import digest, manifest_from_source
from app.models import (
    CatalogVisualPartMap,
    CatalogVisualSource,
    DocumentTemplateVersion,
    GeneratedDocument,
    Machine,
    OfficialDocument,
    OfficialDocumentVersion,
    PartCatalog,
    RequestApprovalSnapshot,
    User,
)
from docx import Document
from sqlalchemy import select


def test_professional_four_workflows_signing_and_controlled_appendices(
    client, auth_headers, machine_ids, issue_payload, finalize_signatures, session_factory, tmp_path, monkeypatch,
):
    if os.name == 'nt' or not shutil.which('libreoffice'):
        pytest.skip('Real bounded conversion runs on Linux/Docker and CI')
    from test_document_studio_publication import (
        test_real_canonical_issue_return_sign_repair_and_kk_use_selected_studio_sources,
    )

    part_id = client.get('/api/catalog/parts', headers=auth_headers).json()[0]['id']
    with session_factory() as db:
        original_ids = {m.id for m in db.scalars(select(Machine))}
        category_id = db.get(Machine, machine_ids['9']).category_id
        synthetic_ids = {}
        for code in ('4', '8', '9'):
            machine = Machine(inventory_number=f'SYNTHETIC-QA-{code}', name='SYNTHETIC QA ONLY',
                category='SYNTHETIC QA CATEGORY', category_id=category_id, brand='SYNTHETIC QA MANUFACTURER',
                manufacturer='SYNTHETIC QA MANUFACTURER', model='SYNTHETIC QA MODEL', serial_number=f'SYNTHETIC-QA-SERIAL-{code}',
                pressure_bar=500, status='READY', ownership='SYNTHETIC QA ONLY')
            db.add(machine)
            db.flush()
            synthetic_ids[code] = machine.id
        from app.catalog.references import ReferenceCreate, create
        part = db.get(PartCatalog, part_id)
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        actor.first_name, actor.middle_name, actor.last_name = 'Синтетичен', 'Тестов', 'Подписващ'
        actor.full_name = 'Синтетичен Тестов Подписващ'
        actor.job_title, actor.department_id = 'SYNTHETIC QA approver and repairer', None
        scheme = db.scalar(select(CatalogVisualSource).where(CatalogVisualSource.source_id == part.source_id,
            CatalogVisualSource.role == 'EXPLODED_SCHEME'))
        parts_list = db.scalar(select(CatalogVisualSource).join(CatalogVisualPartMap,
            CatalogVisualPartMap.visual_source_id == CatalogVisualSource.id).where(CatalogVisualPartMap.part_id == part_id,
                CatalogVisualSource.role == 'SPARE_PARTS_LIST'))
        assert scheme is not None and parts_list is not None
        # Audited approval in this disposable QA DB, through the existing
        # compatibility service; verified source parts are never modified.
        create(db, actor, ReferenceCreate(machine_ids=[synthetic_ids['9']], scheme_id=scheme.id,
            parts_list_id=parts_list.id, part_ids=[part_id], source_revision=scheme.catalog_revision,
            compatibility_confirmed=True, reason='SYNTHETIC QA ONLY isolated reference approval'))
        db.commit()
    output = Path(os.environ.get('ASSETCORE_STUDIO_FIDELITY_WORKFLOW_OUTPUT', str(tmp_path)))
    monkeypatch.setenv('ASSETCORE_STUDIO_QA_OUTPUT', str(tmp_path / 'intermediate'))
    test_real_canonical_issue_return_sign_repair_and_kk_use_selected_studio_sources(
        True, client, auth_headers, synthetic_ids, issue_payload, finalize_signatures, session_factory, tmp_path,
    )
    # Canonical KK uses its existing REQUESTED_BY / APPROVED_BY slots and the
    # normal sequential encrypted signature sessions, including after creation.
    from test_production_hardening import _external_signer, _signature_payload
    with session_factory() as db:
        kk = db.scalar(select(OfficialDocument).where(OfficialDocument.document_type == 'PART_REQUEST',
            OfficialDocument.machine_id == synthetic_ids['9']))
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        kk_id, actor_id = kk.id, actor.id
        approval = db.scalar(select(RequestApprovalSnapshot))
        approval_id, approval_bytes = approval.request_id, approval.template_content
        unsigned = db.get(OfficialDocumentVersion, kk.current_version_id)
        with zipfile.ZipFile(io.BytesIO(unsigned.docx_content)) as package:
            appendix_bytes = {name: package.read(name) for name in package.namelist() if name.startswith('word/media/')}
        old_generated = {d.id: (d.content, d.sha256) for d in db.scalars(select(GeneratedDocument))}
    wrong = _external_signer(client, auth_headers, 'studio-fidelity-wrong')
    rejected = client.post(f'/api/official-documents/{kk_id}/participants', headers=auth_headers, json={'participants': [
        {'slot_code': 'REQUESTED_BY', 'operation_role': 'Synthetic requester', 'external_signer_id': wrong['id']},
        {'slot_code': 'APPROVED_BY', 'operation_role': 'Synthetic approver', 'user_id': actor_id}]})
    assert rejected.status_code == 409 and rejected.json()['detail']['code'] == 'studio_signature_identity_mismatch'
    prepared = client.post(f'/api/official-documents/{kk_id}/participants', headers=auth_headers, json={'participants': [
        {'slot_code': code, 'operation_role': 'Synthetic QA role', 'user_id': actor_id} for code in ('REQUESTED_BY', 'APPROVED_BY')]})
    assert prepared.status_code == 200, prepared.text
    people = {p['slot_code']: p for p in prepared.json()['participants']}
    for index, code in enumerate(('REQUESTED_BY', 'APPROVED_BY'), 1):
        opened = client.post('/api/signatures/sessions', headers=auth_headers, json={'participant_id': people[code]['id']})
        assert opened.status_code == 201
        token = opened.json()['signing_token']
        summary = client.get(f'/api/signing/{token}').json()
        assert client.post(f'/api/signing/{token}', json=_signature_payload(summary['consent_notice'], index)).status_code == 201
        if code == 'APPROVED_BY':
            from app.document_studio import renderer
            from app.template_engine import TemplateValidationError
            with session_factory() as db:
                current = db.get(OfficialDocumentVersion, db.get(OfficialDocument, kk_id).current_version_id)
                before_failure = (current.docx_content, current.pdf_content, current.status)
            def unavailable(_content):
                raise TemplateValidationError('studio_converter_unavailable')
            with monkeypatch.context() as scoped:
                scoped.setattr(renderer, 'pdf_required', unavailable)
                assert client.post(f'/api/signing/{token}/confirm').status_code == 500
            with session_factory() as db:
                current = db.get(OfficialDocumentVersion, db.get(OfficialDocument, kk_id).current_version_id)
                assert (current.docx_content, current.pdf_content, current.status) == before_failure
        confirmed = client.post(f'/api/signing/{token}/confirm')
        assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()['document_status'] == 'SIGNED'
    with session_factory() as db:
        kk = db.get(OfficialDocument, kk_id)
        signed_version_id = kk.current_version_id
        signed = db.get(OfficialDocumentVersion, signed_version_id)
        parsed = Document(io.BytesIO(signed.docx_content))
        with zipfile.ZipFile(io.BytesIO(signed.docx_content)) as package:
            assert appendix_bytes and all(package.read(name) == content for name, content in appendix_bytes.items())
        def tables_in(container):
            for table in container.tables:
                yield table
                for row in table.rows:
                    for cell in row.cells:
                        yield from tables_in(cell)
        slots = next(t for t in tables_in(parsed) if len(t.columns) == 3 and len(t.rows) == 2 and 'Заявител' in t.cell(0, 0).text)
        for row in slots.rows:
            assert len(row.cells[2]._tc.xpath('.//w:drawing')) == 1
            assert 'Синтетичен Тестов Подписващ' in row.cells[1].text
        assert not any('Потвърдени подписи' in p.text for p in parsed.paragraphs)
        signed_bytes = (signed.docx_content, signed.pdf_content, signed.snapshot_sha256, signed.signing_sha256)
        assert db.get(RequestApprovalSnapshot, approval_id).template_content == approval_bytes
        for id_, evidence in old_generated.items():
            row = db.get(GeneratedDocument, id_)
            assert (row.content, row.sha256) == evidence
    company = client.get('/api/studio/company', headers=auth_headers).json()
    changed = client.put('/api/studio/company', headers=auth_headers, json={
        'expected_revision': company['revision'], 'identity': {**company['identity'], 'legal_name': 'SYNTHETIC QA AFTER SIGNING'}})
    assert changed.status_code == 200, changed.text
    with session_factory() as db:
        signed = db.get(OfficialDocumentVersion, signed_version_id)
        assert (signed.docx_content, signed.pdf_content, signed.snapshot_sha256, signed.signing_sha256) == signed_bytes
    folders = {'TRANSFER_ISSUE': 'transfer-issue', 'TRANSFER_RETURN': 'transfer-return',
        'REPAIR_PROTOCOL': 'repair-protocol', 'PART_REQUEST': 'parts-request'}
    reports = []
    with session_factory() as db:
        assert original_ids <= {m.id for m in db.scalars(select(Machine))}
        documents = {}
        statement = select(OfficialDocumentVersion).join(DocumentTemplateVersion,
            DocumentTemplateVersion.id == OfficialDocumentVersion.template_version_id).where(
                DocumentTemplateVersion.layout_contract["studio"].as_integer() == 1).order_by(OfficialDocumentVersion.id)
        for version in db.scalars(statement):
            if manifest_from_source(version.docx_content) is not None:
                documents[db.get(OfficialDocument, version.document_id).document_type] = version
        assert set(documents) == set(folders)
        for workflow, version in documents.items():
            if workflow != 'REPAIR_PROTOCOL':
                role = {'TRANSFER_ISSUE': 'Предал оборудването', 'TRANSFER_RETURN': 'Върнал оборудването', 'PART_REQUEST': 'Заявител'}[workflow]
                signed_table = next(t for t in tables_in(Document(io.BytesIO(version.docx_content)))
                    if len(t.columns) == 3 and len(t.rows) == 2 and role in t.cell(0, 0).text)
                assert all(len(row.cells[2]._tc.xpath('.//w:drawing')) == 1 for row in signed_table.rows)
            folder = output / folders[workflow]
            folder.mkdir(parents=True, exist_ok=True)
            (folder / 'official-synthetic.docx').write_bytes(version.docx_content)
            (folder / 'official-synthetic.pdf').write_bytes(version.pdf_content)
            for stale in folder.glob('official-synthetic-page-*.png'):
                stale.unlink()
            with fitz.open(stream=version.pdf_content, filetype='pdf') as pdf:
                for i, page in enumerate(pdf, 1):
                    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(folder / f'official-synthetic-page-{i}.png')
                text = '\n'.join(p.get_text() for p in pdf)
                assert 'SYNTHETIC-QA-' in text and '[[ASSETCORE_' not in text
                assert all(label not in text for label in ('Фирмено лого', 'Company logo', 'Логотип компании'))
                assert 'Фирмено лого' not in Document(io.BytesIO(version.docx_content)).sections[0].header._element.xml
                if workflow == 'PART_REQUEST':
                    assert len(pdf) >= 3  # Actual approved diagram and parts-list appendices.
                reports.append({'workflow': workflow, 'docx_sha256': digest(version.docx_content), 'pdf_sha256': digest(version.pdf_content),
                    'pages': len(pdf), 'status': version.status, 'template_version_id': version.template_version_id,
                    'synthetic_equipment_and_people': True, 'original_inventory_ids_preserved': True,
                    'recorded_checklist_items': len(version.snapshot.get('equipment_checklist', [])) if workflow.startswith('TRANSFER_') else None,
                    'archive_unchanged_after_profile_and_selection_change': True,
                    'signed_kk_archive_unchanged_after_followup_profile_change': True if workflow == 'PART_REQUEST' else None,
                    'graphic_signatures_in_original_cells': True if workflow != 'REPAIR_PROTOCOL' else None,
                    'catalog_image_bytes_preserved_through_signing': True if workflow == 'PART_REQUEST' else None,
                    'studio_final_conversion_failure_rolled_back': True if workflow == 'PART_REQUEST' else None,
                    'approval_template_bytes_unchanged': True if workflow == 'PART_REQUEST' else None})
    (output / 'official-workflow-proof.json').write_text(json.dumps({'passed': True, 'cases': reports}, indent=2) + '\n', encoding='utf-8')
    from studio_proposal_archive_regression import verify_proposal_archive_preservation
    verify_proposal_archive_preservation(client, auth_headers, session_factory, output)
