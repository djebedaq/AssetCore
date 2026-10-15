"""REQUESTS-01 scenarios use disposable databases and explicitly labelled QA content."""

import hashlib
from decimal import Decimal

import fitz
import pytest
from app.models import GeneratedDocument, PartRequest, RequestNumberCounter
from app.part_requests.request_schemas import FreeFormLine
from sqlalchemy import select


def payload(**changes):
    return {
        "department": "QA отдел",
        "reason": "QA свободна заявка",
        "language": "bg",
        "lines": [
            {"description": "QA делим материал", "quantity": "2.5", "unit": "l"},
            {"description": "QA неделим артикул", "quantity": "20", "unit": "pcs"},
        ],
        **changes,
    }


def create(client, headers, **changes):
    response = client.post("/api/part-requests/free-form", headers=headers, json=payload(**changes))
    assert response.status_code == 201
    return response.json()


def approve(client, headers, request):
    response = client.post(
        f"/api/part-requests/{request['id']}/decision",
        headers=headers,
        json={"decision": "APPROVED"},
    )
    assert response.status_code == 200
    return response.json()


def test_free_form_edit_preview_approval_archive(client, auth_headers, session_factory):
    item = create(client, auth_headers)
    assert item["machine_id"] is None and item["official_number"] is None
    assert item["request_kind"] == "FREE_FORM"
    changed = payload(
        lines=[{"description": "QA редактиран ред", "quantity": "0,75", "unit": "kg"}]
    )
    edited = client.put(
        f"/api/part-requests/{item['id']}/free-form", headers=auth_headers, json=changed
    )
    assert edited.status_code == 200
    assert len(edited.json()["lines"]) == 1
    assert Decimal(edited.json()["lines"][0]["exact_quantity"]) == Decimal("0.75")
    preview = client.get(f"/api/part-requests/{item['id']}/preview", headers=auth_headers)
    assert preview.status_code == 200
    with fitz.open(stream=preview.content, filetype="pdf") as pdf:
        text = "".join(page.get_text() for page in pdf)
        assert "ПРОЕКТ" in text and "KK1000" not in text and "PART No" not in text
    client.post(f"/api/part-requests/{item['id']}/submit", headers=auth_headers)
    approved = approve(client, auth_headers, item)
    assert approved["official_number"] == "KK1000"
    assert approved["generation_status"] == "READY", approved.get("generation_error")
    assert {d["filename"] for d in approved["documents"]} == {"KK1000.pdf", "KK1000.docx"}
    again = approve(client, auth_headers, item)
    assert again["official_number"] == "KK1000"
    blocked = client.put(
        f"/api/part-requests/{item['id']}/free-form", headers=auth_headers, json=changed
    )
    assert blocked.status_code == 409
    with session_factory() as db:
        docs = list(db.scalars(select(GeneratedDocument)))
        assert len(docs) == 2
        for doc in docs:
            first = client.get(f"/api/generated-documents/{doc.id}/download", headers=auth_headers)
            second = client.get(f"/api/generated-documents/{doc.id}/download", headers=auth_headers)
            assert first.status_code == second.status_code == 200
            assert first.content == second.content == doc.content
            assert hashlib.sha256(first.content).hexdigest() == doc.sha256


@pytest.mark.parametrize(
    "amount,unit,valid",
    [
        ("1", "pcs", True),
        ("2.5", "l", True),
        ("0.75", "kg", True),
        ("7.5", "m", True),
        ("1.2", "pcs", False),
        ("0", "l", False),
        ("-1", "kg", False),
        ("NaN", "m", False),
        ("Infinity", "m", False),
        ("bad", "m", False),
        (True, "m", False),
        ("0.0000001", "m", False),
    ],
)
def test_quantities(amount, unit, valid):
    data = {"description": "QA", "quantity": amount, "unit": unit}
    if valid:
        assert FreeFormLine.model_validate(data).quantity > 0
    else:
        with pytest.raises(ValueError):
            FreeFormLine.model_validate(data)


def test_preview_watermark_every_page_and_no_issue(client, auth_headers, session_factory):
    item = create(
        client,
        auth_headers,
        lines=[
            {"description": f"QA дълъг ред {i} " * 10, "quantity": "7.5", "unit": "m"}
            for i in range(65)
        ],
    )
    response = client.get(f"/api/part-requests/{item['id']}/preview", headers=auth_headers)
    assert response.status_code == 200
    with fitz.open(stream=response.content, filetype="pdf") as pdf:
        assert len(pdf) > 1
        assert all("ПРОЕКТ" in page.get_text() for page in pdf)
    with session_factory() as db:
        assert db.get(RequestNumberCounter, 1) is None
        assert not list(db.scalars(select(GeneratedDocument)))
        assert db.get(PartRequest, item["id"]).status == "DRAFT"


def test_delivery_correction_cancellation(client, auth_headers):
    item = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    original_files = {
        record["download_endpoint"]: client.get(
            f"/api{record['download_endpoint']}", headers=auth_headers
        ).content
        for record in item["documents"]
    }
    line = item["lines"][0]
    url = f"/api/part-requests/{item['id']}"
    receipt = {
        "idempotency_key": "qa-receipt",
        "lines": [{"line_id": line["id"], "quantity": "0.1"}],
    }
    first = client.post(url + "/deliveries", headers=auth_headers, json=receipt)
    assert first.status_code == 201
    second = client.post(url + "/deliveries", headers=auth_headers, json=receipt)
    assert len(second.json()["delivery_events"]) == 1
    assert second.json()["lines"][0]["received_quantity"] == "0.1"
    assert second.json()["lines"][0]["remaining_quantity"] == "2.4"
    too_much = client.post(
        url + "/deliveries",
        headers=auth_headers,
        json={
            **receipt,
            "idempotency_key": "qa-over",
            "lines": [{"line_id": line["id"], "quantity": "3"}],
        },
    )
    assert too_much.status_code == 409
    event_id = first.json()["delivery_events"][0]["id"]
    reversed_ = client.post(
        url + f"/deliveries/{event_id}/reverse",
        headers=auth_headers,
        json={"idempotency_key": "qa-reverse", "reason": "QA корекция"},
    )
    assert reversed_.status_code == 201
    assert Decimal(reversed_.json()["lines"][0]["received_quantity"]) == 0
    again = client.post(
        url + "/deliveries", headers=auth_headers, json={**receipt, "idempotency_key": "qa-second"}
    )
    assert again.status_code == 201
    no_confirmation = client.post(
        url + "/cancel", headers=auth_headers, json={"confirmed": False, "reason": "QA"}
    )
    assert no_confirmation.status_code == 422
    cancelled = client.post(
        url + "/cancel", headers=auth_headers, json={"confirmed": True, "reason": "QA причина"}
    )
    assert cancelled.status_code == 200
    body = cancelled.json()
    assert (
        body["official_number"] == item["official_number"]
        and body["approved_at"] == item["approved_at"]
    )
    assert body["status"] == "CANCELLED" and len(body["delivery_events"]) == 3
    assert body["lines"][0]["received_quantity"] == "0.1"
    for endpoint, content in original_files.items():
        response = client.get(f"/api{endpoint}", headers=auth_headers)
        assert response.status_code == 200 and response.content == content


@pytest.mark.parametrize(
    "statement",
    [
        "UPDATE generated_documents SET filename = 'QA overwritten' WHERE part_request_id = :id",
        "DELETE FROM generated_documents WHERE part_request_id = :id",
        "UPDATE request_delivery_events SET note = 'QA overwritten' WHERE request_id = :id",
        "DELETE FROM request_delivery_events WHERE request_id = :id",
    ],
)
def test_sqlite_archive_and_delivery_database_guards(
    client, auth_headers, session_factory, statement
):
    from sqlalchemy import text
    from sqlalchemy.exc import DBAPIError

    item = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    response = client.post(
        f"/api/part-requests/{item['id']}/deliveries",
        headers=auth_headers,
        json={
            "idempotency_key": "qa-evidence-guard",
            "lines": [{"line_id": item["lines"][0]["id"], "quantity": "0.75"}],
        },
    )
    assert response.status_code == 201
    with session_factory() as db, pytest.raises(DBAPIError):
        db.execute(text(statement), {"id": item["id"]})
        db.commit()
    with session_factory() as db:
        request = db.get(PartRequest, item["id"])
        assert request.delivery_events[0].quantities == {str(item["lines"][0]["id"]): "0.75"}
        documents = list(
            db.scalars(
                select(GeneratedDocument).where(GeneratedDocument.part_request_id == item["id"])
            )
        )
        assert len(documents) == 2
        for document in documents:
            assert document.filename in {"KK1000.pdf", "KK1000.docx"}
            assert hashlib.sha256(document.content).hexdigest() == document.sha256


def test_render_failure_preserves_issue_and_retry_snapshot(
    client, auth_headers, monkeypatch, session_factory
):
    from app.documents import request_issue_documents

    original = request_issue_documents.build_pdf
    monkeypatch.setattr(
        request_issue_documents,
        "build_pdf",
        lambda *_a, **_kw: (_ for _ in ()).throw(RuntimeError("QA failure")),
    )
    first = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    assert first["official_number"] == "KK1000" and first["generation_status"] == "FAILED"
    assert not first["documents"]
    with session_factory() as db:
        request = db.get(PartRequest, first["id"])
        before = request.approval_snapshot.sha256
        request.decided_by.full_name = "QA changed identity"
        db.commit()
    monkeypatch.setattr(request_issue_documents, "build_pdf", original)
    response = client.post(f"/api/part-requests/{first['id']}/documents", headers=auth_headers)
    assert response.status_code == 201
    assert response.json()["generation_status"] == "READY"
    with session_factory() as db:
        request = db.get(PartRequest, first["id"])
        assert request.approval_snapshot.sha256 == before
        pdf = db.scalar(select(GeneratedDocument).where(GeneratedDocument.format == "pdf"))
        with fitz.open(stream=pdf.content, filetype="pdf") as rendered:
            assert "QA changed identity" not in "".join(p.get_text() for p in rendered)
    second = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    assert second["official_number"] == "KK1001"


def test_return_reject_pending_and_historical_approval(client, auth_headers, session_factory):
    from app.models import PartRequestLine

    first = create(client, auth_headers, submit_for_approval=True)
    for decision in ("RETURNED_FOR_CHANGES", "REJECTED"):
        response = client.post(
            f"/api/part-requests/{first['id']}/decision",
            headers=auth_headers,
            json={"decision": decision},
        )
        assert response.status_code == 200
        assert response.json()["official_number"] is None
        if decision == "RETURNED_FOR_CHANGES":
            client.post(f"/api/part-requests/{first['id']}/submit", headers=auth_headers)
    with session_factory() as db:
        old = PartRequest(
            part_name="QA historical", request_reference="PR-QA-HISTORICAL", status="APPROVED"
        )
        pending = PartRequest(
            part_name="QA pending", request_reference="PR-QA-PENDING", status="WAITING_APPROVAL"
        )
        pending.lines.append(PartRequestLine(description="QA catalog legacy row", quantity=1))
        db.add_all([old, pending])
        db.commit()
        old_id, pending_id = old.id, pending.id
    approved = approve(client, auth_headers, {"id": pending_id})
    assert approved["official_number"] == "KK1000"
    with session_factory() as db:
        old = db.get(PartRequest, old_id)
        assert old.approval_snapshot is None and old.request_reference == "PR-QA-HISTORICAL"


def test_cancel_permissions_and_locked_attachment(client, auth_headers, session_factory):
    from test_part_request_canonical_workflow import _add_user, _login

    _add_user(session_factory, email="qa-requests-mechanic@assetcore.test", role="mechanic")
    mechanic = _login(client, "qa-requests-mechanic@assetcore.test")
    item = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    response = client.post(
        f"/api/part-requests/{item['id']}/cancel",
        headers=mechanic,
        json={"confirmed": True, "reason": "QA"},
    )
    assert response.status_code == 403
    response = client.patch(
        f"/api/part-requests/{item['id']}/fulfillment",
        headers=auth_headers,
        json={"status": "CANCELLED", "confirmed": "true", "note": "QA"},
    )
    assert response.status_code == 422
    response = client.patch(
        f"/api/part-requests/{item['id']}/fulfillment",
        headers=mechanic,
        json={"status": "CANCELLED", "confirmed": True, "note": "QA"},
    )
    assert response.status_code == 403
    response = client.post(
        f"/api/part-requests/{item['id']}/attachments",
        headers=auth_headers,
        json={"filename": "qa.pdf", "media_type": "application/pdf", "content_base64": "UUE="},
    )
    assert response.status_code in {409, 422}


@pytest.mark.parametrize("language", ["bg", "en", "ru"])
def test_documents_all_languages(client, auth_headers, language):
    import io

    from docx import Document

    item = approve(
        client,
        auth_headers,
        create(client, auth_headers, submit_for_approval=True, language=language),
    )
    assert item["generation_status"] == "READY"
    for record in item["documents"]:
        response = client.get(f"/api{record['download_endpoint']}", headers=auth_headers)
        assert response.status_code == 200
        if record["format"] == "docx":
            document = Document(io.BytesIO(response.content))
            assert len(document.tables[0].columns) == 3
            assert [cell.text for cell in document.tables[0].rows[0].cells] == [
                "Pos",
                "DESCRIPTION",
                "Q'TY ordered",
            ]
            assert "KK1000" in " ".join(p.text for p in document.paragraphs)


def test_all_request_kinds_share_the_sequence_and_search(client, auth_headers, session_factory):
    from app.models import ApprovalDecision, PartRequestLine, RepairKit, User
    from app.part_requests.service import decide_request, load_request

    first = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    assert first["official_number"] == "KK1000"
    with session_factory() as db:
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        kit = RepairKit(code="QA-REQUESTS01", name="QA request kit", created_by_id=actor.id)
        db.add(kit)
        db.flush()
        for index, kind in enumerate(("CATALOG", "UNKNOWN_PART", "REPAIR_KIT"), 1001):
            request = PartRequest(
                part_name="QA other request kind",
                status="WAITING_APPROVAL",
                requested_by_id=actor.id,
                request_reference=f"PR-QA-KIND-{index}",
                repair_kit_id=kit.id if kind == "REPAIR_KIT" else None,
            )
            request.lines.append(
                PartRequestLine(
                    description="QA other row",
                    quantity=1,
                    is_unknown_part=kind == "UNKNOWN_PART",
                    assembly="QA",
                )
            )
            db.add(request)
            db.flush()
            decide_request(
                db, load_request(db, request.id, lock=True), actor, ApprovalDecision.APPROVED, None
            )
            assert request.approval_snapshot.official_number == f"KK{index}"
            assert request.approval_snapshot.data["request_kind"] == kind
            if kind == "REPAIR_KIT":
                assert request.approval_snapshot.data["repair_kit_id"] == kit.id
            assert request.approval_snapshot.data["lines"][0]["reason"] is None
        db.commit()
    result = client.get(
        "/api/workspace/requests?q=KK1000&request_kind=FREE_FORM", headers=auth_headers
    )
    assert result.status_code == 200 and result.json()["total"] == 1
    assert result.json()["items"][0]["id"] == first["id"]


def test_kk_namespace_and_generic_document_correction_are_blocked(
    client, auth_headers, session_factory
):
    from app.models import OfficialDocument

    item = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    with session_factory() as db:
        document = db.scalar(
            select(OfficialDocument).where(
                OfficialDocument.document_number == item["official_number"]
            )
        )
        document_id = document.id
    blocked = client.post(
        f"/api/official-documents/{document_id}/supersede",
        headers=auth_headers,
        json={"reason": "QA forbidden modification", "snapshot": {"reason": "QA changed"}},
    )
    assert blocked.status_code == 409
    blocked = client.post(
        f"/api/official-documents/{document_id}/cancel",
        headers=auth_headers,
        json={"reason": "QA forbidden generic cancellation"},
    )
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["code"] == "approved_request_is_immutable"
    blocked = client.post(
        "/api/official-documents",
        headers=auth_headers,
        json={
            "document_number": "KK1001",
            "document_type": "PART_REQUEST",
            "language": "bg",
            "snapshot": {},
        },
    )
    assert blocked.status_code == 409


def test_approval_retains_header_bytes_across_render_retry(
    client, auth_headers, session_factory, monkeypatch, tmp_path
):
    import zipfile

    from app.documents import request_issue_documents
    from app.documents.part_request_documents import OFFICIAL_HEADER
    from PIL import Image

    original_header_hash = hashlib.sha256(OFFICIAL_HEADER.read_bytes()).hexdigest()
    renderer = request_issue_documents.build_pdf
    monkeypatch.setattr(
        request_issue_documents,
        "build_pdf",
        lambda *_a, **_kw: (_ for _ in ()).throw(RuntimeError("QA retry")),
    )
    item = approve(client, auth_headers, create(client, auth_headers, submit_for_approval=True))
    assert item["generation_status"] == "FAILED"
    changed = tmp_path / "qa-other-header.png"
    Image.new("RGB", (100, 20), "red").save(changed)
    changed_hash = hashlib.sha256(changed.read_bytes()).hexdigest()
    monkeypatch.setattr(request_issue_documents, "OFFICIAL_HEADER", changed)
    monkeypatch.setattr(request_issue_documents, "build_pdf", renderer)
    retry = client.post(f"/api/part-requests/{item['id']}/documents", headers=auth_headers)
    assert retry.status_code == 201 and retry.json()["generation_status"] == "READY"
    with session_factory() as db:
        record = db.scalar(select(GeneratedDocument).where(GeneratedDocument.format == "docx"))
        assert record.snapshot["header_image_sha256"] == original_header_hash
        import io

        with zipfile.ZipFile(io.BytesIO(record.content)) as archive:
            hashes = {
                hashlib.sha256(archive.read(name)).hexdigest()
                for name in archive.namelist()
                if name.startswith("word/media/")
            }
        assert original_header_hash in hashes and changed_hash not in hashes


def test_export_review_examples(client, auth_headers):
    """Export only from a disposable fixture, never a committed database."""
    import json
    import os
    from pathlib import Path

    target = os.environ.get("ASSETCORE_REQUESTS01_QA_OUTPUT")
    if not target:
        pytest.skip("Optional QA examples export is not requested.")
    root = Path(__file__).resolve().parents[1] / ".tmp"
    target = Path(target).resolve()
    assert target.is_relative_to(root.resolve())
    target.mkdir(parents=True, exist_ok=True)
    data = payload(
        department="QA — ПРИМЕР ЗА ПРЕГЛЕД",
        reason="QA пример от заданието; не е производствена заявка.",
        lines=[
            {"description": "Тефлонова лента", "quantity": "20", "unit": "pcs"},
            {"description": "WD-40", "quantity": "3", "unit": "pcs"},
            {"description": "Масло", "quantity": "2.5", "unit": "l"},
            {"description": "Грес", "quantity": "0.75", "unit": "kg"},
            {"description": "Маркуч", "quantity": "7.5", "unit": "m"},
        ],
    )
    draft = create(client, auth_headers, **data)
    response = client.get(f"/api/part-requests/{draft['id']}/preview", headers=auth_headers)
    assert response.status_code == 200
    (target / "request-preview.pdf").write_bytes(response.content)
    client.post(f"/api/part-requests/{draft['id']}/submit", headers=auth_headers)
    item = approve(client, auth_headers, draft)
    assert item["generation_status"] == "READY"
    (target / "free-form-request.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for record in item["documents"]:
        response = client.get(f"/api{record['download_endpoint']}", headers=auth_headers)
        assert response.status_code == 200
        (target / record["filename"]).write_bytes(response.content)
    for filename in ("request-preview.pdf", "KK1000.pdf"):
        with fitz.open(target / filename) as pdf:
            for index, page in enumerate(pdf, 1):
                page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(
                    target / f"{filename}-{index}.png"
                )
