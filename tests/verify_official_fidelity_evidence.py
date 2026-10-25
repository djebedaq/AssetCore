"""Verify committed evidence against actual bytes; never rewrite the baseline."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import subprocess
from pathlib import Path

import fitz
from app.document_studio.files import validate_docx
from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops

CASES = {"transfer-issue", "transfer-return", "repair-protocol", "kk-machine", "kk-free"}


def text_nodes(raw):
    doc = Document(io.BytesIO(raw))
    parts = [doc.element, doc.sections[0].header._element, doc.sections[0].footer._element]
    return [node.text or "" for part in parts for node in part.iter(qn("w:t"))]


def verify(root):
    reports = json.loads((root / "comparison-results.json").read_text(encoding="utf-8"))
    expected = {(case, lang, count) for case in CASES for lang in ("bg", "en", "ru") for count in (1, 5, 30)}
    assert len(reports) == len(expected)
    assert {(r["case"], r["language"], r["rows"]) for r in reports} == expected
    pages = 0
    for report in reports:
        folder = root / report["case"] / f"{report['language']}-{report['rows']}"
        assert json.loads((folder / "comparison.json").read_text()) == report
        raw = [(folder / f"{name}.docx").read_bytes() for name in ("legacy", "studio")]
        for content in raw:
            validate_docx(content)
        assert text_nodes(raw[0]) == text_nodes(raw[1])
        if report["case"].startswith("kk-"):
            index = 0 if report["case"] == "kk-free" else 1
            for content in raw:
                doc = Document(io.BytesIO(content))
                assert len(doc.tables[index].columns) == (3 if index == 0 else 4)
                assert len(doc.tables[index].rows) == report["rows"] + 1
        with fitz.open(folder / "legacy.pdf") as before, fitz.open(folder / "studio.pdf") as after:
            assert len(before) == len(after) == report["legacy_pages"] == report["studio_pages"]
            assert "\n".join(p.get_text() for p in before) == "\n".join(p.get_text() for p in after)
            assert report["pdf_text_equal"] and report["native_text_equal"]
            assert len(report["pages"]) == len(before)
            for i in range(len(before)):
                rendered = []
                for name, pdf in (("legacy", before), ("studio", after)):
                    pix = pdf[i].get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
                    actual = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    with Image.open(folder / f"{name}-page-{i + 1}.png") as saved:
                        assert actual.size == saved.size
                        assert not ImageChops.difference(actual, saved.convert("RGB")).getbbox()
                    rendered.append(actual)
                assert rendered[0].size == rendered[1].size
                assert not ImageChops.difference(*rendered).getbbox()
                assert report["pages"][i] == {"page": i + 1, "changed_pixel_fraction": 0.0}
                with Image.open(folder / f"comparison-page-{i + 1}.png") as comparison:
                    assert comparison.size == (rendered[0].width * 2, rendered[0].height)
                    for side in range(2):
                        crop = comparison.crop((side * rendered[0].width, 0, (side + 1) * rendered[0].width, rendered[0].height))
                        assert not ImageChops.difference(crop.convert("RGB"), rendered[side]).getbbox()
                pages += 1
    preservation = json.loads((root / "preservation.json").read_text(encoding="utf-8"))
    # The report binds each baseline path to its actual Git blob. This also
    # works in the deliberately shallow CI checkout without fetching history.
    tracked = {}
    for record in subprocess.check_output(["git", "ls-tree", "-rz", "HEAD"]).split(b"\0"):
        if record:
            metadata, path = record.split(b"\t", 1)
            tracked[path.decode("utf-8")] = metadata.split()[2].decode()
    for name, evidence in preservation["files"].items():
        assert tracked.get(name) == evidence["baseline_git_blob"]
    changed = set(subprocess.check_output(["git", "diff", "--name-only"], text=True).splitlines())
    assert not changed.intersection(preservation["files"])
    assert not preservation["modified_protected_paths"]
    # Binary resources have no platform line-ending transform.
    for name, evidence in preservation["files"].items():
        if Path(name).suffix.lower() in {".docx", ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif", ".ico", ".xlsx", ".xls", ".zip"}:
            assert hashlib.sha256(Path(name).read_bytes()).hexdigest() == evidence["baseline_sha256"]
    print(f"Verified {len(reports)} actual DOCX/PDF pairs, {pages} page pairs and {len(preservation['files'])} protected paths")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path("docs/qa/official-document-fidelity"))
    verify(parser.parse_args().root)
