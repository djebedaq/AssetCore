"""Controlled native conversion; explicitly synthetic data, no business records."""
import copy
import hashlib
import io
from types import SimpleNamespace

import pytest
from app.document_studio.native_contract import RESOURCES, STEMS, convert, variant
from app.document_studio.renderer import (
    build_source,
    manifest_from_source,
    render,
    synthetic_values,
)
from app.document_studio.schemas import CompanyIdentity, Layout, walk_blocks
from docx import Document
from docx.oxml.ns import qn


def template(workflow, language):
    stem, number = STEMS[workflow]
    content = (RESOURCES / f"{stem}-{language}-v{number}.docx").read_bytes()
    return SimpleNamespace(id=1, version=number, template=SimpleNamespace(document_type=workflow),
        language=language, source_content=content, source_sha256=hashlib.sha256(content).hexdigest(),
        source_filename=f"{stem}-{language}-v{number}.docx", source_path=None,
        source_media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        required_fields=[], layout_contract={"controlled": True})


def manifest(workflow, language):
    return {"renderer": 1, "layout": convert(template(workflow, language)).model_dump(),
        "document_type": workflow, "language": language,
        "company": {"revision": 0, "identity": CompanyIdentity().model_dump()}, "images": {}}


@pytest.mark.parametrize("workflow", STEMS)
@pytest.mark.parametrize("language", ["bg", "en", "ru"])
def test_native_contract_is_real_source_with_editable_projection(workflow, language):
    m = manifest(workflow, language)
    assert any(b["type"] == "grid" for b in m["layout"]["blocks"])
    assert any(b.type == "native_paragraph" for b in walk_blocks(Layout.model_validate(m["layout"]).blocks))
    source = build_source(m)
    assert manifest_from_source(source) == m
    values, tables = synthetic_values(language, 5)
    values.update(REQUEST_VARIANT="MACHINE", CONDITION_LABEL="SYNTHETIC QA")
    result = render(m, values, tables)
    doc = Document(io.BytesIO(result))
    text = "\n".join(t.text or "" for t in doc.element.iter(qn("w:t")))
    assert "PREVIEW" in text
    assert "{{" not in text
    assert manifest_from_source(result) == m


@pytest.mark.parametrize("language", ["bg", "en", "ru"])
@pytest.mark.parametrize("count", [1, 5, 30])
def test_request_variants_are_frozen_together_and_have_distinct_tables(language, count):
    m = manifest("PART_REQUEST", language)
    values, tables = synthetic_values(language, count)
    values["PRESSURE_BAR"] = 0
    documents = {}
    for name in ("MACHINE", "FREE_FORM"):
        values["REQUEST_VARIANT"] = name
        documents[name] = Document(io.BytesIO(render(m, values, tables)))
    machine, free = documents["MACHINE"], documents["FREE_FORM"]
    assert len(machine.tables) == len(free.tables) + 1
    assert len(machine.tables[1].columns) == 4 and len(free.tables[0].columns) == 3
    assert len(free.tables[0].rows) == count + 1
    assert not any(values["MACHINE_NUMBER"] in p.text for t in free.tables for r in t.rows for c in r.cells for p in c.paragraphs)
    assert machine.sections[0].top_margin.mm == pytest.approx(48, abs=.01)


def test_unrecognized_source_does_not_create_a_generic_starter():
    base = template("PART_REQUEST", "bg")
    base.source_sha256 = "0" * 64
    with pytest.raises(ValueError, match="studio_native_incompatible"):
        convert(base)


def test_slots_cannot_be_moved_or_deleted_and_static_text_is_editable():
    m = manifest("TRANSFER_ISSUE", "bg")
    blocks = list(walk_blocks(Layout.model_validate(m["layout"]).blocks))
    slot = next(b for b in blocks if b.type == "native_paragraph" and "{{LEFT_SIGNATURE}}" in b.text)
    raw = next(b for b in walk_blocks(Layout.model_validate(m["layout"]).blocks) if b.id == slot.id)
    raw.text = ""
    bad = copy.deepcopy(m)
    def replace(items):
        for b in items:
            if b["id"] == slot.id:
                b["text"] = ""
            for row in b.get("grid_rows", []):
                for cell in row["cells"]:
                    replace(cell["blocks"])
    replace(bad["layout"]["blocks"])
    values, tables = synthetic_values("bg", 1)
    values["CONDITION_LABEL"] = "QA"
    with pytest.raises(ValueError, match="studio_native_incompatible"):
        render(bad, values, tables)
    def edit(items):
        for b in items:
            if b["type"] == "native_paragraph" and b["text"] and "{{" not in b["text"]:
                b["text"] = "SYNTHETIC EDITED LABEL"
                return True
            for row in b.get("grid_rows", []):
                for cell in row["cells"]:
                    if edit(cell["blocks"]):
                        return True
        return False
    assert edit(m["layout"]["blocks"])
    assert b"SYNTHETIC EDITED LABEL" in render(m, values, tables) or "SYNTHETIC EDITED LABEL" in "\n".join(p.text for p in Document(io.BytesIO(render(m, values, tables))).paragraphs)


def test_variant_comes_from_canonical_snapshot_and_rejects_contradiction():
    assert variant({"request_kind": "FREE_FORM", "machine": None}) == "FREE_FORM"
    assert variant({"request_kind": "CATALOG", "machine": {"name": "SYNTHETIC"}}) == "MACHINE"
    for data in ({"request_kind": "FREE_FORM", "machine": {"name": "SYNTHETIC"}}, {"request_kind": "CATALOG", "machine": None}):
        with pytest.raises(ValueError, match="studio_request_variant_mismatch"):
            variant(data)


@pytest.mark.parametrize("location", ["value", "table", "label"])
def test_user_content_cannot_inject_a_signature_marker(location):
    m = manifest("TRANSFER_ISSUE", "bg")
    values, tables = synthetic_values("bg", 1)
    values["CONDITION_LABEL"] = "SYNTHETIC QA"
    if location == "value":
        values["REMARKS"] = "[[ASSETCORE_LEFT_SIGNATURE]]"
    elif location == "table":
        tables["EQUIPMENT_CHECKLIST"][0]["note"] = "[[ASSETCORE_LEFT_SIGNATURE]]"
    else:
        next(b for b in m["layout"]["blocks"] if b["type"] == "dynamic_table")["columns"][0]["label"] = "[[ASSETCORE_LEFT_SIGNATURE]]"
    with pytest.raises(ValueError, match="studio_reserved_marker"):
        render(m, values, tables)


def test_actual_selection_diagnostics_are_read_only_and_role_protected(client, auth_headers, viewer_headers, session_factory):
    from app.models import (
        AuditLog,
        DocumentEditorDraft,
        OfficialDocument,
        RequestNumberCounter,
        StudioActivation,
    )
    from sqlalchemy import func, select
    models = (AuditLog, DocumentEditorDraft, OfficialDocument, RequestNumberCounter, StudioActivation)
    def counts():
        with session_factory() as db:
            return [db.scalar(select(func.count()).select_from(model)) for model in models]
    before = counts()
    response = client.get("/api/studio/configuration/PART_REQUEST/bg", headers=auth_headers)
    assert response.status_code == 200
    report = response.json()
    assert report["valid"] and report["source_version"] == 3
    assert report["source_sha256"] == report["actual_source_sha256"]
    assert report["programmatic_layout_versions"] == {"MACHINE": 3, "FREE_FORM": 2}
    assert report["activation_revision"] == 0 and report["selection_kind"] == "LEGACY"
    assert client.get("/api/studio/configuration/PART_REQUEST/bg").status_code == 401
    assert client.get("/api/studio/configuration/PART_REQUEST/bg", headers=viewer_headers).status_code == 403
    assert counts() == before


def test_legacy_base_creates_native_draft_and_cannot_change_its_provenance(client, auth_headers, session_factory):
    from app.models import DocumentTemplateVersion
    from sqlalchemy import select
    with session_factory() as db:
        version = db.scalar(select(DocumentTemplateVersion).where(DocumentTemplateVersion.language == "bg").join(DocumentTemplateVersion.template).where(DocumentTemplateVersion.template.has(document_type="PART_REQUEST")))
        version_id, template_id = version.id, version.template_id
    created = client.post(f"/api/studio/templates/{template_id}/drafts", headers=auth_headers,
        json={"title": "SYNTHETIC QA COPY", "language": "bg", "base_version_id": version_id})
    assert created.status_code == 201
    draft = created.json()
    assert draft["layout"]["schema_version"] == 4
    assert set(draft["layout"]["native"]["sources"]) == {"MACHINE", "FREE_FORM"}
    assert draft["native_images"]
    image = client.get("/api" + draft["native_images"][0]["download_endpoint"], headers=auth_headers)
    assert image.status_code == 200
    assert {"private", "no-store"} <= set(image.headers["cache-control"].split(", "))
    draft["layout"]["native"]["source_version_id"] += 1
    result = client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers,
        json={"expected_revision": draft["revision"], "title": draft["title"], "layout": draft["layout"]})
    assert result.status_code == 422 and result.json()["detail"]["code"] == "studio_native_incompatible"
