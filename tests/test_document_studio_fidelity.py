"""Structured forms use synthetic inputs only; no production installation."""
import copy
import io
from types import SimpleNamespace

import pytest
from app.document_studio import renderer
from app.document_studio.schemas import (
    SUPPORTED_TYPES,
    CompanyIdentity,
    Layout,
    validate_publication_fields,
    walk_blocks,
)
from app.document_studio.starter_layouts import professional_layout
from app.documents.templates import _template_version
from app.documents.transfer_documents import CHECKLIST, CHECKLIST_CODES, studio_checklist_rows
from app.models import DocumentTemplate, DocumentTemplateVersion, StudioActivation, User
from app.signature_rendering import _all_paragraphs, _embed_signature_image
from docx import Document
from PIL import Image
from pydantic import ValidationError
from sqlalchemy import select


def fidelity_manifest(workflow, language='bg'):
    return {'renderer': 1, 'layout': professional_layout(workflow, language).model_dump(),
        'document_type': workflow, 'language': language,
        'company': {'revision': 0, 'identity': CompanyIdentity(legal_name='SYNTHETIC QA ONLY').model_dump()}, 'images': {}}


@pytest.mark.parametrize('size', [(240, 200), (200, 240), (480, 160)])
def test_starter_fits_authorized_logos_without_stretching(size):
    layout = professional_layout('TRANSFER_ISSUE', 'bg', {'primary_logo_id': 1}, {1: size})
    image = next(b for b in walk_blocks(layout.blocks) if b.type == 'logo')
    assert image.width_mm / image.height_mm == pytest.approx(size[0] / size[1])
    assert image.width_mm <= 22 and image.height_mm <= 20
    renderer.check_layout(layout, 'TRANSFER_ISSUE')


def test_request_grouping_preserves_frozen_assembly_values_and_original_position_order():
    from app.documents.request_issue_documents import _studio_rows
    manifest = fidelity_manifest('PART_REQUEST')
    block = next(b for b in manifest['layout']['blocks'] if b['type'] == 'dynamic_table')
    block['group_by'] = 'assembly'
    source = {'request_kind': 'CATALOG', 'language': 'bg', 'lines': [
        {'position': str(i), 'part_number': f'SYNTHETIC-PART-{i}', 'description': f'Synthetic row {i}',
         'quantity': '1', 'unit': 'pcs', 'assembly': group, 'note': None}
        for i, group in enumerate(['QA assembly A', 'QA assembly A', '', 'QA assembly B', 'QA assembly A'], 1)]}
    rows = _studio_rows(source)
    assert [r['assembly'] for r in rows] == [r['assembly'] for r in source['lines']]
    values, tables = renderer.synthetic_values('bg', 1)
    tables['REQUEST_LINES'] = rows
    document = Document(io.BytesIO(renderer.render(manifest, values, tables)))
    table = next(t for t in document.tables if t.rows[0].cells[0].text == block['columns'][0]['label'] and len(t.columns) == 4)
    labels = [r.cells[0].text for r in table.rows][1:]
    assert labels == ['QA assembly A', '1', '2', '3', 'QA assembly B', '4', 'QA assembly A', '5']
    assert len(table._tbl.xpath('.//w:gridSpan[@w:val="4"]')) == 3
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert renderer.manifest_from_source(renderer.render(manifest, values, tables)) == manifest


@pytest.mark.parametrize('source,group', [('PARTS_USED', 'assembly'), ('EQUIPMENT_CHECKLIST', 'assembly'), ('REQUEST_LINES', 'arbitrary_expression')])
def test_grouping_is_allowlisted_and_workflow_specific(source, group):
    from app.document_studio.schemas import DynamicTable
    with pytest.raises(ValidationError):
        DynamicTable.model_validate({'id': 'bad_group', 'type': 'dynamic_table', 'source': source, 'group_by': group,
            'columns': [{'key': 'description', 'label': 'Synthetic', 'width_percent': 100}]})


def test_studio_signing_cannot_replace_missing_protected_slots_with_an_annex():
    from app.signature_rendering import finalize_signed_files_from_rows
    from app.template_engine import TemplateValidationError
    manifest = fidelity_manifest('PART_REQUEST')
    values, tables = renderer.synthetic_values('bg', 1)
    source = renderer.render(manifest, values, tables)  # Preview underscores have no protected markers.
    version = SimpleNamespace(docx_content=source, pdf_content=None, language='bg', snapshot={})
    with pytest.raises(TemplateValidationError, match='studio_signature_markers_missing'):
        finalize_signed_files_from_rows(version, [(SimpleNamespace(slot_code='REQUESTED_BY', identity_snapshot={}), None, b'')])
    assert version.docx_content == source and version.pdf_content is None


@pytest.mark.parametrize('language', ['bg', 'en', 'ru'])
@pytest.mark.parametrize('workflow', SUPPORTED_TYPES)
def test_neutral_professional_forms_are_native_editable_and_complete(workflow, language):
    manifest = fidelity_manifest(workflow, language)
    layout = Layout.model_validate(manifest['layout'])
    renderer.check_layout(layout, workflow)
    validate_publication_fields(layout, workflow)
    values, tables = renderer.synthetic_values(language, 5)
    source = renderer.build_source(manifest)
    docx = renderer.render_source(source, values, tables)
    parsed = Document(io.BytesIO(docx))
    assert parsed.sections[0].header.tables[0].columns.__len__() == 3
    assert len(parsed.tables) >= 4
    text = '\n'.join(p.text for p in _all_paragraphs(parsed))
    assert values['MACHINE_NUMBER'] in text and values['SERIAL_NUMBER'] in text
    assert values['PREPARER_NAME'] in text
    assert not parsed.inline_shapes  # No company or certification images are bundled.
    assert renderer.manifest_from_source(docx) == manifest
    fields = {b.field for b in walk_blocks(layout.blocks) if b.type == 'field'}
    if workflow == 'REPAIR_PROTOCOL':
        assert {'SOURCE_RETURN_REFERENCE', 'REPAIR_START', 'REPAIR_END', 'REQUIRED_PARTS', 'REMOVED_PARTS', 'APPROVED_BY_JOB_TITLE', 'CONDITION_AFTER'} <= fields
        assert values['TEST_DETAILS'] in text
    elif workflow.startswith('TRANSFER_'):
        assert all(component in text for component in CHECKLIST[language])


@pytest.mark.parametrize('operation', ['issue', 'return'])
@pytest.mark.parametrize('language', ['bg', 'en', 'ru'])
def test_checklist_uses_exact_operation_snapshot_and_canonical_order(operation, language):
    issued = [{'code': code, 'condition': 'GOOD', 'note': '', 'length_m': 30 if code == 'cable' else None} for code in CHECKLIST_CODES]
    returned = [{**item, 'condition': 'FAULTY', 'note': 'Синтетична дълга бележка ' * 50 if i == 4 else ''} for i, item in enumerate(issued)]
    transfer = SimpleNamespace(issue_checklist=issued, return_checklist=returned)
    rows = studio_checklist_rows(transfer, operation, language)
    assert len(rows) == 10
    assert [r['component'] for r in rows] == CHECKLIST[language]
    assert [r['position'] for r in rows] == list(map(str, range(1, 11)))
    expected = {'bg': ('Добро', 'Неизправно'), 'en': ('Good', 'Faulty'), 'ru': ('Хорошее', 'Неисправно')}[language][operation == 'return']
    assert all(r['condition'] == expected for r in rows)
    assert rows[6]['measurement'] == '30 m'
    assert rows[4]['note'] == (returned[4]['note'] if operation == 'return' else '')
    frozen = copy.deepcopy(rows)
    (returned if operation == 'issue' else issued)[0]['condition'] = 'MISSING'
    assert studio_checklist_rows(transfer, operation, language) == frozen
    custom = SimpleNamespace(issue_checklist=[{'code': 'qa_custom', 'label': 'SYNTHETIC CATEGORY COMPONENT', 'condition': 'Recorded free text'}], return_checklist=[])
    assert studio_checklist_rows(custom, 'issue', language)[0]['component'] == 'SYNTHETIC CATEGORY COMPONENT'
    assert studio_checklist_rows(custom, 'return', language) == []
    assert studio_checklist_rows(SimpleNamespace(issue_checklist=None, return_checklist=None), operation, language) == []


@pytest.mark.parametrize('mutation', ['span', 'width', 'duplicate', 'depth', 'unknown', 'cycle', 'shape', 'signature', 'unauthorized_image'])
def test_nested_layout_validation_cannot_hide_unsafe_children(mutation):
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    layout = manifest['layout']
    grid = next(b for b in layout['blocks'] if b['type'] == 'grid' and b['region'] == 'body')
    if mutation == 'span':
        grid['grid_rows'][0]['cells'][0]['colspan'] = 8
    elif mutation == 'width':
        grid['widths'] = [1, 99]
    elif mutation == 'duplicate':
        grid['grid_rows'][0]['cells'][0]['blocks'].append(copy.deepcopy(grid['grid_rows'][0]['cells'][1]['blocks'][0]))
    elif mutation in {'depth', 'cycle'}:
        nested = {'id': 'deep', 'type': 'grid', 'widths': [100], 'grid_rows': [{'cells': [{'blocks': []}]}]}
        if mutation == 'cycle':
            nested['grid_rows'][0]['cells'][0]['blocks'].append(nested)
        else:
            for i in range(6):
                nested = {'id': f'deep_{i}', 'type': 'grid', 'widths': [100], 'grid_rows': [{'cells': [{'blocks': [nested]}]}]}
        layout['blocks'].append(nested)
    elif mutation == 'unknown':
        grid['grid_rows'][0]['cells'][0]['blocks'].append({'id': 'unsafe', 'type': 'field', 'field': '__import__("os")'})
    elif mutation == 'shape':
        grid['grid_rows'] = ['invalid']
    elif mutation == 'signature':
        def remove_signatures(blocks):
            removed = sum(b['type'] == 'signature' for b in blocks)
            blocks[:] = [b for b in blocks if b['type'] != 'signature']
            for b in blocks:
                for row in b.get('grid_rows', []):
                    for cell in row['cells']:
                        removed += remove_signatures(cell['blocks'])
            return removed
        assert remove_signatures(layout['blocks']) == 2
    else:
        grid['grid_rows'][0]['cells'][0]['blocks'].append({'id': 'guessed', 'type': 'image', 'image_id': 999, 'width_mm': 10, 'height_mm': 10})
    with pytest.raises((ValueError, ValidationError), match='studio_|validation'):
        renderer.manifest_from_source(renderer.build_source(manifest))


def test_nested_protected_signature_cells_preserve_identity_when_images_are_embedded():
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    values, tables = renderer.synthetic_values('bg', 1)
    values.update(LEFT_SIGNATURE='[[ASSETCORE_LEFT_SIGNATURE]]', RIGHT_SIGNATURE='[[ASSETCORE_RIGHT_SIGNATURE]]',
        LEFT_SIGNER_NAME='СИНТЕТИЧНО ДЪЛГО ИМЕ ' * 8, RIGHT_SIGNER_JOB_TITLE='СИНТЕТИЧНА ДЛЪЖНОСТ ' * 8)
    unsigned = renderer.render(manifest, values, tables)
    doc = Document(io.BytesIO(unsigned))
    assert len(doc.inline_shapes) == 0
    stream = io.BytesIO()
    Image.new('RGB', (290, 90), '#17314f').save(stream, format='PNG')
    for side in ('LEFT', 'RIGHT'):
        assert _embed_signature_image(doc, f'[[ASSETCORE_{side}_SIGNATURE]]', stream.getvalue())
    text = '\n'.join(p.text for p in _all_paragraphs(doc))
    assert values['LEFT_SIGNER_NAME'] in text and values['RIGHT_SIGNER_JOB_TITLE'] in text
    assert '[[ASSETCORE_' not in text and len(doc.inline_shapes) == 2
    output = io.BytesIO()
    doc.save(output)
    assert renderer.manifest_from_source(output.getvalue()) == manifest


def test_publication_cannot_hide_signer_identity_inside_a_custom_signature_layout():
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    layout = Layout.model_validate(manifest['layout'])
    validate_publication_fields(layout, 'TRANSFER_ISSUE')
    def strip(blocks):
        for block in blocks:
            if block['type'] == 'grid':
                for row in block['grid_rows']:
                    for cell in row['cells']:
                        cell['blocks'] = [b for b in cell['blocks'] if b.get('field') != 'LEFT_SIGNER_NAME']
                        strip(cell['blocks'])
    strip(manifest['layout']['blocks'])
    with pytest.raises(ValueError, match='studio_required_block'):
        validate_publication_fields(Layout.model_validate(manifest['layout']), 'TRANSFER_ISSUE')


def test_cell_shading_is_inherited_by_default_and_explicit_white_remains_available():
    from docx.oxml.ns import qn
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    cell = manifest['layout']['blocks'][2]['grid_rows'][0]['cells'][0]
    cell['style']['background'] = '#ddeeff'
    values, tables = renderer.synthetic_values('bg', 1)
    parsed = Document(io.BytesIO(renderer.render(manifest, values, tables)))
    actual = parsed.tables[0].cell(0, 0)
    assert actual._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == 'ddeeff'
    assert actual.paragraphs[0]._p.pPr.find(qn('w:shd')) is None
    cell['blocks'][0]['style']['background'] = '#ffffff'
    parsed = Document(io.BytesIO(renderer.render(manifest, values, tables)))
    assert parsed.tables[0].cell(0, 0).paragraphs[0]._p.pPr.find(qn('w:shd')).get(qn('w:fill')) == 'ffffff'


@pytest.mark.parametrize('source', ['static', 'field', 'table'])
def test_reserved_signature_markers_cannot_be_injected_into_document_content(source):
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    values, tables = renderer.synthetic_values('bg', 1)
    marker = '[[ASSETCORE_LEFT_SIGNATURE]]'
    if source == 'static':
        manifest['layout']['blocks'].insert(1, {'id': 'injected', 'type': 'text', 'text': marker})
    elif source == 'field':
        values['REMARKS'] = marker
    else:
        tables['EQUIPMENT_CHECKLIST'][0]['note'] = marker
    with pytest.raises(ValueError, match='studio_reserved_marker'):
        renderer.render(manifest, values, tables)


def test_company_image_authorization_covers_nested_cells(client, auth_headers):
    from test_document_studio import image_payload
    from test_document_studio_publication import make_draft
    image = client.post('/api/studio/images', headers=auth_headers, json=image_payload()).json()
    draft = make_draft(client, auth_headers)
    draft['layout']['blocks'][0]['grid_rows'][0]['cells'][0]['blocks'].append({
        'id': 'nested_guessed_image', 'type': 'image', 'image_id': image['id'], 'width_mm': 10, 'height_mm': 10})
    response = client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers, json={
        'title': draft['title'], 'expected_revision': draft['revision'], 'layout': draft['layout']})
    assert response.status_code == 422 and response.json()['detail']['code'] == 'studio_image_not_allowed'


def test_missing_repair_labor_is_never_replaced_with_elapsed_calendar_time(
    client, auth_headers, machine_ids, session_factory, monkeypatch,
):
    from datetime import datetime

    from app.documents import repair_documents
    from app.models import Repair
    from test_repair_workflow_v14 import _create_repair

    id_ = _create_repair(client, auth_headers, machine_ids['8'])
    captured = {}
    class Captured(Exception):
        pass
    def capture(_template, values, _tables):
        captured.update(values)
        raise Captured
    monkeypatch.setattr(repair_documents, '_template_version', lambda *a: SimpleNamespace(layout_contract={'studio': 1}))
    monkeypatch.setattr(repair_documents, 'render_docx', capture)
    with session_factory() as db:
        repair = db.get(Repair, id_)
        repair.started_at, repair.closed_at = datetime(2026, 10, 19, 8), datetime(2026, 10, 19, 12)
        repair.diagnosis_minutes = repair.repair_minutes = repair.testing_minutes = None
        repair.participants.clear()
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        with pytest.raises(Captured):
            repair_documents.make_repair_documents(db, repair, actor.id)
    assert captured['REPAIR_START'] != captured['REPAIR_END']
    assert captured['TOTAL_WORK_DURATION'] == '' and captured['PARTICIPANT_TOTAL_DURATION'] == ''


def test_real_nested_merged_cells_and_explicit_grid_row_page_break(tmp_path):
    import os
    import shutil

    import fitz

    if os.name == 'nt' or not shutil.which('libreoffice'):
        pytest.skip('Real bounded conversion runs on Linux/Docker and CI')
    manifest = fidelity_manifest('TRANSFER_ISSUE')
    manifest['layout']['blocks'].append({'id': 'page_probe', 'type': 'grid', 'widths': [50, 50], 'grid_rows': [
        {'cells': [{'colspan': 2, 'blocks': [{'id': 'before_break', 'type': 'text', 'text': 'SYNTHETIC BEFORE ROW BREAK'}]}]},
        {'page_break_before': True, 'cells': [{'blocks': [{'id': 'after_break', 'type': 'text', 'text': 'SYNTHETIC AFTER ROW BREAK'}]},
            {'blocks': [{'id': 'nested_probe', 'type': 'grid', 'widths': [100], 'grid_rows': [{'cells': [{'blocks': [{'id': 'nested_field', 'type': 'field', 'field': 'MACHINE_NUMBER'}]}]}]}]}]},
    ]})
    values, tables = renderer.synthetic_values('bg', 1)
    docx = renderer.render(manifest, values, tables)
    assert b'gridSpan' in Document(io.BytesIO(docx)).element.xml.encode()
    pdf = renderer.pdf_required(docx)
    with fitz.open(stream=pdf, filetype='pdf') as document:
        text = [p.get_text() for p in document]
        before = next(i for i, t in enumerate(text) if 'SYNTHETIC BEFORE ROW BREAK' in t)
        after = next(i for i, t in enumerate(text) if 'SYNTHETIC AFTER ROW BREAK' in t)
        assert after > before
    (tmp_path / 'nested-page-break.docx').write_bytes(docx)
    (tmp_path / 'nested-page-break.pdf').write_bytes(pdf)


@pytest.mark.parametrize('language', ['bg', 'en', 'ru'])
def test_studio_kk_signer_uses_approved_identity_without_live_substitution(language):
    from app.signature_rendering import validate_studio_request_signer
    from fastapi import HTTPException
    data = {'approval_snapshot_sha256': 'synthetic-only', 'requested_by_id': 123,
        'requested_by_name': 'SYNTHETIC REQUESTER', 'requested_by_job_title': 'SYNTHETIC ROLE'}
    version = SimpleNamespace(snapshot=data, language=language)
    manifest = fidelity_manifest('PART_REQUEST', language)
    identity = {'user_id': 123, 'display_name': 'SYNTHETIC REQUESTER', 'job_title': 'SYNTHETIC ROLE'}
    validate_studio_request_signer(version, 'REQUESTED_BY', identity, manifest)
    for change in ({'user_id': 124}, {'display_name': 'SYNTHETIC LATER NAME'}, {'job_title': 'SYNTHETIC LATER ROLE'}):
        with pytest.raises(HTTPException) as result:
            validate_studio_request_signer(version, 'REQUESTED_BY', {**identity, **change}, manifest)
        assert result.value.status_code == 409
        assert result.value.detail['code'] == 'studio_signature_identity_mismatch'
        assert result.value.detail['message'] != 'studio.signature_identity_mismatch'
    assert data['requested_by_name'] == 'SYNTHETIC REQUESTER'


@pytest.mark.parametrize('workflow,slots', [
    ('TRANSFER_ISSUE', ('HANDOVER', 'ACCEPTANCE')),
    ('TRANSFER_RETURN', ('RETURNED_BY', 'ACCEPTED_RETURN')),
    ('PART_REQUEST', ('REQUESTED_BY', 'APPROVED_BY')),
])
def test_real_graphic_signatures_fit_long_names_and_jobs_in_multipage_cells(workflow, slots, tmp_path):
    import json
    import os
    import shutil
    from pathlib import Path

    import fitz
    from app.documents.templates import _signature_status
    from app.signature_rendering import finalize_signed_files_from_rows

    if os.name == 'nt' or not shutil.which('libreoffice'):
        pytest.skip('Real bounded conversion runs on Linux/Docker and CI')
    manifest = fidelity_manifest(workflow)
    values, tables = renderer.synthetic_values('bg', 30)
    values.update(LEFT_SIGNATURE='[[ASSETCORE_LEFT_SIGNATURE]]', RIGHT_SIGNATURE='[[ASSETCORE_RIGHT_SIGNATURE]]',
        LEFT_SIGNER_NAME='СИНТЕТИЧНО ДЪЛГО ИМЕ ЗА ПРОВЕРКА ' * 6,
        RIGHT_SIGNER_NAME='СИНТЕТИЧНО ИМЕ НА ПРИЕМАЩИЯ ' * 6,
        LEFT_SIGNER_JOB_TITLE='СИНТЕТИЧНА ДЪЛЖНОСТ ЗА ПРОВЕРКА ' * 6,
        RIGHT_SIGNER_JOB_TITLE='СИНТЕТИЧНА ДЛЪЖНОСТ НА ПРИЕМАЩИЯ ' * 6,
        CONDITION_TEXT='Синтетично техническо състояние. ' * 40, SIGNATURE_STATUS=_signature_status('bg'))
    unsigned = renderer.render(manifest, values, tables)
    assert not Document(io.BytesIO(unsigned)).inline_shapes
    source_pdf = renderer.pdf_required(unsigned)
    version = SimpleNamespace(docx_content=unsigned, pdf_content=source_pdf, language='bg', snapshot={})
    if workflow == 'PART_REQUEST':
        version.snapshot = {'approval_snapshot_sha256': 'synthetic-only', 'requested_by_id': 1, 'approved_by_id': 2,
            'requested_by_name': values['LEFT_SIGNER_NAME'], 'requested_by_job_title': values['LEFT_SIGNER_JOB_TITLE'],
            'approved_by_name': values['RIGHT_SIGNER_NAME'], 'approved_by_job_title': values['RIGHT_SIGNER_JOB_TITLE']}
    image = io.BytesIO()
    Image.new('RGB', (290, 90), '#17314f').save(image, format='PNG')
    rows = [(SimpleNamespace(slot_code=slot, identity_snapshot={'user_id': i + 1,
        'display_name': values[f'{side}_SIGNER_NAME'], 'job_title': values[f'{side}_SIGNER_JOB_TITLE']}), None, image.getvalue())
        for i, (slot, side) in enumerate(zip(slots, ('LEFT', 'RIGHT'), strict=True))]
    finalize_signed_files_from_rows(version, rows)
    assert renderer.manifest_from_source(version.docx_content) == manifest
    assert len(Document(io.BytesIO(version.docx_content)).inline_shapes) == 2
    output = Path(os.environ.get('ASSETCORE_STUDIO_LONG_SIGNATURE_OUTPUT', str(tmp_path))) / {
        'TRANSFER_ISSUE': 'transfer-issue', 'TRANSFER_RETURN': 'transfer-return', 'PART_REQUEST': 'parts-request'}[workflow]
    output.mkdir(parents=True, exist_ok=True)
    with fitz.open(stream=version.pdf_content, filetype='pdf') as pdf:
        assert len(pdf) > 1
        placements = []
        for i, page in enumerate(pdf, 1):
            for picture in page.get_image_info():
                bounds = fitz.Rect(picture['bbox'])
                assert bounds.x0 >= 0 and bounds.y0 >= 0 and bounds.x1 <= page.rect.width and bounds.y1 <= page.rect.height
                for word in page.get_text('words'):
                    overlap = bounds & fitz.Rect(word[:4])
                    assert overlap.is_empty or overlap.height <= 1 or overlap.width <= 1
                placements.append({'page': i, 'bounds': list(bounds)})
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(output / f'long-signed-page-{i}.png')
        assert len(placements) == 2
        assert '[[ASSETCORE_' not in '\n'.join(p.get_text() for p in pdf)
    (output / 'long-signed.docx').write_bytes(version.docx_content)
    (output / 'long-signed.pdf').write_bytes(version.pdf_content)
    (output / 'long-signature-proof.json').write_text(json.dumps({'passed': True, 'source_manifest_unchanged': True,
        'unsigned_had_no_signature_images': True, 'signed_via_existing_finalizer': True, 'placements': placements,
        'docx_sha256': renderer.digest(version.docx_content), 'pdf_sha256': renderer.digest(version.pdf_content)}, indent=2) + '\n', encoding='utf-8')


def test_explicit_legacy_restore_has_cas_confirmation_permissions_and_append_only_history(client, auth_headers, viewer_headers, session_factory):
    with session_factory() as db:
        legacy = _template_version(db, 'PART_REQUEST', 'bg')
        legacy_id, legacy_hash = legacy.id, legacy.source_sha256
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        template = db.scalar(select(DocumentTemplate).where(DocumentTemplate.document_type == 'PART_REQUEST'))
        source = renderer.build_source(fidelity_manifest('PART_REQUEST'))
        version = DocumentTemplateVersion(template_id=template.id, version=300, language='bg', source_content=source,
            source_sha256=renderer.digest(source), source_filename='synthetic.docx', layout_contract={'studio': 1},
            validation_status='PASSED', validation_report={'valid': True}, is_published=True, created_by_id=actor.id)
        db.add(version)
        db.commit()
        version_id = version.id
    activation = {'expected_version_id': None, 'expected_revision': 0, 'reason': 'Synthetic QA activation'}
    assert client.post(f'/api/studio/versions/{version_id}/activate', headers=auth_headers, json=activation).status_code == 201

    endpoint = '/api/studio/selections/PART_REQUEST/bg/restore-legacy'
    payload = {'expected_version_id': version_id, 'expected_revision': 1, 'reason': 'Synthetic operational restore', 'confirmed': True}
    assert client.post(endpoint, headers=viewer_headers, json=payload).status_code == 403
    assert client.post(endpoint, headers=auth_headers, json={**payload, 'confirmed': False}).status_code == 422
    assert client.post(endpoint, headers=auth_headers, json={**payload, 'reason': ''}).status_code == 422
    with session_factory() as db:
        assert _template_version(db, 'PART_REQUEST', 'bg').id == version_id
    response = client.post(endpoint, headers=auth_headers, json=payload)
    assert response.status_code == 201, response.text
    assert response.json()['version_id'] == legacy_id
    assert client.post(endpoint, headers=auth_headers, json=payload).status_code == 409
    assert client.post(f'/api/studio/versions/{version_id}/activate', headers=auth_headers, json=activation).status_code == 409
    with session_factory() as db:
        assert _template_version(db, 'PART_REQUEST', 'bg').id == legacy_id
        assert db.get(DocumentTemplateVersion, legacy_id).source_sha256 == legacy_hash
        assert db.get(DocumentTemplateVersion, version_id).source_content == source
        history = list(db.scalars(select(StudioActivation).order_by(StudioActivation.revision)))
        assert [(s.revision, s.selection_kind, s.version_id) for s in history] == [(1, 'STUDIO', version_id), (2, 'LEGACY', legacy_id)]
    activation['expected_revision'] = 2
    assert client.post(f'/api/studio/versions/{version_id}/activate', headers=auth_headers, json=activation).status_code == 201


def test_real_legacy_studio_legacy_documents_and_approval_bytes_remain_immutable(
    client, auth_headers, session_factory, tmp_path,
):
    import json
    import os
    import shutil
    from pathlib import Path

    from app.models import GeneratedDocument, RequestApprovalSnapshot
    from test_document_studio_publication import activate, make_draft, publish

    if os.name == 'nt' or not shutil.which('libreoffice'):
        pytest.skip('Real bounded conversion is verified on Linux/Docker and CI')
    with session_factory() as db:
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        actor.first_name, actor.middle_name, actor.last_name = 'Синтетичен', 'Тестов', 'Подписващ'
        actor.full_name = 'Синтетичен Тестов Подписващ'
        actor.job_title, actor.department_id = 'SYNTHETIC QA approver', None
        db.commit()
    def issue():
        request = client.post('/api/part-requests/free-form', headers=auth_headers, json={
            'department': 'SYNTHETIC QA ONLY', 'reason': 'Synthetic legacy restore acceptance', 'language': 'bg',
            'lines': [{'description': 'Синтетична техническа позиция', 'quantity': '1', 'unit': 'pcs'}]})
        assert request.status_code == 201, request.text
        id_ = request.json()['id']
        assert client.post(f'/api/part-requests/{id_}/submit', headers=auth_headers, json={}).status_code == 200
        approved = client.post(f'/api/part-requests/{id_}/decision', headers=auth_headers, json={'decision': 'APPROVED', 'note': 'Synthetic QA approval'})
        assert approved.status_code == 200 and approved.json()['generation_status'] == 'READY', approved.text
        with session_factory() as db:
            snapshot = db.scalar(select(RequestApprovalSnapshot).where(RequestApprovalSnapshot.request_id == id_))
            docs = list(db.scalars(select(GeneratedDocument).where(GeneratedDocument.part_request_id == id_)))
            return {'request_id': id_, 'template_id': snapshot.template_version_id, 'template': snapshot.template_content,
                'approval_hash': snapshot.sha256, 'archives': {d.id: (d.format, d.content, d.sha256) for d in docs}}
    legacy = issue()
    draft = make_draft(client, auth_headers)
    version, _ = publish(client, auth_headers, draft)
    activate(client, auth_headers, version['id'])
    studio = issue()
    assert studio['template_id'] == version['id'] and legacy['template_id'] != version['id']
    result = client.post('/api/studio/selections/PART_REQUEST/bg/restore-legacy', headers=auth_headers, json={
        'expected_version_id': version['id'], 'expected_revision': 1, 'confirmed': True, 'reason': 'Synthetic safe operational restore'})
    assert result.status_code == 201, result.text
    restored = issue()
    assert restored['template_id'] == legacy['template_id'] and restored['template'] == legacy['template']
    reports = []
    out = Path(os.environ.get('ASSETCORE_STUDIO_RESTORE_QA_OUTPUT', str(tmp_path)))
    out.mkdir(parents=True, exist_ok=True)
    with session_factory() as db:
        for label, captured in [('legacy-before', legacy), ('studio', studio), ('legacy-restored', restored)]:
            snapshot = db.scalar(select(RequestApprovalSnapshot).where(RequestApprovalSnapshot.request_id == captured['request_id']))
            assert snapshot.template_content == captured['template'] and snapshot.sha256 == captured['approval_hash']
            for id_, (format_, content, sha) in captured['archives'].items():
                archived = db.get(GeneratedDocument, id_)
                assert archived.content == content and archived.sha256 == sha
                (out / f'{label}.{format_}').write_bytes(content)
            reports.append({'selection': label, 'template_id': captured['template_id'], 'approval_hash': captured['approval_hash'],
                'archives': {format_: sha for format_, content, sha in captured['archives'].values()}, 'unchanged': True})
    (out / 'restore-proof.json').write_text(json.dumps({'passed': True, 'transitions': reports}, indent=2), encoding='utf-8')
