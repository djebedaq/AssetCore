"""Real converter acceptance; synthetic data exist only in disposable test DBs."""
import copy
import io
import os
import shutil
from pathlib import Path

import fitz
import pytest
from app.document_studio import publication, renderer
from app.document_studio.schemas import CompanyIdentity, Layout, initial_layout, publication_fields
from app.documents.templates import _template_version
from app.models import (
    AuditLog,
    CompanyProfile,
    DocumentEditorDraft,
    DocumentTemplateVersion,
    GeneratedDocument,
    OfficialDocument,
    OfficialDocumentVersion,
    ProtocolDocument,
    RequestApprovalSnapshot,
    RequestNumberCounter,
    StudioActivation,
    User,
)
from app.seed import seed_database
from docx import Document
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError
from test_document_studio import image_payload


@pytest.fixture
def converter():
    if os.name == 'nt' or not (shutil.which('soffice') or shutil.which('libreoffice')):
        pytest.skip('Real bounded LibreOffice acceptance runs on Linux/Docker and CI')


def make_draft(client, headers, workflow='PART_REQUEST', language='bg'):
    templates = client.get('/api/studio/templates', headers=headers).json()
    template = next(t for t in templates if t['document_type'] == workflow)
    response = client.post(f"/api/studio/templates/{template['id']}/drafts", headers=headers,
        json={'title': 'Synthetic QA editor', 'language': language})
    assert response.status_code == 201, response.text
    return response.json()


def publish(client, headers, draft):
    preview = client.post(f"/api/studio/drafts/{draft['id']}/preview", headers=headers,
        json={'expected_revision': draft['revision'], 'rows': 30})
    assert preview.status_code == 200, preview.text
    result = client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=headers,
        json={'expected_revision': draft['revision'], 'receipt': preview.json()['receipt'], 'change_note': 'Synthetic acceptance'})
    assert result.status_code == 201, result.text
    return result.json(), preview.json()


def activate(client, headers, version_id, previous=None):
    result = client.post(f'/api/studio/versions/{version_id}/activate', headers=headers,
        json={'expected_version_id': previous, 'reason': 'Synthetic QA explicit activation'})
    assert result.status_code == 201, result.text


def manifest(workflow='PART_REQUEST', language='bg'):
    layout = initial_layout(workflow).model_dump()
    layout['blocks'] += [{'id': f.lower(), 'type': 'field', 'field': f} for f in publication_fields(workflow)]
    if workflow == 'PART_REQUEST':
        layout['blocks'].insert(-2, {'id': 'rows', 'type': 'dynamic_table', 'source': 'REQUEST_LINES',
            'columns': [{'key': 'description', 'label': 'Описание / Description / Описание', 'width_percent': 70},
                        {'key': 'quantity', 'label': 'Quantity', 'width_percent': 30}]})
    return {'renderer': 1, 'layout': layout, 'document_type': workflow, 'language': language,
            'company': {'revision': 0, 'identity': CompanyIdentity().model_dump()}, 'images': {}}


@pytest.mark.parametrize('language', ['bg', 'en', 'ru'])
@pytest.mark.parametrize('count', [1, 5, 30])
def test_real_editable_docx_to_pdf_language_tables(converter, language, count):
    data = manifest(language=language)
    data['layout']['blocks'][:0] = [{'id': 'heading', 'type': 'heading', 'text': 'Проверка / Verification / Проверка',
        'style': {'bold': True, 'italic': True, 'font': 'DejaVu Sans', 'color': '#334455', 'background': '#f0f4ff', 'border_pt': .5}}]
    source = renderer.build_source(data)
    assert renderer.build_source(data) == source
    docx, pdf, report = publication.validate_render(data, count)
    assert report['page_count'] >= (2 if count == 30 else 1)
    parsed = Document(io.BytesIO(docx))
    assert any('Проверка' in p.text for p in parsed.paragraphs)
    assert len(parsed.tables[0].rows) == count + 1
    assert not parsed.tables[0]._tbl.xpath('.//w:pBdr')
    assert renderer.manifest_from_source(docx) == renderer.manifest_from_source(source)
    # A normal DOCX edit/appendix/signature save retains the frozen manifest.
    edited = io.BytesIO()
    parsed.save(edited)
    assert renderer.manifest_from_source(edited.getvalue()) == data
    with fitz.open(stream=pdf, filetype='pdf') as document:
        actual = '\n'.join(p.get_text() for p in document)
        assert 'Проверка' in actual and 'PREVIEW' in actual
        assert all(p.get_text().strip() for p in document)
        assert actual.count('Описание / Description / Описание') >= len(document)


def test_guessed_uploaded_image_id_is_not_a_company_allowlist(client, auth_headers):
    draft = make_draft(client, auth_headers)
    image = client.post('/api/studio/images', headers=auth_headers, json=image_payload()).json()
    layout = copy.deepcopy(draft['layout'])
    layout['blocks'].append({'id': 'logo', 'type': 'logo', 'image_id': image['id']})
    result = client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers,
        json={'expected_revision': 1, 'title': draft['title'], 'layout': layout})
    assert result.status_code == 422 and result.json()['detail']['code'] == 'studio_image_not_allowed'
    profile = client.get('/api/studio/company', headers=auth_headers).json()
    profile['identity']['secondary_logo_id'] = image['id']
    assert client.put('/api/studio/company', headers=auth_headers,
        json={'expected_revision': 0, 'identity': profile['identity']}).status_code == 200
    assert client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers,
        json={'expected_revision': 1, 'title': draft['title'], 'layout': layout}).status_code == 200


def test_unavailable_converter_does_not_publish_or_activate(client, auth_headers, session_factory, monkeypatch):
    draft = make_draft(client, auth_headers)
    from app import template_engine
    monkeypatch.setattr(template_engine, 'convert_docx_to_pdf', lambda *args, **kwargs: None)
    with session_factory() as db:
        count = db.scalar(select(func.count(DocumentTemplateVersion.id)))
    response = client.post(f"/api/studio/drafts/{draft['id']}/preview", headers=auth_headers, json={'expected_revision': 1})
    assert response.status_code == 422 and response.json()['detail']['code'] == 'studio_converter_unavailable'
    response = client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=auth_headers,
        json={'expected_revision': 1, 'receipt': 'x' * 65, 'change_note': 'QA blocked'})
    assert response.status_code == 422
    with session_factory() as db:
        assert db.scalar(select(func.count(DocumentTemplateVersion.id))) == count
        assert db.scalar(select(func.count(StudioActivation.id))) == 0
        assert db.get(DocumentEditorDraft, draft['id']).revision == 1


@pytest.mark.parametrize('failure', ['converter', 'new_company'])
def test_publication_failure_rolls_back_every_write_after_a_valid_preview_receipt(
    client, auth_headers, session_factory, monkeypatch, failure,
):
    draft = make_draft(client, auth_headers)
    with session_factory() as db:
        actor = db.scalar(select(User).where(User.is_system_owner.is_(True)))
        item = db.get(DocumentEditorDraft, draft['id'])
        frozen = publication.snapshot(db, item, 'PART_REQUEST')
        token = publication.receipt(item, frozen, actor.id)
        before = db.scalar(select(func.count(DocumentTemplateVersion.id)))
        actor_id = actor.id
    def interrupted(manifest):
        if failure == 'converter':
            raise ValueError('studio_pdf_overflow')
        with session_factory() as db:
            db.add(CompanyProfile(id=1, revision=1, identity=CompanyIdentity(legal_name='Synthetic concurrent QA company').model_dump(), updated_by_id=actor_id))
            db.commit()
        # Transaction fixture only: real conversion is covered by converter tests.
        return renderer.build_source(manifest), {'valid': True, 'manifest_sha256': renderer.digest(renderer.canonical(manifest))}
    monkeypatch.setattr(publication, 'publication_source', interrupted)
    response = client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=auth_headers,
        json={'expected_revision': 1, 'receipt': token, 'change_note': 'Synthetic interrupted publication'})
    assert response.status_code == (422 if failure == 'converter' else 409), response.text
    with session_factory() as db:
        assert db.scalar(select(func.count(DocumentTemplateVersion.id))) == before
        assert db.scalar(select(func.count(StudioActivation.id))) == 0
        assert db.get(DocumentEditorDraft, draft['id']).revision == 1


@pytest.mark.parametrize('mutation', ['columns', 'quantity', 'remarks', 'header'])
def test_unproven_layout_bindings_and_overflow_are_rejected_without_conversion(mutation):
    data = manifest()
    layout = copy.deepcopy(data['layout'])
    table = next(b for b in layout['blocks'] if b['type'] == 'dynamic_table')
    if mutation == 'columns':
        table['columns'][1]['key'] = 'date'
    elif mutation == 'quantity':
        table['columns'] = [{'key': 'description', 'label': 'QA', 'width_percent': 100}]
    elif mutation == 'remarks':
        from app.document_studio.schemas import validate_workflow
        repair = initial_layout('REPAIR_PROTOCOL').model_dump()
        repair['blocks'].append({'id': 'unbound', 'type': 'field', 'field': 'REMARKS'})
        with pytest.raises(ValueError, match='studio_incompatible_field'):
            validate_workflow(Layout.model_validate(repair), 'REPAIR_PROTOCOL')
        return
    else:
        layout['blocks'].append({'id': 'header', 'type': 'field', 'field': 'COMPANY_NAME', 'region': 'header', 'style': {'size_pt': 8, 'spacing_mm': 0}})
        data['company']['identity']['legal_name'] = 'Дълго фирмено наименование ' * 9
    data['layout'] = layout
    with pytest.raises(ValueError, match='studio_'):
        publication.validate_render(data)


def test_preview_has_no_official_side_effects_publication_is_immutable_and_activation_explicit(
    converter, client, auth_headers, session_factory,
):
    draft = make_draft(client, auth_headers)
    with session_factory() as db:
        old = _template_version(db, 'PART_REQUEST', 'bg').id
        tables = (OfficialDocument, OfficialDocumentVersion, RequestApprovalSnapshot, RequestNumberCounter, GeneratedDocument, AuditLog)
        before = {t.__tablename__: db.scalar(select(func.count()).select_from(t)) for t in tables}
    response = client.post(f"/api/studio/drafts/{draft['id']}/preview", headers=auth_headers, json={'expected_revision': 1, 'rows': 30})
    assert response.status_code == 200, response.text
    with session_factory() as db:
        assert {t.__tablename__: db.scalar(select(func.count()).select_from(t)) for t in tables} == before
    version, preview = publish(client, auth_headers, draft)
    with session_factory() as db:
        assert _template_version(db, 'PART_REQUEST', 'bg').id == old
        published = db.get(DocumentTemplateVersion, version['id'])
        original = published.source_content
        assert published.validation_report['valid'] and len(published.validation_report['cases']) == 3
    activate(client, auth_headers, version['id'])
    with session_factory() as db:
        assert _template_version(db, 'PART_REQUEST', 'bg').id == version['id']
        # Prepare/restart must never unpublish or rewrite Studio evidence.
        seed_database(db)
        assert db.get(DocumentTemplateVersion, version['id']).source_content == original
        assert db.get(DocumentTemplateVersion, version['id']).is_published
    with session_factory() as db, pytest.raises(DBAPIError, match='immutable_studio_evidence'):
        db.execute(text('UPDATE document_template_versions SET change_note = :note WHERE id = :id'), {'id': version['id'], 'note': 'must fail'})
    with session_factory() as db, pytest.raises(DBAPIError, match='immutable_studio_evidence'):
        db.execute(text('DELETE FROM studio_activations'))
    # A replay cannot publish the same workspace revision a second time.
    replay = client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=auth_headers,
        json={'expected_revision': 1, 'receipt': preview['receipt'], 'change_note': 'replay'})
    assert replay.status_code == 409


def test_profile_changes_invalidate_receipt_but_never_change_published_bytes(converter, client, auth_headers):
    draft = make_draft(client, auth_headers)
    version, preview = publish(client, auth_headers, draft)
    original = client.get(f"/api/studio/versions/{version['id']}/source", headers=auth_headers).content
    profile = client.get('/api/studio/company', headers=auth_headers).json()
    profile['identity']['legal_name'] = 'Synthetic QA later company'
    client.put('/api/studio/company', headers=auth_headers, json={'expected_revision': 0, 'identity': profile['identity']})
    assert client.get(f"/api/studio/versions/{version['id']}/source", headers=auth_headers).content == original
    # Same-revision valid receipt is bound to the company snapshot too.
    draft = client.get(f"/api/studio/drafts/{draft['id']}", headers=auth_headers).json()
    token = client.post(f"/api/studio/drafts/{draft['id']}/preview", headers=auth_headers, json={'expected_revision': draft['revision']}).json()['receipt']
    profile['identity']['legal_name'] = 'Synthetic QA another revision'
    client.put('/api/studio/company', headers=auth_headers, json={'expected_revision': 1, 'identity': profile['identity']})
    response = client.post(f"/api/studio/drafts/{draft['id']}/publish", headers=auth_headers,
        json={'expected_revision': draft['revision'], 'receipt': token, 'change_note': 'must fail'})
    assert response.status_code == 422 and response.json()['detail']['code'] == 'studio_preview_required'


def test_real_canonical_issue_return_sign_repair_and_kk_use_selected_studio_sources(
    converter, client, auth_headers, machine_ids, issue_payload, finalize_signatures, session_factory, tmp_path,
):
    versions = {}
    for workflow in ('TRANSFER_ISSUE', 'TRANSFER_RETURN', 'REPAIR_PROTOCOL', 'PART_REQUEST'):
        draft = make_draft(client, auth_headers, workflow)
        layout = copy.deepcopy(draft['layout'])
        layout['blocks'][:0] = [
            {'id': 'qa_only', 'type': 'heading', 'text': 'SYNTHETIC QA ONLY / СИНТЕТИЧНА ПРОВЕРКА', 'style': {'size_pt': 14, 'bold': True}},
            {'id': 'frozen_company', 'type': 'field', 'field': 'COMPANY_NAME'},
        ]
        saved = client.put(f"/api/studio/drafts/{draft['id']}", headers=auth_headers,
            json={'expected_revision': draft['revision'], 'title': draft['title'], 'layout': layout})
        assert saved.status_code == 200, saved.text
        draft = saved.json()
        version, _ = publish(client, auth_headers, draft)
        activate(client, auth_headers, version['id'])
        versions[workflow] = version['id']
    issued = client.post('/api/transfers/bulk-issue', headers=auth_headers, json=issue_payload(machine_ids['4']))
    assert issued.status_code == 201, issued.text
    finalize_signatures(client, issued)
    item = issued.json()['transfers'][0]
    returned = client.post('/api/transfers/bulk-return', headers=auth_headers, json={'document_language': 'bg', 'items': [{
        'transfer_id': item['transfer_id'], 'machine_id': item['machine_id'], 'condition_text': 'QA condition',
        'result_text': 'QA inspected', 'notes': 'QA', 'next_status': 'READY'}]})
    assert returned.status_code == 200, returned.text
    finalize_signatures(client, returned)
    from test_repair_workflow_v14 import _advance_to_final_stage, _create_repair
    repair = _create_repair(client, auth_headers, machine_ids['8'])
    _advance_to_final_stage(client, auth_headers, repair)
    completed = client.patch(f'/api/repair-cases/{repair}', headers=auth_headers, json={
        'status': 'COMPLETED', 'test_passed': True, 'test_method': 'QA', 'test_details': 'QA test details', 'functional_test_result': 'QA test',
        'result': 'QA completed', 'condition_after': 'QA ready', 'testing_minutes': 20})
    assert completed.status_code == 200, completed.text
    from test_part_request_canonical_workflow import _create_approved_request
    _, _, request = _create_approved_request(client, auth_headers, machine_ids['9'])
    with session_factory() as db:
        approval = db.scalar(select(RequestApprovalSnapshot).where(RequestApprovalSnapshot.request_id == request['id']))
        assert approval.template_version_id == versions['PART_REQUEST']
        assert approval.template_content == db.get(DocumentTemplateVersion, versions['PART_REQUEST']).source_content
        generated = list(db.scalars(select(GeneratedDocument).where(GeneratedDocument.part_request_id == request['id'])))
        assert {d.format for d in generated} == {'pdf', 'docx'}
        assert all(d.template_version_id == versions['PART_REQUEST'] for d in generated)
        official = list(db.scalars(select(OfficialDocumentVersion).where(OfficialDocumentVersion.template_version_id.in_(versions.values()))))
        assert {v.template_version_id for v in official} >= set(versions.values())
        assert all(renderer.manifest_from_source(v.docx_content) for v in official)
        for v in official:
            renderer.inspect_pdf(v.pdf_content)
            with fitz.open(stream=v.pdf_content, filetype='pdf') as pdf:
                assert '[[ASSETCORE_' not in '\n'.join(p.get_text() for p in pdf)
        archived = {v.id: (v.docx_content, v.pdf_content, v.signing_sha256) for v in official}
        transfer_docs = list(db.scalars(select(ProtocolDocument)))
        assert transfer_docs
    # Changes to a profile and a new active template never rewrite old signed
    # protocols or approval snapshots. Existing archive API returns stored bytes.
    identity = CompanyIdentity(legal_name='Synthetic QA future company').model_dump()
    client.put('/api/studio/company', headers=auth_headers, json={'expected_revision': 0, 'identity': identity})
    another = make_draft(client, auth_headers)
    new, _ = publish(client, auth_headers, another)
    activate(client, auth_headers, new['id'], versions['PART_REQUEST'])
    with session_factory() as db:
        for id_, evidence in archived.items():
            v = db.get(OfficialDocumentVersion, id_)
            assert (v.docx_content, v.pdf_content, v.signing_sha256) == evidence
        approval = db.scalar(select(RequestApprovalSnapshot).where(RequestApprovalSnapshot.request_id == request['id']))
        assert approval.template_version_id == versions['PART_REQUEST']
        assert renderer.manifest_from_source(approval.template_content)['company']['identity']['legal_name'] == 'AssetCore'
        out = Path(os.environ.get('ASSETCORE_STUDIO_QA_OUTPUT', str(tmp_path)))
        out.mkdir(parents=True, exist_ok=True)
        for v in db.scalars(select(OfficialDocumentVersion).where(OfficialDocumentVersion.template_version_id.in_(versions.values()))):
            name = v.document.document_type.lower() if hasattr(v, 'document') else str(v.id)
            for extension in ('docx', 'pdf'):
                (out / f'{name}-{v.id}.{extension}').write_bytes(getattr(v, extension + '_content'))
