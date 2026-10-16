"""USER-PROFILE-01: disposable identities, official archives and signed workflows."""

import hashlib
import io
import os
from pathlib import Path

import fitz
import pytest
from app.models import DocumentParticipant, GeneratedDocument, OfficialDocumentVersion, User
from docx import Document
from sqlalchemy import select
from test_bulk_transfers import complete_signing
from test_requests01 import approve, create
from test_user_accounts import _create_payload, _login


@pytest.mark.parametrize("middle", [None, "", "   ", "  Middle  "])
def test_initial_protected_owner_accepts_optional_middle_without_rewriting_existing_data(
    tmp_path, monkeypatch, middle
):
    from app.database import Base
    from app.documents.templates import _preparer_values
    from app.models import Machine
    from app.seed import _seed_verified_registry
    from app.settings import settings
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session

    monkeypatch.setattr(settings, "owner_first_name", " QA ")
    monkeypatch.setattr(settings, "owner_middle_name", middle)
    monkeypatch.setattr(settings, "owner_last_name", " Bootstrap ")
    monkeypatch.setattr(settings, "owner_job_title", " QA operator ")
    engine = create_engine(f"sqlite:///{tmp_path / 'fresh-owner.db'}")
    try:
        with engine.begin() as connection:
            connection.exec_driver_sql("BEGIN")
            Base.metadata.create_all(connection)
        with Session(engine) as db:
            _seed_verified_registry(db)
            db.commit()
            owner = db.scalar(select(User).where(User.is_system_owner.is_(True)))
            expected = "QA Middle Bootstrap" if middle == "  Middle  " else "QA Bootstrap"
            assert owner.full_name == expected
            assert owner.profile_status == "PROFILE_COMPLETE"
            assert owner.role == "administrator" and owner.is_active
            assert not owner.legal_name_exception
            assert _preparer_values(db, owner.id) == {
                "PREPARER_NAME": expected, "PREPARER_JOB_TITLE": "QA operator"}
            inventory = db.execute(select(Machine.__table__).order_by(Machine.id)).all()
            assert len(inventory) == 19
            monkeypatch.setattr(settings, "owner_first_name", "Later configuration")
            monkeypatch.setattr(settings, "owner_middle_name", "Later middle")
            _seed_verified_registry(db)
            db.commit()
            assert owner.full_name == expected  # Never ordinarily edit the designated owner.
            assert db.execute(select(Machine.__table__).order_by(Machine.id)).all() == inventory
    finally:
        engine.dispose()


@pytest.mark.parametrize("middle", [None, "", "   ", "  Automation  ", "omitted"])
@pytest.mark.parametrize("role", ["administrator", "director", "mechanic", "observer"])
def test_create_login_and_edit_optional_middle(client, auth_headers, middle, role):
    payload = _create_payload("optional@example.invalid", role)
    payload.update(first_name="  QA  ", last_name=" Profile  ", middle_name=middle)
    if middle == "omitted":
        payload.pop("middle_name")
    created = client.post("/api/users", headers=auth_headers, json=payload)
    assert created.status_code == 201
    body = created.json()
    expected = "QA Automation Profile" if middle == "  Automation  " else "QA Profile"
    assert body["full_name"] == expected
    assert body["middle_name"] == ("Automation" if "Automation" in expected else None)
    assert body["profile_status"] == "PROFILE_COMPLETE"
    assert not body["legal_name_exception"]
    _, logged_in = _login(client, payload["email"], payload["temporary_password"])
    assert logged_in["user"]["profile_status"] == "PROFILE_COMPLETE"
    for value, name in [("  Other  ", "QA Other Profile"), ("", "QA Profile"),
                        ("Again", "QA Again Profile"), (None, "QA Profile")]:
        updated = client.patch(f"/api/users/{body['id']}", headers=auth_headers,
                               json={"middle_name": value})
        assert updated.status_code == 200
        assert updated.json()["full_name"] == name
        assert updated.json()["profile_status"] == "PROFILE_COMPLETE"
    unchanged = client.patch(f"/api/users/{body['id']}", headers=auth_headers,
                             json={"job_title": "QA revised job"})
    assert unchanged.status_code == 200 and unchanged.json()["full_name"] == "QA Profile"


@pytest.mark.parametrize("field,value", [("first_name", " "), ("last_name", None), ("job_title", " ")])
def test_optional_middle_does_not_relax_required_identity(client, auth_headers, field, value):
    payload = _create_payload("required@example.invalid", "mechanic")
    payload[field] = value
    assert client.post("/api/users", headers=auth_headers, json=payload).status_code == 422
    payload[field] = "Required"
    response = client.post("/api/users", headers=auth_headers, json=payload)
    assert response.status_code == 201
    assert client.patch(f"/api/users/{response.json()['id']}", headers=auth_headers,
                        json={field: value}).status_code == 422


@pytest.mark.parametrize("middle", [None, "", " ", "omitted"])
def test_self_profile_completion_without_legal_exception(client, session_factory, middle):
    from app.security import hash_password

    with session_factory() as db:
        db.add(User(email="completion@example.invalid", full_name="QA incomplete",
                    password_hash=hash_password("StrongPass123!"), role="mechanic",
                    profile_status="PROFILE_INCOMPLETE"))
        db.commit()
    headers, _ = _login(client, "completion@example.invalid")
    payload = {"first_name": " QA ", "middle_name": middle, "last_name": " Completed ",
               "job_title": "QA operator"}
    if middle == "omitted":
        payload.pop("middle_name")
    completed = client.put("/api/users/me/profile", headers=headers, json=payload)
    assert completed.status_code == 200
    assert completed.json()["full_name"] == "QA Completed"
    assert completed.json()["profile_status"] == "PROFILE_COMPLETE"
    assert completed.json()["legal_name_exception_approved_by_id"] is None
    assert client.get("/api/transfers/availability", headers=headers).status_code == 200


def test_editing_a_legacy_identity_does_not_reintroduce_double_spaces(
    client, auth_headers, session_factory
):
    created = client.post("/api/users", headers=auth_headers,
                          json=_create_payload("legacy-space@example.invalid", "mechanic"))
    assert created.status_code == 201
    with session_factory() as db:
        user = db.get(User, created.json()["id"])
        user.first_name, user.middle_name, user.last_name = " QA  Compound ", "  Middle  Name ", " Legacy "
        db.commit()
    edited = client.patch(f"/api/users/{created.json()['id']}", headers=auth_headers,
                          json={"job_title": "QA revised job"})
    assert edited.status_code == 200
    assert edited.json()["full_name"] == "QA Compound Middle Name Legacy"


def test_two_name_preparer_approval_documents_signing_and_immutable_archive(
    client, auth_headers, session_factory, issue_payload, machine_ids
):
    payload = _create_payload("preparer@example.invalid", "administrator")
    payload.update(first_name="QA", last_name="Preparer", middle_name=None)
    created = client.post("/api/users", headers=auth_headers, json=payload)
    assert created.status_code == 201
    headers, _ = _login(client, payload["email"], payload["temporary_password"])
    changed = client.post("/api/auth/change-password", headers=headers, json={
        "current_password": payload["temporary_password"], "new_password": "Permanent789!",
        "confirm_password": "Permanent789!",
    })
    assert changed.status_code == 200
    headers, _ = _login(client, payload["email"], "Permanent789!")
    item = create(client, headers)
    assert client.post(f"/api/part-requests/{item['id']}/submit", headers=headers).status_code == 200
    approved = approve(client, headers, item)
    assert approved["generation_status"] == "READY", approved.get("generation_error")
    with session_factory() as db:
        archived = list(db.scalars(select(GeneratedDocument).where(GeneratedDocument.part_request_id == item["id"])))
        hashes = {document.id: document.sha256 for document in archived}
        assert {document.format for document in archived} == {"pdf", "docx"}
        for document in archived:
            if document.format == "docx":
                parsed = Document(io.BytesIO(document.content))
                text = " ".join([p.text for p in parsed.paragraphs] + [c.text for t in parsed.tables for r in t.rows for c in r.cells])
            else:
                with fitz.open(stream=document.content, filetype="pdf") as pdf:
                    text = " ".join(page.get_text() for page in pdf)
            assert "QA Preparer" in text
            output = os.environ.get("ASSETCORE_PROFILE01_QA_OUTPUT")
            if output:
                path = Path(output)
                path.mkdir(parents=True, exist_ok=True)
                (path / f"two-name-preparer.{document.format}").write_bytes(document.content)
    issued = client.post("/api/transfers/bulk-issue", headers=headers, json=issue_payload(machine_ids["4"]))
    assert issued.status_code == 201
    complete_signing(client, issued)
    with session_factory() as db:
        snapshots = [p.identity_snapshot for p in db.scalars(select(DocumentParticipant).where(
            DocumentParticipant.user_id == created.json()["id"]))]
        assert snapshots and all(p["display_name"] == "QA Preparer" for p in snapshots)
        signed = [(v.id, v.snapshot_sha256, v.signing_sha256, v.docx_sha256, v.pdf_sha256)
                  for v in db.scalars(select(OfficialDocumentVersion))]
    assert client.patch(f"/api/users/{created.json()['id']}", headers=auth_headers,
                        json={"middle_name": "Later"}).status_code == 200
    with session_factory() as db:
        assert hashes == {d.id: hashlib.sha256(d.content).hexdigest() for d in db.scalars(
            select(GeneratedDocument).where(GeneratedDocument.part_request_id == item["id"]))}
        assert signed == [(v.id, v.snapshot_sha256, v.signing_sha256, v.docx_sha256, v.pdf_sha256)
                          for v in db.scalars(select(OfficialDocumentVersion))]
        assert all(p.identity_snapshot["display_name"] == "QA Preparer" for p in db.scalars(
            select(DocumentParticipant).where(DocumentParticipant.user_id == created.json()["id"])))
